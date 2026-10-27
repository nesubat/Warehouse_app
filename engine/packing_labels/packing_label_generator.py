import re
import functools
from collections import Counter
import openpyxl
import io
import math
import base64
import zipfile
try:  # the xlsx parts are read from uploaded files: defusedxml refuses entity tricks (billion laughs)
    from defusedxml import ElementTree as ET
except ImportError:  # pragma: no cover - requirements.txt installs it
    import xml.etree.ElementTree as ET
import pymupdf as fitz
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
from engine.label_shuffler.pdf_engine import code128_modules, _draw_code128, BARCODE_QUIET_MODULES, BARCODE_MAX_MODULE

def extract_rich_value_images(excel_path, sheet_name):
    ns = {
        'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
        'xlrd': 'http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel'
    }
    image_map = {}
    try:
        with zipfile.ZipFile(excel_path, 'r') as z:
            wb_xml = z.read('xl/workbook.xml')
            wb_root = ET.fromstring(wb_xml)
            sheet_id = None
            for sheet in wb_root.findall('.//main:sheet', ns):
                if sheet.get('name') == sheet_name:
                    sheet_id = sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                    break
            if not sheet_id: return {}
            
            wb_rels_xml = z.read('xl/_rels/workbook.xml.rels')
            wb_rels_root = ET.fromstring(wb_rels_xml)
            sheet_path = None
            for rel in wb_rels_root.findall('.//{http://schemas.openxmlformats.org/package/2006/relationships}Relationship'):
                if rel.get('Id') == sheet_id:
                    sheet_path = 'xl/' + rel.get('Target')
                    break
            if not sheet_path: return {}
            
            sheet_xml = z.read(sheet_path)
            sheet_root = ET.fromstring(sheet_xml)
            cell_vm_map = {} 
            for row in sheet_root.findall('.//main:row', ns):
                r_idx = int(row.get('r'))
                for c in row.findall('main:c', ns):
                    vm = c.get('vm')
                    if vm:
                        ref = c.get('r')
                        col_letter = "".join(filter(str.isalpha, ref))
                        cell_vm_map[(r_idx, col_letter)] = vm
            if not cell_vm_map: return {}
                
            meta_xml = z.read('xl/metadata.xml')
            meta_root = ET.fromstring(meta_xml)
            vm_nodes = meta_root.find('.//main:valueMetadata', ns)
            vm_to_rvb = {}
            if vm_nodes is not None:
                for i, bk in enumerate(vm_nodes.findall('main:bk', ns)):
                    rc = bk.find('main:rc', ns)
                    if rc is not None: vm_to_rvb[str(i+1)] = rc.get('v')
                        
            rv_xml = z.read('xl/richData/richValueRel.xml')
            rv_root = ET.fromstring(rv_xml)
            rvb_to_rid = {}
            for i, rel in enumerate(rv_root.findall('.//xlrd:rel', ns)):
                rvb_to_rid[str(i)] = rel.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                
            rv_rels_xml = z.read('xl/richData/_rels/richValueRel.xml.rels')
            rv_rels_root = ET.fromstring(rv_rels_xml)
            rid_to_path = {}
            for rel in rv_rels_root.findall('.//{http://schemas.openxmlformats.org/package/2006/relationships}Relationship'):
                rid_to_path[rel.get('Id')] = rel.get('Target').replace('../', 'xl/')
                
            for (r, col_letter), vm in cell_vm_map.items():
                rvb = vm_to_rvb.get(vm)
                if rvb:
                    rid = rvb_to_rid.get(rvb)
                    if rid:
                        media_path = rid_to_path.get(rid)
                        if media_path:
                            col_idx = column_index_from_string(col_letter)
                            image_map[(r, col_idx)] = io.BytesIO(z.read(media_path))
    except Exception:
        pass
    return image_map

EMU_PER_PT = 12700
EMU_PER_PX = 9525


def _row_height_emu(ws, row):
    h = ws.row_dimensions[row].height if row in ws.row_dimensions else None
    return (h or ws.sheet_format.defaultRowHeight or 15) * EMU_PER_PT


def _col_width_emu(ws, col):
    letter = get_column_letter(col)
    w = ws.column_dimensions[letter].width if letter in ws.column_dimensions else None
    w = w or ws.sheet_format.defaultColWidth or 8.43
    return (w * 7 + 5) * EMU_PER_PX


def _cell_at(offset_emu, size_of):
    """1-based row/col index containing a sheet-absolute EMU offset."""
    idx, edge = 1, 0
    while edge + size_of(idx) <= offset_emu and idx < 100000:
        edge += size_of(idx)
        idx += 1
    return idx


def _start_emu(index0, offset, size_of):
    return sum(size_of(i) for i in range(1, index0 + 1)) + offset


def _floating_image_cell(ws, anchor):
    """Cell holding the image's centre, so an image that pokes into a neighbouring row still maps to its own row."""
    row_size = lambda r: _row_height_emu(ws, r)
    col_size = lambda c: _col_width_emu(ws, c)
    top = _start_emu(anchor._from.row, anchor._from.rowOff, row_size)
    left = _start_emu(anchor._from.col, anchor._from.colOff, col_size)
    if getattr(anchor, 'to', None) is not None:
        bottom = _start_emu(anchor.to.row, anchor.to.rowOff, row_size)
        right = _start_emu(anchor.to.col, anchor.to.colOff, col_size)
    elif getattr(anchor, 'ext', None) is not None:
        bottom = top + anchor.ext.height
        right = left + anchor.ext.width
    else:
        return anchor._from.row + 1, anchor._from.col + 1
    return _cell_at((top + bottom) / 2, row_size), _cell_at((left + right) / 2, col_size)


def _row_ranges(rows):
    """[5, 6, 7, 10] -> '5-7, 10'"""
    parts, start, prev = [], None, None
    for r in sorted(rows):
        if start is None:
            start = prev = r
        elif r == prev + 1:
            prev = r
        else:
            parts.append(f"{start}-{prev}" if prev > start else str(start))
            start = prev = r
    if start is not None:
        parts.append(f"{start}-{prev}" if prev > start else str(start))
    return ", ".join(parts)


ADDRESS_FIELDS = ('address_1', 'address_2', 'suburb', 'state', 'postcode', 'country')
INSTALL_YES = {'y', 'yes', 'true'}


def is_install_flag(value):
    return str(value or '').strip().lower() in INSTALL_YES


def normalize_postcode(postcode, country=''):
    """Australian and New Zealand postcodes are always 4 digits. Excel drops the leading zero
    (0803 -> 803, or 803.0), and the courier portal rejects a 3-digit postcode, so pad it back."""
    value = str(postcode if postcode is not None else '').strip()
    if value.endswith('.0') and value[:-2].isdigit():
        value = value[:-2]
    if value.isdigit() and len(value) == 3 and str(country or 'AU').strip().upper() in ('AU', 'NZ', 'AUSTRALIA', 'NEW ZEALAND'):
        value = value.zfill(4)
    return value


def _norm(text):
    """Comparison key: ignores capitals, spacing and punctuation ('12 Bourke Rd.' == '12  BOURKE RD')."""
    return " ".join(re.sub(r"[^0-9a-z]+", " ", str(text or '').lower()).split())


NOT_A_PLACE = ('email', 'e-mail', 'phone', 'mobile', 'contact', 'attn', 'attention', 'fax')
STORE_WORDS = ('store', 'retailer', 'outlet', 'site name')
RECEIVER_WORDS = ('receiver', 'consignee', 'ship to', 'deliver to', 'delivery name', 'recipient')
# 'Store Type', 'Store Size', 'Store Code', 'Receiver No' … describe the store or receiver: they aren't its name
NOT_A_NAME = {'type', 'size', 'code', 'number', 'no', 'num', 'id', 'format', 'group', 'level', 'cluster', 'region',
              'manager', 'email', 'phone', 'mobile', 'count', 'status', 'category', 'tier', 'class', 'grade', 'date',
              'channel', 'owner', 'brand', 'area', 'zone', 'hours', 'open', 'opening'}


# A dimension split over several columns (vertical) or rows (horizontal): which part a header or row label is, and the
# letter it's printed with. Whole words (misspellings too) anywhere in the label; a bare letter or 'wide' / 'long' …
# only when that's all the label says once a unit and 'size' / 'dimension' are taken off ('Size W (mm)' is W).
DIM_PARTS = (  # letter, words that name it anywhere, words that name it only on their own
    ('W', r'width|widht|wdth', r'w|wide'),
    ('B', r'breadth|bredth|breath', r'b'),
    ('L', r'length|lenght|lengh|lgth', r'l|long'),
    ('H', r'height|hight|heigth|hgt', r'h|high|tall'),
    ('D', r'depth|dept', r'd|deep'),
)
DIM_UNIT = re.compile(r'\b(mm|cm|m|in|inch|inches)\b', re.I)


def dim_part(label):
    """'W', 'B', 'L', 'H' or 'D' for a header or row label that names one part of a dimension, else None:
    'Width', 'Item Length (mm)', 'Bredth', 'Size W', 'L' … but not 'Description' or 'Long sleeve tee'."""
    text = " ".join(str(label or '').lower().replace('_', ' ').split())
    if not text:
        return None
    for letter, words, alone in DIM_PARTS:
        if re.search(rf'\b(?:{words})\b', text):
            return letter
    bare = re.sub(r'[()\[\]:.,/\-]', ' ', DIM_UNIT.sub(' ', text))
    bare = re.sub(r'\b(size|sizes|dim|dims|dimension|dimensions|item|in)\b', ' ', bare).strip()
    for letter, _, alone in DIM_PARTS:
        if re.fullmatch(alone, bare):
            return letter
    return None


def dim_unit(label):
    """The unit a header or row label gives its values ('Width (mm)' -> 'mm'), or ''."""
    m = DIM_UNIT.search(str(label or ''))
    return m.group(1).lower() if m else ''


def join_dimensions(parts):
    """One line for the label from the parts found: {'W': ('600', 'mm'), 'L': ('900', 'mm')} -> 'W 600mm × L 900mm',
    in the order W, B, L, H, D. A unit from the label goes onto a plain number; a value that has its own stays as is."""
    out = []
    for letter, _, _ in DIM_PARTS:
        value, unit = parts.get(letter, ('', ''))
        value = str(value or '').strip()
        if not value:
            continue
        if unit and re.fullmatch(r'\d+(?:[.,]\d+)?', value):
            value += unit
        out.append(f"{letter} {value}")
    return " × ".join(out)


def _header_field(header):
    """Which column a header is. Three names: 'Store Name' / 'Store' / 'Retailer' is the store (it names the label
    and starts the Item Reference); 'Receiver Name' / 'Receiver' / 'Consignee' / 'Ship To' is who the courier
    delivers to; 'Attn' / 'Attention' / 'Contact' the person to attend. 'Address Line 1', 'Receiver Address Line 1'
    or 'Store Address' is the street: address parts are checked before the name words, so 'Receiver Suburb' is the
    suburb, not the receiver. A name header with a describing word ('Store Type', 'Store Code') isn't a name."""
    h = " ".join(str(header).strip().lower().replace('_', ' ').split())
    words = set(re.findall(r'[a-z]+', h))
    place_ok = not any(w in h for w in NOT_A_PLACE)
    if words & {'attn', 'atnn', 'attention'} or h in ('contact', 'contact name', 'contact person', 'receiver contact',
                                                        'receiver contact name', 'attention to'):
        return 'contact'
    if 'packing spec' in h: return 'packing_spec'
    if 'job' in h: return 'job_no'
    if 'qty' in h or 'quantity' in h: return 'qty'
    if 'desc' in h: return 'desc'
    if 'thumb' in h or 'image' in h or 'picture' in h or 'art' in words or 'artwork' in words: return 'thumbnail'
    if 'dimension' in h: return 'dim_combined'
    part = dim_part(h)
    if part: return f"dim_{part.lower()}"  # Width / Breadth / Length / Height / Depth (dim_part)
    if place_ok and ('address' in h or 'street' in h):
        return 'address_2' if re.search(r'(line|address|addr)\s*2\b', h) else 'address_1'
    if place_ok and ('suburb' in h or 'town' in words or 'city' in words or 'locality' in words): return 'suburb'
    if place_ok and 'state' in words: return 'state'
    if place_ok and ('postcode' in h or 'post code' in h or 'postal code' in h or 'zip' in words): return 'postcode'
    if place_ok and 'country' in h: return 'country'
    if 'material' in h: return 'material'
    if 'note' in h: return 'notes'
    if 'install' in h: return 'install'
    describes = bool(words & NOT_A_NAME) and 'name' not in words
    if place_ok and not describes and any(w in h for w in RECEIVER_WORDS): return 'receiver_name'
    if place_ok and not describes and any(w in h for w in STORE_WORDS): return 'store_name'
    return None


# The columns the preview's Column mapping can set by hand: (field, label, required)
MAPPABLE_FIELDS = (
    ('packing_spec', 'Packing Spec', True), ('job_no', 'Job Number', True), ('install', 'Install', True),
    ('store_name', 'Store Name', False), ('receiver_name', 'Receiver Name', False), ('contact', 'Attn', False),
    ('qty', 'Quantity', False), ('desc', 'Description', False),
    ('thumbnail', 'Thumbnail', False), ('dim_combined', 'Dimensions', False), ('dim_w', 'Width', False),
    ('dim_b', 'Breadth', False), ('dim_l', 'Length', False), ('dim_h', 'Height', False), ('dim_d', 'Depth', False), ('address_1', 'Address Line 1', False), ('address_2', 'Address Line 2', False),
    ('suburb', 'Suburb', False), ('state', 'State', False), ('postcode', 'Postcode', False),
    ('country', 'Country', False), ('material', 'Material', False), ('notes', 'Notes', False),
)


def detect_columns(headers):
    """{field: column number} from the header row's cells [(column number, text)], by their names (_header_field).
    A later matching column wins."""
    cols = {}
    for col_idx, val in headers:
        if not val:
            continue
        field = _header_field(str(val))
        if field:
            cols[field] = col_idx
    return cols


def apply_column_choices(cols, choices):
    """The auto-detected columns with the ones set by hand in the preview on top. choices: {field: column letter},
    '-' for "not used", '' (or missing) for automatic."""
    cols = dict(cols)
    known = {f for f, _, _ in MAPPABLE_FIELDS}
    for field, letter in (choices or {}).items():
        letter = str(letter or '').strip().upper()
        if field not in known or not letter:
            continue
        if letter == '-':
            cols.pop(field, None)
            continue
        try:
            cols[field] = column_index_from_string(letter)
        except ValueError:
            pass
    return cols


def header_columns(excel_path, header_row, sheet_name=None):
    """The header row's cells for the preview's Column mapping: [(column letter, header text)], empty cells left
    out. Read without loading the whole sheet."""
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ()) if header_row >= 1 else ()
        return [(get_column_letter(i), " ".join(str(v).split())) for i, v in enumerate(row, start=1)
                if v is not None and str(v).strip()]
    finally:
        wb.close()


class PackCheckError(ValueError):
    """Blocking problems found in the sheet; carries the structured issues for the preview table."""
    def __init__(self, issues, warnings):
        self.issues, self.warnings = issues, warnings
        super().__init__("\n".join(_issue_line(i) for i in issues))


def _issue(rows, col, text, detail='', major=False):
    """A problem found in the file, with its Excel rows and column. major: generation goes on, but labels, units or
    consignments go missing or don't add up (a Packing Spec with no quantities, no address, no barcode): listed first,
    in red, in the preview's To check list."""
    return {'rows': rows, 'col': col, 'text': text, 'detail': detail, 'major': major}


def group_issues(issues):
    """A tab's problems as the page lists them: the same problem once, with every place it's at ('Packing Spec but
    no quantities: no label made' on fifteen rows is one line, 'Pack 1: rows 11, 13, 19 …'), major ones first.
    [{'text', 'detail', 'major', 'count', 'where': ['Pack 1: rows 11, 13', …]}]"""
    groups = {}
    for i in issues or []:
        g = groups.setdefault((i['text'], i.get('detail', '')), {'text': i['text'], 'detail': i.get('detail', ''),
                                                                 'major': bool(i.get('major')), 'count': 0, 'cols': {}})
        g['count'] += 1
        rows = g['cols'].setdefault(i.get('col') or '', [])
        if i.get('stores'):  # each row with its store: '11 (Macpac Adelaide)'
            rows += [f"{r} ({name})" for r, name in i['stores']]
        elif i.get('rows'):
            rows.append(str(i['rows']))
    out = []
    for g in groups.values():
        where = []
        for col, rows in g.pop('cols').items():
            rows_text = f"row{'s' if len(rows) > 1 or any(c in r for r in rows for c in ',-') else ''} {', '.join(rows)}" if rows else ''
            where.append(": ".join(x for x in (col, rows_text) if x))
        out.append({**g, 'where': [w for w in where if w]})
    return sorted(out, key=lambda g: not g['major'])


def _rows_in(text):
    """'5-7, 10' -> [5, 6, 7, 10]: the rows an issue names (the other way from _row_ranges)."""
    rows = []
    for part in re.split(r'[,;]', str(text or '')):
        m = re.fullmatch(r'\s*(\d+)\s*(?:-\s*(\d+))?\s*', part)
        if m:
            a, b = int(m.group(1)), int(m.group(2) or m.group(1))
            if 0 <= b - a <= 2000:
                rows += range(a, b + 1)
    return rows


def with_stores(issues, store_of):
    """Puts the store of each row an issue names with it, so a problem reads 'Row 11 Macpac Adelaide', not only
    'Row 11': issue['stores'] = [[row, store], ...]. store_of(row) -> the store's name in the file ('' when none)."""
    for i in issues:
        found = []
        for r in _rows_in(i.get('rows')):
            try:
                name = " ".join(str(store_of(r) or '').split())
            except Exception:
                name = ''
            if name:
                found.append([r, name])
        if found:
            i['stores'] = found
    return issues


def issue_where(i, limit=6):
    """Where a problem is, with its stores: 'Row 11 Macpac Adelaide', 'Rows 11-13: Macpac Adelaide, Macpac Belconnen,
    Macpac Bondi', or 'Row 11' when the file names no store there. '' when it names no rows."""
    rows, stores = str(i.get('rows') or ''), i.get('stores') or []
    if not rows:
        return ''
    several = ',' in rows or '-' in rows
    if not stores:
        return f"Row{'s' if several else ''} {rows}"
    if not several:
        return f"Row {rows} {stores[0][1]}"
    names = list(dict.fromkeys(name for _, name in stores))
    more = f" and {len(names) - limit} more" if len(names) > limit else ''
    return f"Rows {rows}: {', '.join(names[:limit])}{more}"


def _issue_line(i):
    return " · ".join(x for x in (issue_where(i), i['col'], i['text']) if x)


def issue_text(i):
    """One problem as one readable line for the log (and the Logs page): where (rows and stores), the column, what's
    wrong, then the detail: 'Row 52 Corio Village · Pack 1 · marked x but has quantities — 4 units in L52'."""
    main = " · ".join(str(x) for x in (issue_where(i), i.get('col'), i.get('text')) if x)
    detail = str(i.get('detail') or '').strip()
    return f"{main} — {detail}" if detail else main


def _span(rows):
    return _row_ranges(list(range(rows[0]['row'], rows[-1]['row'] + 1)))


def _short(text, limit=40):
    return text if len(text) <= limit else text[:limit - 1].rstrip() + "…"


STREET_TYPES = {
    'street': 'st', 'road': 'rd', 'avenue': 'ave', 'boulevard': 'blvd', 'highway': 'hwy',
    'drive': 'dr', 'parade': 'pde', 'place': 'pl', 'court': 'ct', 'crescent': 'cres',
    'lane': 'ln', 'terrace': 'tce', 'arcade': 'arc',
}


def _address_key(text):
    """Like _norm, plus 'Street' == 'St' etc. Only after a name, so '12 St Kilda Rd' keeps its Saint."""
    words = _norm(text).split()
    return " ".join(
        STREET_TYPES.get(w, w) if i and words[i - 1].isalpha() else w
        for i, w in enumerate(words)
    )


STATE_WORDS = {'vic', 'nsw', 'qld', 'sa', 'wa', 'tas', 'nt', 'act', 'victoria', 'queensland', 'tasmania', 'australia', 'au'}
UNIT_WORDS = {'shop', 'unit', 'level', 'lvl', 'suite', 'lot', 'tenancy', 'kiosk', 'store', 'building', 'bldg', 'floor'}


def similar_addresses(a, b):
    """True when two address texts are probably one place spelled differently: '12 Smith St, Richmond VIC 3121' ~
    '12 Smith Street Richmond 3121' ~ 'Shop 1, 12 Smith St, Richmond' ~ '12 Smyth St, Richmond VIC 3121'.
    The postcodes must match (when both have one), the street numbers mustn't contradict (12 vs 14 is another
    place; a shop number on one side only is fine), and the street and suburb words must be close (a typo or
    two is fine, 'Smith' vs 'Jones' isn't)."""
    from difflib import SequenceMatcher
    def parts(text):
        words = _address_key(text).split()
        postcode = next((w for w in reversed(words) if re.fullmatch(r'\d{4}', w)), '')
        numbers = {w for w in words if any(ch.isdigit() for ch in w) and w != postcode}
        names = [w for w in words if w.isalpha() and w not in STATE_WORDS and w not in UNIT_WORDS
                 and w not in STREET_TYPES.values()]
        return postcode, numbers, " ".join(names)
    pa, na, wa = parts(a)
    pb, nb, wb = parts(b)
    if pa and pb and pa != pb:
        return False
    if (na - nb) and (nb - na):
        return False  # each has a number the other lacks: another street number or another shop
    if not wa or not wb:
        return False
    return wa == wb or SequenceMatcher(None, wa, wb).ratio() >= 0.8


def _distinct(rows, key):
    found = {}
    for r in rows:
        if r[key]:
            found.setdefault(r[key], []).append(r['row'])
    return found


def _variants(rows, key, display):
    """'4-5: 49 Church Street… · 58: 49 Church St…'"""
    by_value = {}
    for r in rows:
        if r[key]:
            by_value.setdefault(r[key], (r[display], []))[1].append(r['row'])
    # Skip the leading words a variant shares with the first one, so the detail shows where they differ
    def shared(a, b):
        n = 0
        while n < min(len(a), len(b)) and _norm(a[n]) == _norm(b[n]):
            n += 1
        return max(0, n - 1)

    variants = [(text.split(), rs) for text, rs in by_value.values()]
    ref = variants[0][0]
    parts = []
    for idx, (words, rs) in enumerate(variants):
        same = max(shared(ref, w) for w, _ in variants[1:]) if idx == 0 else shared(ref, words)
        parts.append(f"{_row_ranges(rs)}: {'…' if same else ''}{_short(' '.join(words[same:]))}")
    return " · ".join(parts)


def _check_pack_consistency(pack_groups, pack_rows, has_address, col):
    errors, warnings = [], []
    no_address_packs = []

    for pack_id, rows in pack_rows.items():
        span = _span(rows)
        pack = pack_groups[pack_id]['pack_spec_name']

        stores = _distinct(rows, 'store_key')
        if len(stores) > 1:
            errors.append(_issue(span, col['store_name'], f"Store name differs in pack {pack}", _variants(rows, 'store_key', 'store')))
        elif stores and any(not r['store_key'] for r in rows):
            blank = [r['row'] for r in rows if not r['store_key']]
            errors.append(_issue(_row_ranges(blank), col['store_name'], f"Store name blank in pack {pack}"))

        receivers = _distinct(rows, 'receiver_key')
        if len(receivers) > 1 and col['receiver_name'] != col['store_name']:
            errors.append(_issue(span, col['receiver_name'], f"Receiver name differs in pack {pack}", _variants(rows, 'receiver_key', 'receiver')))

        if len({r['install'] for r in rows}) > 1:
            flagged = [r['row'] for r in rows if r['install']]
            errors.append(_issue(span, col['install'], f"Install mixed in pack {pack}", f"Y on {_row_ranges(flagged)}"))

        if has_address:
            addresses = _distinct(rows, 'address_key')
            if len(addresses) > 1:
                errors.append(_issue(span, col['address'], f"Address differs in pack {pack}", _variants(rows, 'address_key', 'address')))
            elif addresses and any(not r['address_key'] for r in rows):
                blank = [r['row'] for r in rows if not r['address_key']]
                errors.append(_issue(_row_ranges(blank), col['address'], f"Address blank in pack {pack}"))
            elif not addresses:
                no_address_packs.append(rows[0]['row'])

            if any(r['crosses_pack'] for r in rows):
                warnings.append(_issue(span, col['address'], "Merged address shared with next pack"))

    if no_address_packs:
        warnings.append(_issue(_row_ranges(no_address_packs), col['address'], "No address", major=True))

    # Across packs of the same store: store deliveries must agree, installer deliveries may differ
    if has_address:
        by_store = {}
        for rows in pack_rows.values():
            first = rows[0]
            if first['store_key'] and first['address_key']:
                by_store.setdefault(first['store_key'], []).append(first)
        for store_rows in by_store.values():
            for install in (False, True):
                group = [r for r in store_rows if r['install'] == install]
                if len(_distinct(group, 'address_key')) > 1:
                    rows_text = _row_ranges([r['row'] for r in group])
                    name = group[0]['store']
                    if install:
                        warnings.append(_issue(rows_text, col['address'], f"{name}: installer addresses differ", _variants(group, 'address_key', 'address')))
                    elif all(similar_addresses(group[0]['address'], r['address']) for r in group[1:]):
                        # The same place written differently: the consignment table offers to merge them into one
                        warnings.append(_issue(rows_text, col['address'], f"{name}: store address written differently between packs "
                                               f"- merge them into one in the consignment table", _variants(group, 'address_key', 'address')))
                    else:
                        warnings.append(_issue(rows_text, col['address'], f"{name}: store address differs between packs - they're side by "
                                               f"side in the consignment table: merge them there if they're one place", _variants(group, 'address_key', 'address')))

    return errors, warnings


def parse_packing_data(excel_path, header_row, sheet_name=None, columns=None):
    """columns: the preview's Column mapping, {field: column letter or '-'} on top of the columns found by their
    header names (apply_column_choices)."""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    
    # Every image in a cell is kept: {(row, col): [(reading-order key, image bytes), ...]}
    image_map = {cell: [((0, 0, 0, 0), data)] for cell, data in extract_rich_value_images(excel_path, ws.title).items()}

    for image in getattr(ws, '_images', []):
        try:
            r, c, order = None, None, (0, 0, 0, 0)
            if hasattr(image, 'anchor'):
                if hasattr(image.anchor, '_from'):
                    r, c = _floating_image_cell(ws, image.anchor)
                    f = image.anchor._from
                    order = (f.row, f.rowOff, f.col, f.colOff)  # top-to-bottom, then left-to-right
                elif isinstance(image.anchor, str):
                    c_letter, r_str = coordinate_from_string(image.anchor)
                    c = column_index_from_string(c_letter)
                    r = int(r_str)

            if r and c:
                img_bytes = None
                if hasattr(image, 'ref') and hasattr(image.ref, 'getvalue'):
                    img_bytes = image.ref.getvalue()
                elif hasattr(image, '_data'):
                    img_bytes = image._data() if callable(image._data) else image._data

                if img_bytes:
                    image_map.setdefault((r, c), []).append((order, io.BytesIO(img_bytes)))
        except Exception:
            continue

    cols = apply_column_choices(
        detect_columns([(i, ws.cell(row=header_row, column=i).value) for i in range(1, ws.max_column + 1)]), columns)

    found_headers = {
        'Packing Spec': 'packing_spec' in cols,
        'Store Name': 'store_name' in cols or 'receiver_name' in cols,
        'Job Number': 'job_no' in cols,
        'Quantity': 'qty' in cols,
        'Description': 'desc' in cols,
        'Thumbnail': 'thumbnail' in cols,
        'Dimensions': any(f in cols for f in ('dim_combined', 'dim_w', 'dim_b', 'dim_l', 'dim_h', 'dim_d')),
        'Address Info': any(k in cols for k in ADDRESS_FIELDS),
        'Material': 'material' in cols,
        'Notes': 'notes' in cols,
        'Install': 'install' in cols
    }

    if 'packing_spec' not in cols:
        raise ValueError(f"Could not find a 'Packing Spec' column in Row {header_row}. Pick it in the Column mapping below.")
    if 'job_no' not in cols:
        raise ValueError(f"Could not find a 'Job Number' column in Row {header_row}. Pick it in the Column mapping below.")
    if 'install' not in cols:
        raise ValueError(f"Could not find an 'Install' column in Row {header_row}. Pick it in the Column mapping below, or add one: Y for packs going to an installer, N or blank for the store.")

    def get_cell_info(row, col):
        for merged_range in ws.merged_cells.ranges:
            if merged_range.min_col <= col <= merged_range.max_col and merged_range.min_row <= row <= merged_range.max_row:
                val = ws.cell(row=merged_range.min_row, column=merged_range.min_col).value
                return val, merged_range.min_row, merged_range.max_row
        return ws.cell(row=row, column=col).value, row, row

    pack_groups = {}
    pack_rows = {}
    rows_missing_job = []
    rows_no_barcode = []
    rows_without_image = []
    address_keys = [k for k in ADDRESS_FIELDS if k in cols]
    col_ref = {k: f"{get_column_letter(c)} · {str(ws.cell(row=header_row, column=c).value or '').strip().title()}" for k, c in cols.items()}
    addr_letters = [get_column_letter(cols[k]) for k in address_keys]
    col_ref['address'] = f"{addr_letters[0]}–{addr_letters[-1]} · Address" if len(addr_letters) > 1 else (f"{addr_letters[0]} · Address" if addr_letters else "Address")
    # No Store Name column: the Receiver Name stands in for it (names the label, starts the Item Reference)
    col_ref.setdefault('store_name', col_ref.get('receiver_name', "Store Name"))
    col_ref.setdefault('receiver_name', "Receiver Name")
    col_ref.setdefault('install', "Install")
    col_ref.setdefault('thumbnail', "Thumbnail")
    current_row = header_row + 1
    last_processed_row = current_row

    while True:
        pack_spec_val, pack_top_row, pack_bottom_row = get_cell_info(current_row, cols['packing_spec'])

        if pack_spec_val is None or str(pack_spec_val).strip() == "": break

        packing_spec = str(pack_spec_val).strip()
        pack_id = f"ROW_{pack_top_row}_{packing_spec}"

        def get_val(key):
            if key in cols:
                val, _, _ = get_cell_info(current_row, cols[key])
                return str(val).strip() if val is not None else ""
            return ""

        receiver_name = get_val('receiver_name')
        store_name = get_val('store_name') or receiver_name  # no store name: the receiver's names the label
        address_parts = {k: get_val(k) for k in address_keys}
        install = is_install_flag(get_val('install'))

        # Address cells merged past this pack's edges (shared with a neighbouring pack)
        crosses_pack = any(
            lo < pack_top_row or hi > pack_bottom_row
            for _, lo, hi in (get_cell_info(current_row, cols[k]) for k in address_keys)
        )

        if pack_id not in pack_groups:
            pack_groups[pack_id] = {
                'pack_spec_name': packing_spec,
                'store_name': store_name,
                'receiver_name': receiver_name,  # who the courier delivers to; else from the address, else the store
                'contact': get_val('contact'),   # the Attn column
                'install': install,
                **{k: address_parts.get(k, '') for k in ADDRESS_FIELDS},
                'excel_rows': [],
                'items': []
            }
            pack_rows[pack_id] = []

        pack_rows[pack_id].append({
            'row': current_row,
            'store': store_name,
            'store_key': _norm(store_name),
            'receiver': receiver_name,
            'receiver_key': _norm(receiver_name),
            'address': ", ".join(v for v in address_parts.values() if v),
            'address_key': _address_key(" ".join(address_parts.values())),
            'install': install,
            'crosses_pack': crosses_pack,
        })

        # One Dimensions column as written; else the parts in their own columns, each with its letter and the unit
        # its header gives: 'W 600mm × H 900mm' (join_dimensions, the same as horizontal files' rows)
        dimension = get_val('dim_combined') or join_dimensions({
            f[4:].upper(): (get_val(f), dim_unit(ws.cell(row=header_row, column=cols[f]).value))
            for f in ('dim_w', 'dim_b', 'dim_l', 'dim_h', 'dim_d') if f in cols})

        thumbnails = []
        if 'thumbnail' in cols:
            tc = cols['thumbnail']
            # Images come from the thumbnail cell itself, never a neighbouring cell. A merged
            # thumbnail cell is one cell: every row it covers gets all the images inside it.
            area = next((m for m in ws.merged_cells.ranges
                         if m.min_row <= current_row <= m.max_row and m.min_col <= tc <= m.max_col), None)
            cells = ([(r, c) for r in range(area.min_row, area.max_row + 1) for c in range(area.min_col, area.max_col + 1)]
                     if area else [(current_row, tc)])
            found = sorted((entry for cell in cells for entry in image_map.get(cell, [])), key=lambda e: e[0])
            thumbnails = [data for _, data in found]
            if not thumbnails:
                rows_without_image.append(current_row)

        item = {
            'job_no': get_val('job_no'),
            'qty': get_val('qty') or "1",
            'desc': get_val('desc'),
            'dimension': dimension,
            'material': get_val('material'),
            'install': "INSTALLER" if install else "",
            'notes': get_val('notes'),
            'thumbnails': thumbnails
        }

        if item['job_no'] and code128_modules(item['job_no']) is None:
            rows_no_barcode.append(current_row)
        if not item['job_no']:
            rows_missing_job.append(current_row)
        pack_groups[pack_id]['items'].append(item)
        pack_groups[pack_id]['excel_rows'].append(current_row)
        last_processed_row = current_row
        current_row += 1

    rows_missing_spec = []
    for row in range(current_row, ws.max_row + 1):
        spec_val, _, _ = get_cell_info(row, cols['packing_spec'])
        job_val, _, _ = get_cell_info(row, cols['job_no'])
        if (spec_val is None or str(spec_val).strip() == "") and job_val is not None and str(job_val).strip():
            rows_missing_spec.append(row)

    wb.close()

    errors = []
    if rows_missing_spec:
        errors.append(_issue(_row_ranges(rows_missing_spec), col_ref['packing_spec'], "Packing Spec empty"))
    if rows_missing_job:
        errors.append(_issue(_row_ranges(rows_missing_job), col_ref['job_no'], "Job Number empty"))
    if not pack_groups and not errors:
        errors.append(_issue('', '', f"No data found below row {header_row}"))

    warnings = []
    if rows_without_image:
        warnings.append(_issue(_row_ranges(rows_without_image), col_ref['thumbnail'], "No image"))
    if rows_no_barcode:
        warnings.append(_issue(_row_ranges(rows_no_barcode), col_ref['job_no'],
                               "No barcode: Job Number has a line break or a character a barcode can't hold", major=True))

    pack_errors, pack_warnings = _check_pack_consistency(pack_groups, pack_rows, bool(address_keys), col_ref)
    errors += pack_errors
    warnings += pack_warnings
    if 'store_name' in cols or 'receiver_name' in cols:  # 'Row 11 Macpac Adelaide', not only 'Row 11'
        with_stores(errors + warnings, lambda r: get_cell_info(r, cols.get('store_name') or cols['receiver_name'])[0])

    if errors:
        raise PackCheckError(errors, warnings)

    return pack_groups, last_processed_row, found_headers, warnings

TEXT_STYLES = {
    'desc': (9, "helv", (0, 0, 0)),
    'dimension': (9, "hebo", (0, 0, 0)),
    'material': (9, "helv", (0, 0, 0)),
    'install': (9, "hebo", (0.8, 0, 0)),
    'notes': (9, "helv", (0, 0, 0)),
}
LINE = 1.2


def _clean(val):
    val = str(val or '').strip()
    return '' if val.lower() == 'none' else val


@functools.lru_cache(maxsize=200_000)
def _text_len(text, fontname, size):
    """fitz.get_text_length, remembered: the same job numbers, descriptions and words are measured thousands of times
    in a big job (a 214-label file measured 177,000 strings)."""
    return fitz.get_text_length(text, fontname=fontname, fontsize=size)


def _wrap(text, fontname, size, width):
    return list(_wrap_cached(text, fontname, size, width))


@functools.lru_cache(maxsize=50_000)
def _wrap_cached(text, fontname, size, width):
    lines = []
    for paragraph in text.splitlines() or ['']:
        line = ''
        for word in paragraph.split():
            while _text_len(word, fontname, size) > width:
                cut = len(word)
                while cut > 1 and _text_len(word[:cut], fontname, size) > width:
                    cut -= 1
                if line:
                    lines.append(line)
                    line = ''
                lines.append(word[:cut])
                word = word[cut:]
            candidate = f"{line} {word}".strip()
            if line and _text_len(candidate, fontname, size) > width:
                lines.append(line)
                line = word
            else:
                line = candidate
        if line:
            lines.append(line)
    return tuple(lines)


def _draw_lines(page, lines, x0, x1, y, fontname, size, color):
    for line in lines:
        w = _text_len(line, fontname, size)
        page.insert_text((x0 + (x1 - x0 - w) / 2, y + size * 0.92), line,
                         fontname=fontname, fontsize=size, color=color)
        y += size * LINE
    return y


# Artwork used on many labels (one file: 53 images on 1,760 item boxes) goes into the PDF once, exactly as it is in the
# Excel file (same pixels, same quality), and every other box points at that same image. Embedding it again for each
# box made a 340 MB PDF that took minutes to write; its size is read once too.
def _art(doc, data):
    """This artwork as this PDF uses it: {'stream', 'size' (px), 'xref' (0 until embedded)}. Kept on the document, so
    two Generates at once never share one."""
    from PIL import Image
    store = getattr(doc, '_wh_art', None)
    if store is None:
        store = doc._wh_art = {}
    art = store.get(id(data))
    if art is None:
        raw = data.getvalue()
        try:
            size = Image.open(io.BytesIO(raw)).size  # the header only: nothing decoded or changed
        except Exception:
            pix = fitz.Pixmap(raw)  # a format only PyMuPDF reads
            size = (pix.width, pix.height)
        art = store[id(data)] = {'data': data, 'stream': raw, 'size': size, 'xref': 0}  # 'data' keeps its id its own
    return art


def _fit_rect(slot, size):
    """The largest rect with the image's own proportions (size: width, height) that fits inside slot, centred in it."""
    scale = min(slot.width / size[0], slot.height / size[1])
    w, h = size[0] * scale, size[1] * scale
    x0, y0 = slot.x0 + (slot.width - w) / 2, slot.y0 + (slot.height - h) / 2
    return fitz.Rect(x0, y0, x0 + w, y0 + h)


def _layout_cell(item, attribute_order, cell, text_scale, thumb_scale, draw_page=None, slots=None, heights=None, _y=None):
    """Stacks the chosen blocks top-down; returns the total height used. Draws only when draw_page is given.
    slots ({block: height}, plan_cells): each block takes the same height in every item box of the job, empty or not,
    so job numbers, barcodes and quantities sit at the same place in every box. heights: a dict filled with what
    each block of this item needs (plan_cells reads it)."""
    pad = 4
    inner_x0, inner_x1 = cell.x0 + pad, cell.x1 - pad
    inner_w = inner_x1 - inner_x0
    gap = 3 * text_scale
    y = cell.y0 + pad if _y is None else _y
    for attr in attribute_order:
        if slots is not None:
            if not slots.get(attr):
                continue  # no item of the job has this block
            top = y
            _layout_cell(item, [attr], cell, text_scale, thumb_scale, draw_page, heights=None, _y=top)
            y = top + slots[attr] + gap
            continue
        start = y
        if attr == 'thumbnail':
            images = item.get('thumbnails') or []
            if not images:
                continue
            h = cell.height * 0.35 * thumb_scale
            if draw_page:
                # One image: centred, 70% of the box width. Several: share the full inner width,
                # up to 3 side by side, then further rows; each keeps its own proportions.
                n = len(images)
                per_row = min(n, 3)
                rows = math.ceil(n / per_row)
                area_w = cell.width * 0.7 if n == 1 else inner_w
                x0 = cell.x0 + (cell.width - area_w) / 2
                slot_w = (area_w - (per_row - 1) * 2) / per_row
                slot_h = (h - (rows - 1) * 2) / rows
                for i, data in enumerate(images):
                    sx = x0 + (i % per_row) * (slot_w + 2)
                    sy = y + (i // per_row) * (slot_h + 2)
                    rect = fitz.Rect(sx, sy, sx + slot_w, sy + slot_h)
                    try:
                        art = _art(draw_page.parent, data)
                        where = _fit_rect(rect, art['size'])
                        if art['xref']:
                            draw_page.insert_image(where, xref=art['xref'])
                        else:
                            art['xref'] = draw_page.insert_image(where, stream=art['stream'])
                    except Exception:
                        draw_page.insert_textbox(rect, "Image Error", fontsize=8, align=1)
            y += h + gap

        elif attr == 'job_no':
            val = _clean(item.get('job_no'))
            if not val:
                continue
            size = 14 * text_scale
            lines = _wrap(val, "hebo", size, inner_w - 4)
            h = len(lines) * size * LINE + 6 * text_scale
            if draw_page:
                _round_rect(draw_page, fitz.Rect(cell.x0 + 3, y, cell.x1 - 3, y + h), 3, fill=(0, 0, 0))
                _draw_lines(draw_page, lines, inner_x0, inner_x1, y + 3 * text_scale, "hebo", size, (1, 1, 1))
            y += h + gap

        elif attr == 'barcode':
            if item.get('is_code'):
                continue  # a packing-sheet code box: nothing to scan
            val = item.get('barcode') or _clean(item.get('job_no'))
            modules = code128_modules(val, compact=True) if val else None
            if not modules:
                continue
            # Code 128 needs 10 blank modules either side; the bars get as wide as the box allows
            module_w = min(BARCODE_MAX_MODULE, inner_w / (len(modules) + 2 * BARCODE_QUIET_MODULES))
            bar_h = 22 * text_scale
            size = 6.5 * text_scale
            h = bar_h + 2 + size * LINE
            if draw_page:
                bx = cell.x0 + (cell.width - len(modules) * module_w) / 2
                _draw_code128(draw_page, bx, y, module_w, bar_h, modules)
                _draw_lines(draw_page, [val], inner_x0, inner_x1, y + bar_h + 2, "helv", size, (0, 0, 0))
            y += h + gap

        elif attr == 'qty':
            val = _clean(item.get('qty'))
            if not val:
                continue
            size = 20 * text_scale
            border = 2.5 * text_scale
            lines = _wrap(val, "hebo", size, inner_w - 2 * border - 8)
            widest = max(_text_len(l, "hebo", size) for l in lines)
            box_w = min(inner_w, max(70 * text_scale, widest + 16 * text_scale))
            h = len(lines) * size * LINE + 6 * text_scale
            if draw_page:
                bx0 = cell.x0 + (cell.width - box_w) / 2
                half = border / 2
                draw_page.draw_rect(fitz.Rect(bx0 + half, y + half, bx0 + box_w - half, y + h - half), color=(0, 0, 0), width=border)
                _draw_lines(draw_page, lines, bx0, bx0 + box_w, y + 3 * text_scale, "hebo", size, (0, 0, 0))
            y += h + gap

        elif attr in TEXT_STYLES or attr.startswith('row_'):
            # A row of the file picked for the item box ('row_5', horizontal files) reads like a description
            val = _clean(item.get(attr))
            if not val:
                continue
            base, fontname, color = TEXT_STYLES.get(attr, TEXT_STYLES['desc'])
            size = base * text_scale
            lines = _wrap(val, fontname, size, inner_w)
            if draw_page:
                _draw_lines(draw_page, lines, inner_x0, inner_x1, y, fontname, size, color)
            y += len(lines) * size * LINE + gap
        if heights is not None and y > start:
            heights[attr] = y - start - gap
    return y - gap - cell.y0 + pad


def _scales(step):
    """Step 0..70 of shrinking: text first (down to 65%), then the thumbnail."""
    s = 1.0 - step * 0.01
    return max(s, 0.65), (1.0 if s >= 0.65 else s / 0.65)


def plan_cells(items, attribute_order, cell):
    """One layout for every item box of the job: each chosen block gets the height the tallest item needs for it, at
    one text and thumbnail scale (the largest at which that stack fits the box). An item without a value for a block
    leaves its space empty, so every block starts at the same place in every box. Returns (text_scale, thumb_scale,
    slots). Found by halving the steps (the stack only gets shorter as the scales shrink)."""
    distinct = list({repr((sorted((k, str(v)) for k, v in it.items() if k != 'thumbnails'), len(it.get('thumbnails') or []))): it
                     for it in items}.values())

    def slots_at(step):
        ts, th = _scales(step)
        slots = {}
        for it in distinct:
            used = {}
            _layout_cell(it, attribute_order, cell, ts, th, heights=used)
            for attr, h in used.items():
                slots[attr] = max(slots.get(attr, 0), h)
        gap = 3 * ts
        total = 8 + sum(h + gap for h in slots.values()) - (gap if slots else 0)
        return ts, th, slots, total

    lo, hi = 0, 70
    if slots_at(0)[3] <= cell.height:
        hi = 0
    else:
        while lo < hi:
            mid = (lo + hi) // 2
            if slots_at(mid)[3] <= cell.height:
                hi = mid
            else:
                lo = mid + 1
    ts, th, slots, _ = slots_at(hi)
    return ts, th, slots


def _draw_cell(page, item, attribute_order, cell, plan=None):
    if plan:  # the job's shared layout (plan_cells)
        ts, th, slots = plan
        _layout_cell(item, attribute_order, cell, ts, th, draw_page=page, slots=slots)
        return
    # On its own: shrink text first (down to 65%), then the thumbnail, until the whole stack fits inside the cell
    text_scale = thumb_scale = 1.0
    for step in range(71):
        text_scale, thumb_scale = _scales(step)
        if _layout_cell(item, attribute_order, cell, text_scale, thumb_scale) <= cell.height:
            break
    _layout_cell(item, attribute_order, cell, text_scale, thumb_scale, draw_page=page)


def _round_rect(page, rect, r, color=None, fill=None, width=0):
    radius = min(0.5, r / max(1, min(rect.width, rect.height)))
    page.draw_rect(rect, color=color, fill=fill, width=width, radius=radius)


def _spec_panel(page, rect, spec, max_size=22, tag=""):
    """Black rounded panel with the Packing Spec in large white type. Returns its bottom y."""
    inner_w = rect.width - 16
    size = max_size
    lines = _wrap(spec, "hebo", size, inner_w)
    while len(lines) > 2 and size > 11:
        size -= 1
        lines = _wrap(spec, "hebo", size, inner_w)
    lines = lines[:3]
    h = 14 + len(lines) * size * LINE + 6
    panel = fitz.Rect(rect.x0, rect.y0, rect.x1, rect.y0 + h)
    _round_rect(page, panel, 6, fill=(0, 0, 0))
    page.insert_text((panel.x0 + 8, panel.y0 + 11), "PACKING SPEC", fontname="hebo", fontsize=7, color=(1, 1, 1))
    if tag:  # horizontal files: the pack and its packing-sheet code, top right of the panel
        tag_size = 9  # its own name: 'size' is the Packing Spec's, drawn below
        while tag_size > 6 and fitz.get_text_length(tag, fontname="hebo", fontsize=tag_size) > panel.width - 90:
            tag_size -= 0.5
        width = fitz.get_text_length(tag, fontname="hebo", fontsize=tag_size)
        page.insert_text((panel.x1 - 8 - width, panel.y0 + 12), tag, fontname="hebo", fontsize=tag_size, color=(1, 0.82, 0.2))
    _draw_lines(page, lines, panel.x0 + 8, panel.x0 + 8 + inner_w, panel.y0 + 15, "hebo", size, (1, 1, 1))
    return panel.y1


def _counter_chip(page, x0, y0, text, size=11):
    w = fitz.get_text_length(text, fontname="hebo", fontsize=size) + 16
    h = size * LINE + 8
    chip = fitz.Rect(x0, y0, x0 + w, y0 + h)
    _round_rect(page, chip, h / 2, color=(0, 0, 0), width=1.5)
    page.insert_text((x0 + 8, y0 + 4 + size * 0.92), text, fontname="hebo", fontsize=size)
    return chip


def barcode_text(job, kind):
    """What an item's barcode encodes: 'J477161-26 K3', or just 'J477161-54' for a job number on one row."""
    return f"{job} K{kind}" if kind else job


def assign_barcodes(pack_groups):
    """Gives every item its barcode text. Same idea as the Distribution Mapper's kinds, with rows in
    place of columns: a job number on more than one row gets one kind per row, numbered K1, K2, K3...
    in reading order (first tab to last, then top to bottom), so K1 is its first row in the first tab.
    An item may carry 'source_key' (horizontal files: its tab and item column): items sharing one are the same
    kind, so an item column keeps its kind on every store's label. Call it with all selected tabs combined, in tab order."""
    items = [item for g in pack_groups.values() for item in g['items']]
    key = lambda item: str(item.get('job_no') or '').strip().upper()
    source = lambda item: item.get('source_key') or id(item)
    kinds_of = {}  # job -> {source: kind}, in reading order
    for item in items:
        if key(item):
            kinds = kinds_of.setdefault(key(item), {})
            kinds.setdefault(source(item), len(kinds) + 1)
    for item in items:
        job = str(item.get('job_no') or '').strip()
        if not job:
            item['barcode'] = ''
            continue
        kinds = kinds_of[key(item)]
        item['barcode'] = barcode_text(job, kinds[source(item)] if len(kinds) > 1 else None)


def assign_label_numbers(pack_groups):
    """'Label X of Y': each store's pack groups numbered in file order (store names compared loosely)."""
    totals = Counter(_norm(g['store_name']) for g in pack_groups.values())
    seen = Counter()
    for g in pack_groups.values():
        key = _norm(g['store_name'])
        seen[key] += 1
        g['label_no'], g['label_total'] = seen[key], totals[key]


# ---------- the PACK AUDIT page: before each pack's labels (horizontal files) or each tab's (vertical files) ----------

AUDIT_GREEN, AUDIT_AMBER = (0.18, 0.63, 0.36), (0.85, 0.55, 0.0)
AUDIT_INK, AUDIT_MUTED, AUDIT_RULE = (0.2, 0.2, 0.2), (0.45, 0.45, 0.45), (0.82, 0.82, 0.82)


def audit_text(page, x, y, text, size=10, bold=False, color=AUDIT_INK, right=None):
    font = 'hebo' if bold else 'helv'
    text = _latin(text)
    if right is not None:
        x = right - _text_len(text, font, size)
    page.insert_text((x, y), text, fontname=font, fontsize=size, color=color)


def audit_fit(text, size, width, bold=False):
    """The text cut to the width, ending in '...' when it had to be cut."""
    font = 'hebo' if bold else 'helv'
    text = _latin(text)
    if _text_len(text, font, size) <= width:
        return text
    while text and _text_len(text + '...', font, size) > width:
        text = text[:-1]
    return text + '...' if text else ''


def spec_summary(specs):
    """'Packing Specs: OB1070170170: 30 labels · 2 x CS450320150: 1 label' from a Counter of Packing Specs, most first."""
    return "Packing Specs: " + " · ".join(f"{spec or '(none)'}: {n} label{'s' if n != 1 else ''}" for spec, n in specs.most_common())


def draw_audit_page(doc, title, figures, notes, issues, all_good, specs=None):
    """A PACK AUDIT page: 'PACK AUDIT' over the title ('Tab 1 - Pack 1', 'Tab 1: VIC'), the figures to check against
    the packing sheet ([(label, value)]), grey notes (lines), then what needs a look, each with its row and store
    ('Row 11 Macpac Adelaide · Pack 1: ...'), as many as fit and then how many more. Green when nothing needs a look
    (all_good says so), amber when something should be checked. specs (a Counter of Packing Specs) are drawn as badges
    under the figures, the way the preview shows them: 'OB1070170170 x 21' (labels with that spec), most first."""
    m = 12 * MM2PT
    colour = AUDIT_AMBER if issues else AUDIT_GREEN
    page = doc.new_page(width=A4_W, height=A4_H)
    page.draw_rect(fitz.Rect(0, 0, A4_W, 54), color=colour, fill=colour)
    audit_text(page, m, 24, "PACK AUDIT", 9, True, (1, 1, 1))
    audit_text(page, m, 44, audit_fit(title, 18, A4_W - 2 * m - 190, bold=True), 18, True, (1, 1, 1))
    audit_text(page, 0, 44, "Check before packing" if issues else "All checks passed", 12, True, (1, 1, 1), right=A4_W - m)
    x = m
    for label, value in figures:
        value = str(value)
        w = max(70, _text_len(value, 'hebo', 15) + 22)
        page.draw_rect(fitz.Rect(x, 68, x + w - 8, 112), color=AUDIT_RULE, fill=(0.97, 0.97, 0.97), radius=0.15, width=0.6)
        audit_text(page, x + 8, 84, label.upper(), 7.5, True, AUDIT_MUTED)
        audit_text(page, x + 8, 104, value, 15, True)
        x += w
    y = 134
    if specs:
        audit_text(page, m, y, "PACKING SPECS", 7.5, True, AUDIT_MUTED)
        y += 8
        x = m
        for spec, n in specs.most_common():
            text = _latin(f"{spec or '(none)'} x {n}")
            w = _text_len(text, 'hebo', 10) + 18
            if x + w > A4_W - m and x > m:  # the next badge goes on a new line
                x, y = m, y + 24
            page.draw_rect(fitz.Rect(x, y, x + w, y + 19), color=AUDIT_RULE, fill=(0.97, 0.97, 0.97), radius=0.3, width=0.6)
            audit_text(page, x + 9, y + 13.5, text, 10, True)
            x += w + 6
        y += 19 + 20
    for line in notes:
        audit_text(page, m, y, audit_fit(line, 9, A4_W - 2 * m), 9, False, AUDIT_MUTED)
        y += 14
    y += 14
    if not issues:
        audit_text(page, m, y, all_good, 9.5, False, AUDIT_GREEN)
        return page
    audit_text(page, m, y, "Needs a look", 11, True, AUDIT_AMBER)
    room = int((A4_H - m - 16 - y) // 15)  # the lines that fit, keeping one for 'and N more'
    for n, i in enumerate(issues):
        if n >= room and len(issues) > room + 1:
            y += 15
            audit_text(page, m, y, f"... and {len(issues) - n} more: see the preview's To check list and the job report",
                       9.5, True, AUDIT_AMBER)
            break
        y += 15
        where = " · ".join(p for p in (issue_where(i), i.get('col', '')) if p)  # 'Row 11 Macpac Adelaide · Pack 1'
        audit_text(page, m, y, audit_fit(f"- {where}: {i.get('text', '')}" if where else f"- {i.get('text', '')}",
                                         9.5, A4_W - 2 * m), 9.5)
    return page


REPORT_COLOURS = {'bad': (0.75, 0.15, 0.13), 'warn': (0.85, 0.55, 0.0), 'info': (0.16, 0.44, 0.64)}


def _latin(text):
    """The built-in PDF fonts only have Latin-1: swap the few other characters the app writes."""
    swaps = {'—': '-', '–': '-', '→': '->', '✓': 'OK', '✏️': '', '’': "'", '‘': "'", '“': '"', '”': '"', '…': '...', '×': 'x'}
    for a, b in swaps.items():
        text = str(text).replace(a, b)
    return text.encode('latin-1', 'replace').decode('latin-1')


def draw_report_pages(doc, report, at=0):
    """Inserts the job report in front of the labels (A4 landscape), as many pages as it needs. Returns the count.

    report: {'title': 'Packing Labels Only', 'job': 'J477161', 'project': 'Lux Test 30', 'when': '05/10/2026 20:25',
             'facts': [(label, value) or (label, value, 'bad'|'warn')...],
             'sections': [{'title', 'level': 'bad'|'warn'|'info', 'items': [text...]}]}
    Compact: the job number in a white-on-black panel with the project beside it, the facts as one row of chips
    (red / amber when they need a look), then each finding as a coloured heading with its items run together."""
    W, H = fitz.paper_size("a4-l")
    M = 30
    pages = []

    def new_page():
        page = doc.new_page(pno=at + len(pages), width=W, height=H)
        pages.append(page)
        page.insert_text((M, M + 10), _latin(f"{report['title'].upper()}  ·  REPORT"), fontname="hebo", fontsize=9,
                         color=(0.35, 0.35, 0.35))
        page.insert_text((W - M - 40, M + 10), _latin(f"Page {len(pages)}"), fontname="helv", fontsize=8, color=(0.5, 0.5, 0.5))
        return page, M + 20

    page, y = new_page()
    # Job number: the thing to check first. The project name beside it needs no caption.
    job = _latin(report.get('job') or '(no job number)')
    size = 30
    job_w = fitz.get_text_length(job, fontname="hebo", fontsize=size) + 30
    panel = fitz.Rect(M, y, M + job_w, y + 48)
    _round_rect(page, panel, 8, fill=(0, 0, 0))
    page.insert_text((M + 8, y + 12), "JOB NUMBER", fontname="hebo", fontsize=6.5, color=(1, 1, 1))
    page.insert_text((M + 15, y + 40), job, fontname="hebo", fontsize=size, color=(1, 1, 1))
    x = panel.x1 + 16
    name = (_wrap(_latin(report.get('project') or ''), "hebo", 20, W - M - x) or [''])[0]
    page.insert_text((x, y + 24), name, fontname="hebo", fontsize=20)
    page.insert_text((x, y + 42), _latin(report.get('when', '')), fontname="helv", fontsize=9, color=(0.35, 0.35, 0.35))
    y = panel.y1 + 10

    # Key facts as chips; the ones that need a look are coloured
    cx = M
    for fact in report.get('facts', []):
        label, value = fact[0], fact[1]
        colour = REPORT_COLOURS.get(fact[2]) if len(fact) > 2 and fact[2] else None
        text = _latin(f"{label}: {value}")
        w = fitz.get_text_length(text, fontname="hebo", fontsize=8.5) + 14
        if cx + w > W - M:
            cx, y = M, y + 20
        _round_rect(page, fitz.Rect(cx, y, cx + w, y + 16), 8, color=colour or (0, 0, 0), width=1.2 if colour else 0.8,
                    fill=colour)
        page.insert_text((cx + 7, y + 11.5), text, fontname="hebo", fontsize=8.5, color=(1, 1, 1) if colour else (0, 0, 0))
        cx += w + 6
    y += 28

    sections = [s for s in report.get('sections', []) if s.get('items')]
    if not sections:
        page.insert_text((M, y + 12), "Nothing needs a look: every label, address and courier detail checked out.",
                         fontname="hebo", fontsize=11, color=(0.12, 0.52, 0.29))
    for section in sections:
        colour = REPORT_COLOURS.get(section.get('level'), REPORT_COLOURS['info'])
        # The items run on as one paragraph, separated by dots, instead of a line each
        body = _wrap(_latin("   ·   ".join(section['items'])), "helv", 8.5, W - 2 * M - 12)
        lines = [(True, _latin(f"{section['title']} ({len(section['items'])})"))] + [(False, t) for t in body]
        top = y
        for bold, text in lines:
            if y + 12 > H - M:
                page.draw_rect(fitz.Rect(M, top, M + 3, y), color=None, fill=colour)
                page, y = new_page()
                top = y
            if bold:
                page.insert_text((M + 9, y + 9.5), text, fontname="hebo", fontsize=9.5, color=colour)
                y += 13
            else:
                page.insert_text((M + 9, y + 9), text, fontname="helv", fontsize=8.5)
                y += 11
        page.draw_rect(fitz.Rect(M, top, M + 3, y), color=None, fill=colour)
        y += 7
    return len(pages)


# Label page geometry (A4 landscape), shared by the PDF and the page's live cell preview (cell_preview)
MM2PT = 2.83465
MARGIN = 4 * MM2PT
GAP = 2.5 * MM2PT
CELL_RADIUS = 3 * MM2PT
A4_W, A4_H = fitz.paper_size("a4-l")

ZONE_A_W = 107 * MM2PT
ZONE_A_H = 150 * MM2PT
CHIP_H = 11 * LINE + 8
# The row under the address on a label's first page: LABEL / PAGE (/ BOX) chips on the left and, on the right, the slot
# Stitch Labels puts the courier service's logo in, as tall as the chips (the logos are strips about 6.3 : 1, so the
# chips there are 9.5 pt to leave the strip its width). Its place goes in the label map.
CHIP1_SIZE = 9.5
LOGO_ROW_H = CHIP1_SIZE * LINE + 8   # the chips' height
LOGO_RATIO = 6.3
LOGO_W = LOGO_ROW_H * LOGO_RATIO     # about 43 x 6.8 mm
ROWS = 3

# One box size for every page: width from page 1's 4 columns beside the courier region,
# height from the later pages, which lose the footer strip at the bottom
FIRST_COLS = 4
CELL_W = (A4_W - 2 * MARGIN - ZONE_A_W - GAP - (FIRST_COLS - 1) * GAP) / FIRST_COLS
CELL_H = (A4_H - 2 * MARGIN - CHIP_H - GAP - (ROWS - 1) * GAP) / ROWS
NEXT_COLS = int((A4_W - 2 * MARGIN + GAP) // (CELL_W + GAP))
FIRST_PAGE_CELLS, NEXT_PAGE_CELLS = FIRST_COLS * ROWS, NEXT_COLS * ROWS


def _thumb_data_url(data, max_px):
    """An item image shrunk to fit max_px, as a JPEG data URL for the browser; None when it can't be read
    (the PDF shows "Image Error" for those too). See-through parts go white, as they print on the label."""
    from PIL import Image
    try:
        img = Image.open(io.BytesIO(data.getvalue()))
        img.load()
        img.thumbnail((max_px, max_px))
        img = img.convert('RGBA')
        flat = Image.new('RGB', img.size, (255, 255, 255))
        flat.paste(img, mask=img.getchannel('A'))
        buf = io.BytesIO()
        flat.save(buf, 'JPEG', quality=85)
        return f"data:image/jpeg;base64,{base64.b64encode(buf.getvalue()).decode('ascii')}"
    except Exception:
        return None


def cell_preview(pack_groups, max_px=360):
    """The first pack group's item cells for the page's live layout preview (script.js section 10). The browser
    draws them with the same sizes and rules as _layout_cell/_draw_cell, so this hands it the cell size in pt and,
    per item, its texts, barcode modules and images (small data URLs). Barcodes are given to every pack of the tab
    first (assign_barcodes), as the PDF does, so the K numbers match. Only page 1's worth of cells is sent.
    A horizontal file's label opens with its code boxes (group['code_cells']), so the preview does too."""
    if not pack_groups:
        return None
    assign_barcodes(pack_groups)
    group = next(iter(pack_groups.values()))
    cells = list(group.get('code_cells') or []) + group['items']
    items = []
    for item in cells[:FIRST_PAGE_CELLS]:
        # A code box has nothing to scan: no barcode, as _layout_cell draws it
        barcode = '' if item.get('is_code') else item.get('barcode') or _clean(item.get('job_no'))
        items.append({
            **{k: _clean(item.get(k)) for k in ('job_no', 'qty', 'desc', 'dimension', 'material', 'install', 'notes')},
            **{k: _clean(v) for k, v in item.items() if k.startswith('row_')},  # rows of the file offered as blocks
            'barcode': barcode,
            'is_code': bool(item.get('is_code')),  # a code box: left out when the labels aren't grouped by allocation
            'modules': code128_modules(barcode, compact=True) if barcode else None,
            'thumbnails': [_thumb_data_url(d, max_px) for d in item.get('thumbnails') or []],
        })
    return {
        'cell_pt': [round(CELL_W, 2), round(CELL_H, 2)],
        'radius_pt': round(CELL_RADIUS, 2),
        'barcode_quiet': BARCODE_QUIET_MODULES,
        'barcode_max_module': BARCODE_MAX_MODULE,
        'store': group.get('store_name', ''),
        'spec': group.get('pack_spec_name', ''),
        'items': items,
        'total_items': len(cells),
    }


def generate_packing_labels(pack_groups, output_pdf_path="packing_labels.pdf", attribute_order=None, addresses=None,
                            report=None, serials=None, tabs=None, tab_notes=None, on_progress=None):
    """Draws the PDF. Returns where each pack landed: {'pages': {pack_key: [0-based page indexes]}, ...}.

    addresses: {pack_key: {'address', 'receiver', 'contact', 'sent'}} as finalised in the consignment preview
    (address book, edits); a plain address string also works. A pack going to an installer gets a red
    "Installer: <receiver> · Attn <contact>" line; any other pack whose
    receiver isn't the store (or has an Attn) gets "Deliver to: …". A pack left out of the courier CSV
    (sent=False) says so in the courier label region; one collected instead ('collect': 'KYLE TO COLLECT',
    courier_export.collect_note) gets that note and 'PACK 2 OF 3' there in red on yellow, counted over its store's
    boxes only. A pack missing from addresses shows its Excel address.
    report: the job report (draw_report_pages) put in front as page 1; the page numbers returned count it.
    A Packing Spec with a box count in front ('2 x OB170170170', packing_specs.split_count) gets one full set of
    pages per box, each with a diagonal "BOX 1 OF 2" watermark over the item grid (region C) and a box chip, so the
    packers make that many boxes and each box gets its own courier label. 'copies' in the result lists each box's
    pages: {pack_key: [[pages of box 1], [pages of box 2], ...]}.
    serials: {pack_key: '01T1'}, the label's serial in its tab, written before the store name ('01T1. Blue Star
    Moorabbin') on every page of the pack. tabs: {pack_key: (tab number, tab name)}; a divider page goes between
    the labels of one tab and the next (none before the first tab). tab_notes: {tab number: [issues]} (vertical files):
    a PACK AUDIT page goes before each tab's labels instead (tab_audit), the tab's problems on it.
    on_progress(done, total, step): called after each pack's labels and before saving, for the loading overlay's bar."""
    from engine.courier.packing_specs import split_count
    addresses = addresses or {}
    doc = fitz.open()

    if not attribute_order:
        attribute_order = ['thumbnail', 'desc', 'dimension', 'job_no', 'barcode', 'qty']

    assign_label_numbers(pack_groups)
    assign_barcodes(pack_groups)
    # 'PACK 2 OF 3' on a collected store's labels: its boxes counted, that store's only (a '2 x' pack is two)
    store_boxes = Counter()
    for g in pack_groups.values():
        store_boxes[_norm(g['store_name'])] += split_count(g.get('pack_spec_name', ''))[0]
    boxes_seen = Counter()
    # One item-box layout for the whole job (every box is CELL_W x CELL_H): the same block at the same place in each
    plan = plan_cells([it for g in pack_groups.values() for it in list(g.get('code_cells') or []) + g['items']],
                      attribute_order, fitz.Rect(0, 0, CELL_W, CELL_H))
    page_map, copies = {}, {}
    serials, tabs = serials or {}, tabs or {}

    def divider(tab_no, tab_name):
        """A page between two tabs' labels: the tab, its name and the serials that follow. Not a label."""
        keys = [k for k in pack_groups if tabs.get(k, (None,))[0] == tab_no]
        first, last = serials.get(keys[0], ''), serials.get(keys[-1], '')
        page = doc.new_page(width=A4_W, height=A4_H)
        frame = fitz.Rect(MARGIN * 3, MARGIN * 3, A4_W - MARGIN * 3, A4_H - MARGIN * 3)
        page.draw_rect(frame, color=(0.2, 0.2, 0.2), width=1.5, dashes="[8 6] 0")
        page.insert_text((frame.x0 + 16, frame.y0 + 28), "DIVIDER - NOT A LABEL", fontname="hebo", fontsize=12, color=(0.45, 0.45, 0.45))
        lines = [(f"TAB {tab_no}", 120, (0, 0, 0)), (_latin(tab_name or ''), 40, (0.15, 0.15, 0.15)),
                 (f"Labels {first} - {last}  ·  {len(keys)} pack{'s' if len(keys) != 1 else ''}" if first else '', 20, (0.3, 0.3, 0.3))]
        y = frame.y0 + frame.height * 0.42
        for text, size, colour in lines:
            if not text:
                continue
            width = fitz.get_text_length(text, fontname="hebo", fontsize=size)
            page.insert_text(((A4_W - width) / 2, y), text, fontname="hebo", fontsize=size, color=colour)
            y += size * 1.25
        return page

    previous_tab = None

    def tab_audit(tab_no, tab_name):
        """Vertical files: the tab's PACK AUDIT page before its labels ('PACK AUDIT' over 'Tab 1: VIC'), the figures
        to check against the packing sheet, its Packing Specs and what needs a look, row and store."""
        keys = [k for k in pack_groups if tabs.get(k, (None,))[0] == tab_no]
        boxes = sum(split_count(pack_groups[k].get('pack_spec_name', ''))[0] for k in keys)
        specs = Counter(pack_groups[k].get('pack_spec_name', '') for k in keys)
        first, last = serials.get(keys[0], ''), serials.get(keys[-1], '')
        draw_audit_page(doc, f"Tab {tab_no}: {tab_name}" if tab_name else f"Tab {tab_no}",
                        [("Labels", len(keys)), ("Boxes", boxes), ("Stores", len({_norm(pack_groups[k]['store_name']) for k in keys})),
                         ("Items", sum(len(pack_groups[k]['items']) for k in keys))],
                        [f"Order: as in the file, store by store{f' (labels {first} - {last})' if first else ''}."],
                        tab_notes.get(tab_no, []), "Every pack in this tab has its label.", specs=specs)

    def watermark(page, area, box, boxes, spec):
        """'BOX 1 OF 2' written diagonally across the item grid, see-through, so the cells stay readable."""
        angle = math.degrees(math.atan2(area.height, area.width))
        centre = fitz.Point((area.x0 + area.x1) / 2, (area.y0 + area.y1) / 2)
        lines = ((f"BOX {box} OF {boxes}", min(96, area.width / 6.2)),
                 (f"{boxes} BOXES: {spec} - ONE LABEL PER BOX", min(22, area.width / 30)))
        for n, (text, size) in enumerate(lines):
            width = fitz.get_text_length(text, fontname="hebo", fontsize=size)
            # Both lines centred on the grid's centre along the diagonal, the second under the first
            start = fitz.Point(centre.x - width / 2, centre.y + lines[0][1] * 0.35 + (0 if n == 0 else size * 1.9))
            page.insert_text(start, text, fontname="hebo", fontsize=size, color=(0.78, 0.08, 0.08),
                             fill_opacity=0.22, morph=(centre, fitz.Matrix(-angle)), overlay=True)

    for done, (pack_key, group_data) in enumerate(pack_groups.items()):
        if on_progress:
            on_progress(done, len(pack_groups), f"Drawing the packing labels: {done + 1} of {len(pack_groups)}")
        tab_no, tab_name = tabs.get(pack_key, (None, ''))
        if tab_notes is not None and tab_no is not None and tab_no != previous_tab:
            tab_audit(tab_no, tab_name)
        elif tab_no is not None and previous_tab is not None and tab_no != previous_tab:
            divider(tab_no, tab_name)
        previous_tab = tab_no if tab_no is not None else previous_tab
        # Horizontal files: the packing-sheet code boxes go first. Drawn only, so they never count as items,
        # units, job numbers or barcode kinds anywhere else
        items = list(group_data.get('code_cells') or []) + group_data['items']
        total_items = len(items)
        store = group_data['store_name']
        # '01T1. Blue Star Moorabbin': the label's serial in its tab before the name, on every page of the pack
        shown_name = f"{serials[pack_key]}. {store}" if serials.get(pack_key) else store
        spec = group_data.get('pack_spec_name', '')
        label_text = f"LABEL {group_data['label_no']} OF {group_data['label_total']}"
        info = addresses.get(pack_key) or {}
        if isinstance(info, str):
            info = {'address': info}
        installer = bool(group_data.get('install'))
        first_page = len(doc)

        total_pages = 1 if total_items <= FIRST_PAGE_CELLS else 1 + math.ceil((total_items - FIRST_PAGE_CELLS) / NEXT_PAGE_CELLS)
        boxes, one_box_spec = split_count(spec)
        copies[pack_key] = []
        for box in range(1, boxes + 1):
            box_first_page = len(doc)
            boxes_seen[_norm(store)] += 1
            current_item_idx = 0
            page_num = 1

            while current_item_idx < total_items or (total_items == 0 and page_num == 1):
                page = doc.new_page(width=A4_W, height=A4_H)
                page_text = f"PAGE {page_num} OF {total_pages}"

                if page_num == 1:
                    zone_a_rect = fitz.Rect(MARGIN, MARGIN, MARGIN + ZONE_A_W, MARGIN + ZONE_A_H)
                    page.draw_rect(zone_a_rect, color=(0.8, 0.8, 0.8), width=0.5, dashes="[3] 0")
                    if info.get('collect'):
                        # Collected, not sent ('Kyle collect from warehouse'): who collects, the store the packs are
                        # for and which of its packs this is, in red on yellow where the courier label would go
                        panel = fitz.Rect(zone_a_rect.x0 + 8, zone_a_rect.y0 + 8, zone_a_rect.x1 - 8, zone_a_rect.y1 - 8)
                        page.draw_rect(panel, color=(0.8, 0.0, 0.0), fill=(1.0, 0.92, 0.25), width=2)
                        red, room = (0.8, 0.0, 0.0), panel.width - 20

                        def fitted(text, size, least, most_lines):
                            # Biggest size (down to least) at which no word runs past the panel
                            while size > least and max(fitz.get_text_length(w, fontname="hebo", fontsize=size)
                                                       for w in text.split()) > room:
                                size -= 1
                            return [(line, size) for line in _wrap(text, "hebo", size, room)[:most_lines]]

                        blocks = [fitted(_latin(info['collect']), 26, 12, 3),
                                  fitted(_latin(" ".join(str(store).split())), 18, 10, 3),
                                  fitted(f"PACK {boxes_seen[_norm(store)]} OF {store_boxes[_norm(store)]}", 22, 12, 1)]
                        gap = 14
                        height = sum(sz * 1.2 for b in blocks for _, sz in b) + gap * (len(blocks) - 1)
                        y = panel.y0 + (panel.height - height) / 2
                        for b in blocks:
                            for line, sz in b:
                                y += sz * 1.2
                                page.insert_text((panel.x0 + (panel.width - fitz.get_text_length(line, fontname="hebo", fontsize=sz)) / 2,
                                                  y - sz * 0.25), line, fontname="hebo", fontsize=sz, color=red)
                            y += gap
                    elif info.get('sent', True):
                        page.insert_textbox(zone_a_rect, "Courier Label Region\n(107mm x 150mm)", fontsize=10, color=(0.6, 0.6, 0.6), align=1)
                    else:
                        # Left out of the courier CSV: no courier label will come for this box
                        note = fitz.Rect(zone_a_rect.x0 + 10, zone_a_rect.y0 + zone_a_rect.height / 2 - 30, zone_a_rect.x1 - 10, zone_a_rect.y1)
                        page.insert_textbox(note, "NOT IN COURIER CSV\nAddress incomplete: no courier label for this box",
                                            fontname="hebo", fontsize=11, color=(0.75, 0.1, 0.1), align=1)

                    zone_b = fitz.Rect(MARGIN, zone_a_rect.y1 + GAP, MARGIN + ZONE_A_W, A4_H - MARGIN)
                    y = _spec_panel(page, zone_b, spec, tag=_latin(group_data.get('pack_tag') or '')) + 6

                    if store and store.lower() != 'none':
                        for line in _wrap(shown_name, "hebo", 13, zone_b.width)[:2]:
                            page.insert_text((zone_b.x0, y + 12), line, fontname="hebo", fontsize=13)
                            y += 13 * LINE

                    # Who it's really going to, when that isn't the store itself (installers, Attn names)
                    receiver, contact = _clean(info.get('receiver')), _clean(info.get('contact'))
                    if installer or (receiver and _norm(receiver) != _norm(store)) or contact:
                        to = ("Installer: " if installer else "Deliver to: ") + (receiver or store or '')
                        to += f" · Attn {contact}" if contact else ''
                        # One line when it fits at 8 pt or more, so the address keeps its room above the logo row
                        to_size = 9.5
                        while to_size > 8 and len(_wrap(to, "hebo", to_size, zone_b.width)) > 1:
                            to_size -= 0.5
                        for line in _wrap(to, "hebo", to_size, zone_b.width)[:2]:
                            page.insert_text((zone_b.x0, y + to_size), line, fontname="hebo", fontsize=to_size,
                                             color=(0.75, 0.1, 0.1) if installer else (0, 0, 0))
                            y += to_size * LINE

                    address = info.get('address')
                    if not address:
                        addr_parts = [_clean(group_data.get(k)) for k in ('address_1', 'address_2', 'suburb', 'state', 'postcode', 'country')]
                        address = ", ".join(p for p in addr_parts if p)
                    if address:
                        address_lines = _wrap(address, "helv", 9, zone_b.width)
                        room = int((zone_b.y1 - LOGO_ROW_H - 3 - y) // (9 * LINE))
                        for line in address_lines[:max(0, min(3, room))]:
                            page.insert_text((zone_b.x0, y + 9), line, fontname="helv", fontsize=9)
                            y += 9 * LINE

                    # The chips, centred in the logo row and kept clear of the logo slot (smaller type if they must)
                    chips = [label_text, page_text] + ([f"BOX {box} OF {boxes}"] if boxes > 1 else [])
                    room_w = zone_b.width - LOGO_W - 8
                    chip_size = CHIP1_SIZE
                    while chip_size > 7 and sum(_text_len(t, "hebo", chip_size) + 16 for t in chips) + 6 * (len(chips) - 1) > room_w:
                        chip_size -= 0.5
                    chip_y = zone_b.y1 - LOGO_ROW_H + (LOGO_ROW_H - (chip_size * LINE + 8)) / 2
                    x = zone_b.x0
                    for t in chips:
                        x = _counter_chip(page, x, chip_y, t, size=chip_size).x1 + 6

                    grid_x0 = zone_a_rect.x1 + GAP
                    cols = FIRST_COLS
                else:
                    footer_y = A4_H - MARGIN - CHIP_H
                    x = MARGIN
                    if store and store.lower() != 'none':
                        page.insert_text((x, footer_y + 4 + 12 * 0.92), shown_name, fontname="hebo", fontsize=12)
                        x += fitz.get_text_length(shown_name, fontname="hebo", fontsize=12) + 10
                    chip = _counter_chip(page, x, footer_y, label_text)
                    chip = _counter_chip(page, chip.x1 + 6, footer_y, page_text)
                    if boxes > 1:
                        _counter_chip(page, chip.x1 + 6, footer_y, f"BOX {box} OF {boxes}")

                    grid_x0 = MARGIN
                    cols = NEXT_COLS


                for r in range(ROWS):
                    for c in range(cols):
                        if current_item_idx >= total_items: break
                        cx0 = grid_x0 + c * (CELL_W + GAP)
                        cy0 = MARGIN + r * (CELL_H + GAP)
                        cell = fitz.Rect(cx0, cy0, cx0 + CELL_W, cy0 + CELL_H)
                        _round_rect(page, cell, CELL_RADIUS, color=(0, 0, 0), width=0.75)
                        _draw_cell(page, items[current_item_idx], attribute_order, cell, plan)
                        current_item_idx += 1
                if boxes > 1:  # region C: the item grid, right of the courier region on page 1, above the footer after
                    grid = fitz.Rect(grid_x0, MARGIN, A4_W - MARGIN, A4_H - MARGIN - (0 if page_num == 1 else CHIP_H + GAP))
                    watermark(page, grid, box, boxes, one_box_spec)
                page_num += 1
                if total_items == 0: break
            copies[pack_key].append(list(range(box_first_page, len(doc))))
        page_map[pack_key] = list(range(first_page, len(doc)))

    report_pages = draw_report_pages(doc, report, at=0) if report else 0
    if report_pages:
        page_map = {k: [p + report_pages for p in v] for k, v in page_map.items()}
        copies = {k: [[p + report_pages for p in box] for box in v] for k, v in copies.items()}
    if on_progress:
        on_progress(len(pack_groups), len(pack_groups), f"Saving the PDF ({len(doc)} pages)")
    doc.save(output_pdf_path, garbage=1, deflate=True)  # artwork is shared already: no duplicate search needed
    doc.close()
    return {
        'pages': page_map,
        'copies': copies,
        'report_pages': report_pages,
        'page_size_mm': [round(A4_W / MM2PT, 1), round(A4_H / MM2PT, 1)],
        'courier_region_mm': {'x': round(MARGIN / MM2PT, 1), 'y': round(MARGIN / MM2PT, 1),
                              'width': round(ZONE_A_W / MM2PT, 1), 'height': round(ZONE_A_H / MM2PT, 1)},
        # Where Stitch Labels puts the courier service's logo on each label's first page (LOGO_ROW_H, LOGO_W)
        'logo_region_mm': {'x': round((MARGIN + ZONE_A_W - LOGO_W) / MM2PT, 2), 'y': round((A4_H - MARGIN - LOGO_ROW_H) / MM2PT, 2),
                           'width': round(LOGO_W / MM2PT, 2), 'height': round(LOGO_ROW_H / MM2PT, 2)},
    }
