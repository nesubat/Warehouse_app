"""Horizontal distribution labels: reading, checks, packing-sheet order, references, the PDF, stitching codes,
and the whole Packing Labels page (upload -> preview -> generate) through Flask's test client.

    pip install -r requirements-dev.txt
    pytest -q
"""
import copy
import json
import os
import random
import re
import sys

import pymupdf as fitz
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fixtures import build_final_file, build_signature_links  # noqa: E402
from engine.packing_labels.horizontal_distribution import read_horizontal, read_signature_file, item_references  # noqa: E402
from engine.packing_labels.packing_label_generator import PackCheckError, assign_barcodes  # noqa: E402

FINAL = "Final J476699 Spring rollout.xlsx"
SIG = "Signature links_J476699.xlsx"


def serial_of(code):
    """'03T1P2QWXK' -> '03T1P2': the serial printed on the packing label, without this Generate's code."""
    return re.sub(r"[A-Z]{4,}$", "", code)


@pytest.fixture()
def files(tmp_path):
    codes = build_final_file(tmp_path / FINAL)
    build_signature_links(tmp_path / SIG, codes)
    return tmp_path / FINAL, tmp_path / SIG, codes


def read(path, tab="Spring Sale", sig=None, tab_no=1):
    return read_horizontal(str(path), 8, tab, None, tab_no, read_signature_file(str(sig)) if sig else None)


def test_codes_read_from_code_columns_in_packing_sheet_order(files):
    final, sig, _ = files
    groups, _, found, warnings = read(final, sig=sig)
    order = [(g["pack_no"], g["pack_code"], g["store_name"]) for g in groups.values()]
    assert order[:3] == [(1, "A", "Westfield Parramatta"), (1, "A", "Carindale"), (1, "B", "Corio Village")]
    assert [p for p, _, _ in order] == sorted(p for p, _, _ in order)  # pack by pack
    assert found["plugin"] and not warnings


def test_reference_is_store_number_tab_pack(files):
    final, sig, _ = files
    groups, *_ = read(final, sig=sig)
    refs = dict(item_references(groups))
    assert "03T1P2 Chadstone" in refs and refs["03T1P2 Chadstone"] == "03T1P2"
    easter, *_ = read(final, "Easter", sig, tab_no=2)
    assert {g["serial"] for g in easter.values()} == {"01T2P1", "02T2P1", "03T2P1", "04T2P1"}


def test_subgroup_codes_kept_and_shown(files):
    final, sig, _ = files
    groups, *_ = read(final, sig=sig)
    chad = next(g for g in groups.values() if g["serial"] == "03T1P2")
    assert chad["sub_codes"] == [("1-2", "C")]
    # The codes are in the label's code boxes, Signature Links or not (the packing spec box has no pack or code tag)
    assert [c["qty"] for c in chad["code_cells"]] == [chad["pack_code"], "C"]
    alone, *_ = read(final)
    chad = next(g for g in alone.values() if g["serial"] == "03T1P2")
    assert [c["qty"] for c in chad["code_cells"]] == [chad["pack_code"], "C"]


def test_without_code_columns_codes_come_from_signature_links_then_quantities(tmp_path):
    codes = build_final_file(tmp_path / FINAL, codes=False)
    build_signature_links(tmp_path / SIG, codes)
    with_sig, *_ = read(tmp_path / FINAL, sig=tmp_path / SIG)
    alone, _, found, _ = read(tmp_path / FINAL)
    assert [g["pack_code"] for g in with_sig.values()] == [g["pack_code"] for g in alone.values()]
    assert "worked out" in found["codes"][0]


def test_signature_links_decide_the_groups_and_differences_only_warn(tmp_path):
    codes = build_final_file(tmp_path / FINAL)
    build_signature_links(tmp_path / SIG, codes, tamper=("Spring Sale", "1", 0, "C"))
    groups, _, _, warnings = read(tmp_path / FINAL, sig=tmp_path / SIG)  # no PackCheckError: Generate goes ahead
    corio = next(g for g in groups.values() if g["store_name"] == "Corio Village" and g["pack_no"] == 1)
    assert corio["pack_code"] == "C"  # Signature Links' code is the one used
    texts = [w["text"] for w in warnings if w.get("major")]
    assert "Code differs from the Signature Links file" in texts
    assert "Signature Links code doesn't match the quantities" in texts
    assert all("Pack 1" in w["col"] for w in warnings if "Signature Links" in w["text"])  # on pack 1's audit page


def test_quantities_changed_after_packing_sheets_stop_generate(tmp_path):
    build_final_file(tmp_path / FINAL, change=("Spring Sale", 0, 0, 7))
    with pytest.raises(PackCheckError) as e:
        read(tmp_path / FINAL)
    assert any("doesn't match the quantities" in i["text"] for i in e.value.issues)


def test_every_unit_lands_on_a_label(files):
    final, sig, _ = files
    groups, *_ = read(final, sig=sig)
    from openpyxl import load_workbook
    ws = load_workbook(final, data_only=True)["Spring Sale"]
    jobs = [c for c in range(1, ws.max_column + 1) if re.match(r"J\d", str(ws.cell(row=7, column=c).value or ""))]
    grid = sum(ws.cell(row=r, column=c).value or 0 for r in range(9, 17) for c in jobs)
    assert grid == sum(float(i["qty"]) for g in groups.values() for i in g["items"])


def test_barcode_kinds_follow_item_columns(files):
    final, sig, _ = files
    groups, *_ = read(final, sig=sig)
    assign_barcodes(groups)
    kinds = {i["barcode"] for g in groups.values() for i in g["items"] if i["job_no"] == "J476699-03"}
    assert kinds == {"J476699-03 K1", "J476699-03 K2"}  # two columns share the job: two kinds, on every store


def test_barcode_kinds_of_vertical_files_unchanged():
    """No source_key (vertical files): every row is its own kind, as before."""
    rnd = random.Random(4)
    for _ in range(200):
        groups = {f"g{i}": {"items": [{"job_no": rnd.choice(["J1-01", "J1-02", "j1-02 ", ""])} for _ in range(rnd.randint(0, 5))]}
                  for i in range(rnd.randint(1, 6))}
        expected = copy.deepcopy(groups)
        seen, totals = {}, {}
        for g in expected.values():
            for it in g["items"]:
                k = it["job_no"].strip().upper()
                if k:
                    totals[k] = totals.get(k, 0) + 1
        for g in expected.values():
            for it in g["items"]:
                k, job = it["job_no"].strip().upper(), it["job_no"].strip()
                if not k:
                    it["barcode"] = ""
                    continue
                seen[k] = seen.get(k, 0) + 1
                it["barcode"] = f"{job} K{seen[k]}" if totals[k] > 1 else job
        assign_barcodes(groups)
        assert groups == expected


def test_stitcher_reads_horizontal_codes():
    from engine.packing_labels.label_stitcher import find_code
    known = {"03T1P2", "12T2P1"}
    assert find_code("Item Ref: 03T1P2 Chadstone", known) == "03T1P2"
    assert find_code("Item Reference 12T2P1 Rundle Mall", known) == "12T2P1"
    assert find_code("Item Ref: 103T1P2 Chadstone", known) is None
    assert find_code("Item Ref: 03T1P21 Chadstone", known) is None


def test_signature_links_tabs_skipped_by_shuffler():
    from engine.core.core_math import NOT_STORE_TABS, LABEL_PLUGIN_SHEET, DIVIDER_BARCODE_SHEET
    assert set(NOT_STORE_TABS) == {LABEL_PLUGIN_SHEET, DIVIDER_BARCODE_SHEET}


# ------------------------------------------------------------------ the page, end to end

@pytest.fixture()
def client(tmp_path, monkeypatch):
    import app as app_module
    from engine.addresses.address_book import AddressBook
    from engine.courier.packing_specs import SpecStore
    monkeypatch.setitem(app_module.app.config, "UPLOAD_FOLDER", str(tmp_path / "projects"))
    monkeypatch.setattr(app_module, "PROJECTS_FOLDER", str(tmp_path / "projects"))  # the Label Shuffler reads it
    os.makedirs(tmp_path / "projects", exist_ok=True)
    # Every store the page reads or writes on temp files of the test's own: Generate saves addresses, service codes
    # and kept pages, and the preview adds unknown packing specs; none of it may reach the real data/ folder
    data = tmp_path / "data"
    os.makedirs(data, exist_ok=True)
    monkeypatch.setattr(app_module, "ADDRESS_BOOK", AddressBook(str(data / "address_book.db")))
    monkeypatch.setattr(app_module, "SPEC_STORE", SpecStore(str(data / "packing_specs.json")))
    monkeypatch.setattr(app_module, "SERVICE_USAGE_FILE", str(data / "service_code_usage.json"))
    from engine.courier.courier_services import ServiceStore
    monkeypatch.setattr(app_module, "SERVICE_STORE", ServiceStore(str(data / "courier_services.json"), str(data / "courier logos")))
    monkeypatch.setattr(app_module, "EXPORT_FILE", str(data / "export_details.json"))
    monkeypatch.setattr(app_module, "VIEW_FOLDER", str(data / "page_views"))
    for name in ("close_if_open_elsewhere", "save_if_open_elsewhere"):  # Excel-only helpers
        monkeypatch.setattr(app_module, name, lambda *a, **k: None)
    monkeypatch.setattr(app_module, "convert_legacy_excel_to_xlsx", lambda p: p)
    # Packing Sheets runs in Excel: here it writes its two files and keeps what it was asked for (app_module.sheets_runs)
    app_module.sheets_runs = []

    def packing_sheets(file_path, original, tabs, inputs, tab_data, folder):
        app_module.sheets_runs.append({'tabs': tabs, 'inputs': inputs, 'tab_data': tab_data})
        for name in ("Packing Sheet_J476699.xlsx", "Signature links_J476699.xlsx"):
            with open(os.path.join(folder, name), "wb") as f:
                f.write(b"xlsx")
        return "Final x.xlsx", "Packing Sheet_J476699.xlsx", "Signature links_J476699.xlsx", []
    monkeypatch.setattr(app_module, "generate_all_outputs", packing_sheets)
    app_module.app.config["TESTING"] = True
    return app_module.app.test_client(), app_module, tmp_path / "projects"


def run_page(client, final, sig, dividers=True, extra=None):
    c, app_module, projects = client
    with open(final, "rb") as f, open(sig, "rb") as s:
        r = c.post("/packing-labels/horizontal", data={"file": (f, final.name), "signature_file": (s, sig.name)},
                   content_type="multipart/form-data", follow_redirects=True)
    assert r.status_code == 200, r.data[:500]
    page = r.data.decode()
    assert 'value="8"' in page  # header rows filled in from the Label Plugin
    filename = app_module.secure_filename(final.name)
    form = {"filename": filename, "all_tabs": ["Spring Sale", "Easter"], "selected_tabs": ["Spring Sale", "Easter"],
            "header_row_Spring Sale": "8", "header_row_Easter": "8"}
    r = c.post("/packing-labels/horizontal", data={**form, "preview": "1"}, follow_redirects=True)
    preview = r.data.decode()
    assert "Cannot Generate" not in preview, preview[preview.find("issue"):][:800]
    assert "Signature Links loaded" in preview
    service = app_module.SERVICE_STORE.services()[0]['code']
    gen = {**form, "generate": "1", "who_pays": "S", "service_code": service, "reference": "J476699",
           "project_name": "Spring rollout", "courier_csv": ["open360", "openfreight"],
           "attribute_order": "thumbnail,desc,dimension,job_no,barcode,qty"}
    if dividers:
        gen["code_dividers"] = "1"
    gen.update(extra or {})
    r = c.post("/packing-labels/horizontal", data=gen, follow_redirects=True)
    assert r.status_code == 200 and "Labels Successfully Generated" in r.data.decode(), r.data.decode()[:1500]
    folder = next(p for p in projects.iterdir() if p.is_dir())
    return folder


def test_page_generates_labels_csv_and_label_map(client, files):
    final, sig, _ = files
    folder = run_page(client, final, sig)
    names = os.listdir(folder)
    pdf = next(n for n in names if n.startswith("Packing Labels Only"))
    open360 = next(n for n in names if n.endswith("Open360.csv"))
    labelmap = json.load(open(folder / next(n for n in names if n.endswith(".labelmap.json"))))
    assert any(n.startswith("Used Signature Links") for n in names)
    rows = open(folder / open360, encoding="utf-8").read()
    assert re.search(r"03T1P2[A-Z]{4} Chadstone", rows) and re.search(r"01T2P1[A-Z]{4} Corio Village", rows)
    doc = fitz.open(folder / pdf)
    texts = [p.get_text() for p in doc]
    # report, then per pack: audit page, divider per code, labels
    audit = texts[labelmap.get("report_pages", 1)]
    assert "PACK AUDIT" in audit and "Tab 1 - Pack 1" in audit  # the header: PACK AUDIT over Tab 1 - Pack 1
    codes = {pk["open360_code"] for packs in labelmap["stores"].values() for pk in packs}
    assert "03T1P2" in {serial_of(c) for c in codes} and all(c != serial_of(c) for c in codes)  # serial + a code
    for packs in labelmap["stores"].values():  # every label-map page really is that pack's packing label
        for pk in packs:
            for p in pk["pdf_pages"] + pk["box_pages"]:
                assert serial_of(pk["open360_code"]) + "." in texts[p - 1], (pk["open360_code"], texts[p - 1][:200])
    chad = next(pk for pk in labelmap["stores"]["Chadstone"] if serial_of(pk["open360_code"]) == "03T1P2")
    assert chad["boxes"] == 2 and len(set(chad["box_pages"])) == 2  # one courier label per box


def test_page_without_dividers_has_no_divider_sheets(client, files):
    final, sig, _ = files
    folder = run_page(client, final, sig, dividers=False)
    pdf = next(n for n in os.listdir(folder) if n.startswith("Packing Labels Only"))
    assert not any("Matched Labels" in p.get_text() for p in fitz.open(folder / pdf))


def test_cross_site_actions_refused(client):
    c, *_ = client
    assert c.post("/packing-labels/horizontal", headers={"Sec-Fetch-Site": "cross-site"}).status_code == 403
    assert c.get("/open_upload/x.xlsx", headers={"Sec-Fetch-Site": "cross-site"}).status_code == 403
    assert c.post("/packing-labels/horizontal", headers={"Origin": "http://evil.example"}).status_code == 403
    assert c.get("/packing-labels/horizontal", headers={"Sec-Fetch-Site": "cross-site"}).status_code == 200  # just a page


def test_courier_labels_stitch_onto_their_packing_labels(client, files):
    """Courier labels booked from the Open360 CSV (Item Ref '03T1P2 Chadstone'), in any order, land on their own
    packing label; a pack of two boxes takes two."""
    import csv
    from engine.packing_labels.label_stitcher import stitch
    final, sig, _ = files
    folder = run_page(client, final, sig)
    names = os.listdir(folder)
    labelmap = json.load(open(folder / next(n for n in names if n.endswith(".labelmap.json"))))
    rows = list(csv.DictReader(open(folder / next(n for n in names if n.endswith("Open360.csv")), encoding="utf-8")))
    refs = [r["Item Reference"] for r in rows for _ in range(int(r["Item Quantity"] or 1))]
    assert all(re.match(r"\d{2}T\dP\d[A-Z]{4} ", ref) for ref in refs)  # serial + this Generate's code
    random.Random(7).shuffle(refs)
    courier = fitz.open()
    for ref in refs:  # a 4x6 courier label with its Item Ref
        page = courier.new_page(width=288, height=432)
        page.insert_text((20, 40), "TIG Freight - Open360", fontsize=12)
        page.insert_text((20, 300), f"Item Ref: {ref}", fontsize=10)
    out = folder / "Complete.pdf"
    result = stitch(str(folder), labelmap, [("courier.pdf", courier.tobytes())], str(out))
    assert not result["unmatched"] and len(result["placed"]) == len(refs)
    assert fitz.open(out).page_count >= len(refs)


def test_packing_spec_box_has_no_pack_number_and_a_long_spec_still_fits(tmp_path):
    """The Packing Spec box (the foot row's, beside the chips) shows the spec only: no 'PACK 2' or code in it. A spec
    too long for one line goes on two, inside the box."""
    from engine.packing_labels.packing_label_generator import generate_packing_labels
    for n, spec in enumerate(("2 x CS450320150", "Bespoke timber crate with foam inserts")):
        group = {"g": {"pack_spec_name": spec, "store_name": "Chadstone", "receiver_name": "", "install": False,
                       "address_1": "", "address_2": "", "suburb": "", "state": "", "postcode": "", "country": "",
                       "pack_no": 2, "pack_code": "C", "excel_rows": [9],
                       "items": [{"job_no": "J1-01", "qty": "1", "desc": "", "dimension": "", "material": "", "install": "",
                                  "notes": "", "thumbnails": []}]}}
        out = tmp_path / f"l{n}.pdf"
        generate_packing_labels(group, str(out))
        page = fitz.open(out)[0]
        words = [w for w in page.get_text("words") if w[1] > 500]  # the foot of the label block
        assert not any(w[4] == "PACK" for w in words)
        box = next(d['rect'] for d in page.get_drawings() if d.get('fill') == (0, 0, 0) and d['rect'].y0 > 500)
        inside = [w[4] for w in words if box.contains(fitz.Rect(w[:4]))]
        assert sorted(inside) == sorted(["PACKING", "SPEC"] + spec.split())  # the whole spec, in the box


def test_labels_carry_code_boxes_and_audit_page_is_short(client, files):
    """Each label has an item box per code: the code column's header in the job number's place, the store's code
    in the quantity's, no barcode. Code boxes never count as job numbers. The audit page has no code table."""
    final, sig, _ = files
    folder = run_page(client, final, sig)
    names = os.listdir(folder)
    doc = fitz.open(folder / next(n for n in names if n.startswith("Packing Labels Only")))
    labelmap = json.load(open(folder / next(n for n in names if n.endswith(".labelmap.json"))))
    chad = next(pk for pk in labelmap["stores"]["Chadstone"] if serial_of(pk["open360_code"]) == "03T1P2")
    page = " ".join(doc[chad["pdf_pages"][0] - 1].get_text().split())  # a long header wraps in its box
    assert page.count("Code for Pack 2") == 1 and "Code for pack 2: 1-2" in page  # header once: no barcode caption
    words = [w[4] for w in doc[chad["pdf_pages"][0] - 1].get_text("words")]
    assert words.count("C") >= 2  # the pack code and the sub-group code, each in its quantity box
    assert all(not str(j).startswith("Code for") for packs in labelmap["stores"].values() for pk in packs
               for j in pk.get("job_numbers", []))
    audit = next(p.get_text() for p in doc if "PACK AUDIT" in p.get_text() and "Tab 1 - Pack 2" in p.get_text())
    assert "J476699-05" not in audit and "Spring Sale" not in audit and "LABELS" in audit


def test_code_box_named_after_pack_when_codes_come_from_signature_links(tmp_path):
    codes = build_final_file(tmp_path / FINAL, codes=False)
    build_signature_links(tmp_path / SIG, codes)
    groups, *_ = read(tmp_path / FINAL, sig=tmp_path / SIG)
    g = next(iter(groups.values()))
    assert [c["job_no"] for c in g["code_cells"]] == ["Code for Pack 1"] and g["code_cells"][0]["qty"] == g["pack_code"]
    assert all(not i.get("is_code") for i in g["items"])


def test_found_names_the_cell_layout_blocks(files):
    """The page's cell layout offers a block only when read_horizontal() says its details were found, by the same
    names parse_packing_data() uses. Without them every slot showed "Empty Field" for horizontal files."""
    final, sig, _ = files
    _, _, found, _ = read(final, sig=sig)
    for name in ("Job Number", "Quantity", "Description", "Dimensions", "Thumbnail", "Install",
                 "Packing Spec", "Store Name", "Address Info"):
        assert found[name], name
    assert not found["Material"] and not found["Notes"]  # no Material or Notes row in the file


def test_page_offers_layout_blocks_and_previews_code_box_first(client, files):
    final, sig, _ = files
    c, app_module, _ = client
    with open(final, "rb") as f, open(sig, "rb") as s:
        c.post("/packing-labels/horizontal", data={"file": (f, final.name), "signature_file": (s, sig.name)},
               content_type="multipart/form-data", follow_redirects=True)
    form = {"filename": app_module.secure_filename(final.name), "all_tabs": ["Spring Sale"],
            "selected_tabs": ["Spring Sale"], "header_row_Spring Sale": "8", "preview": "1"}
    page = c.post("/packing-labels/horizontal", data=form, follow_redirects=True).data.decode()
    cell = page[page.index('id="attribute-list"'):page.index('id="attribute_order"')]
    for block in ("thumbnail", "desc", "dimension", "job_no", "barcode", "qty"):
        assert f'data-id="{block}"' in cell, block
    assert 'data-id="empty"' not in cell  # no "Empty Field" slots
    assert 'data-id="install"' in page[page.index('id="mapped-headers-list"'):page.index('id="attribute-list"')]
    preview = json.loads(re.search(r"data-cells='(.*?)'", page, re.DOTALL).group(1))
    first, item = preview["items"][0], preview["items"][1]
    assert first["job_no"].startswith("Code for Pack") and not first["modules"]  # the code box: no barcode
    assert item["job_no"].startswith("J476699") and item["modules"]


# ------------------------------------------------------------------ problems ranked and grouped

def test_issues_grouped_with_major_first():
    """The same problem on many rows is one line listing the rows; the ones that lose labels or units come first."""
    from engine.packing_labels.packing_label_generator import _issue, group_issues
    issues = [_issue("4", "C", "No artwork image"),
              _issue("11", "Pack 1", "Packing Spec but no quantities: no label made", major=True),
              _issue("13", "Pack 1", "Packing Spec but no quantities: no label made", major=True),
              _issue("19", "Pack 2", "Packing Spec but no quantities: no label made", major=True)]
    groups = group_issues(issues)
    assert [g["text"] for g in groups] == ["Packing Spec but no quantities: no label made", "No artwork image"]
    assert groups[0]["major"] and groups[0]["count"] == 3
    assert groups[0]["where"] == ["Pack 1: rows 11, 13", "Pack 2: row 19"]
    assert not groups[1]["major"] and groups[1]["where"] == ["C: row 4"]


def _spec_without_quantities(tmp_path, monkeypatch):
    """The sample file with Carindale given pack 2's spec but no units in it, and Garden City pack 1's."""
    import fixtures
    tabs = copy.deepcopy(fixtures.TABS)
    stores = tabs["Spring Sale"]["stores"]
    for i, st in enumerate(stores):
        if st[0] == "Carindale":
            stores[i] = st[:7] + (["OB600100100", "CS450320150"], st[8])
        if st[0] == "Garden City":
            stores[i] = st[:7] + (["OB600100100", "CS450320150"], st[8])
    monkeypatch.setattr(fixtures, "TABS", tabs)
    codes = build_final_file(tmp_path / FINAL)
    build_signature_links(tmp_path / SIG, codes)
    return tmp_path / FINAL, tmp_path / SIG


def test_packing_spec_without_quantities_is_major(tmp_path, monkeypatch):
    final, sig = _spec_without_quantities(tmp_path, monkeypatch)
    _, _, _, warnings = read(final, sig=sig)
    found = [w for w in warnings if w["text"].startswith("Packing Spec but no quantities")]
    assert {(w["rows"], w["col"]) for w in found} == {("13", "Pack 2"), ("15", "Pack 1")}
    assert all(w["major"] for w in found)


def test_preview_ranks_file_problems_and_has_one_open_excel_button(client, tmp_path, monkeypatch):
    """The preview's To check list gets the file's problems, grouped, major first; the page has one floating Open
    Excel button and none fixed in the heading."""
    final, sig = _spec_without_quantities(tmp_path, monkeypatch)
    c, app_module, _ = client
    with open(final, "rb") as f, open(sig, "rb") as s:
        c.post("/packing-labels/horizontal", data={"file": (f, final.name), "signature_file": (s, sig.name)},
               content_type="multipart/form-data", follow_redirects=True)
    form = {"filename": app_module.secure_filename(final.name), "all_tabs": ["Spring Sale"],
            "selected_tabs": ["Spring Sale"], "header_row_Spring Sale": "8", "preview": "1"}
    page = c.post("/packing-labels/horizontal", data=form, follow_redirects=True).data.decode()
    import html as html_lib
    review = json.loads(html_lib.unescape(re.search(r'id="pl-issues" data-review="([^"]*)"', page).group(1)))
    first = review["file"][0]
    assert first["major"] and first["text"].startswith("Packing Spec but no quantities") and first["count"] == 2
    assert first["tab"] == "Spring Sale"
    assert "1 needs attention" in page  # the tab card's heading says so too, in red (the two rows are one problem)
    assert page.count('id="fx-open"') == 1 and "Open Excel File" not in page
    # Auto-Map stays under the tabs; Update Previews floats with Open Excel and sends the same form
    assert '>Auto-Map</button>' in page
    floating = page[page.index('id="fx-open"'):page.index('</div>', page.index('id="fx-open"'))]
    assert 'form="pl-form" name="preview"' in floating and "Update Previews" in floating


def test_every_row_above_the_stores_is_a_block_for_the_item_box(tmp_path):
    """A row written in the label column (or with values under the items) that isn't one of the known ones becomes a
    block of its own, a label written twice included; picked for the cell, its value prints on the label."""
    import openpyxl
    from fixtures import ADDRESS, JOB_ROW
    from engine.packing_labels.packing_label_generator import generate_packing_labels
    final = tmp_path / FINAL
    build_final_file(final)
    wb = openpyxl.load_workbook(final)
    ws = wb["Spring Sale"]
    items = [c for c in range(1, ws.max_column + 1) if re.match(r'J\d', str(ws.cell(row=JOB_ROW, column=c).value or ''))]
    ws.cell(row=5, column=len(ADDRESS), value="Kinds")
    ws.cell(row=2, column=len(ADDRESS), value="Description")  # written twice: row 3 says Description too
    for c in items:
        ws.cell(row=5, column=c, value="2x Kinds")
        ws.cell(row=2, column=c, value="Second description")
    wb.save(final)
    groups, _, found, _ = read(final)
    # The first Description is the Description block; the other one is a block of its own
    assert [x['label'] for x in found['Extra rows']] == ["Description (row 3)", "Kinds (row 5)"]
    item = next(i for g in groups.values() for i in g['items'] if not i.get('is_code'))
    assert item['row_5'] == "2x Kinds" and item['desc'] == "Second description" and item['row_3'] == "A1 poster"
    pdf = tmp_path / "labels.pdf"
    generate_packing_labels(dict(list(groups.items())[:1]), str(pdf), ['row_5', 'job_no', 'qty'])
    assert "2x Kinds" in fitz.open(pdf)[0].get_text()


def test_each_artwork_is_embedded_once_at_full_quality(files, tmp_path):
    """An artwork used on many labels goes into the PDF once, unchanged, and every box points at it: embedding it
    again for each box made a 340 MB PDF that took minutes to write."""
    from engine.packing_labels.packing_label_generator import generate_packing_labels
    from PIL import Image
    import io as _io
    final, sig, _ = files
    groups, _, _, _ = read(final, sig=sig)
    distinct = {id(d): d for g in groups.values() for i in g['items'] for d in i.get('thumbnails') or []}
    placed = sum(len(i.get('thumbnails') or []) for g in groups.values() for i in g['items'])
    assert placed > len(distinct)  # the sample repeats its artwork across stores
    pdf = tmp_path / "labels.pdf"
    generate_packing_labels(groups, str(pdf), ['thumbnail', 'job_no', 'qty'])
    doc = fitz.open(pdf)
    xrefs = {img[0] for page in doc for img in page.get_images()}
    assert len(xrefs) == len(distinct)
    sizes = {Image.open(_io.BytesIO(d.getvalue())).size for d in distinct.values()}
    assert {(doc.extract_image(x)['width'], doc.extract_image(x)['height']) for x in xrefs} == sizes  # full size kept


def test_page_without_grouping_keeps_sheet_order_and_every_code_column(client, files):
    """'Don't group': each pack's labels in the sheet's store order, no code in the Packing Spec panel, no dividers,
    but every label keeps its code boxes ('Code for Pack …' is a column of the pack: none is left off, Signature Links
    or not); the Item References (and so the CSVs and Stitch Labels) are the same as grouped."""
    final, sig, _ = files
    folder = run_page(client, final, sig, extra={"label_grouping": "store", "code_dividers": "1"})  # dividers ignored
    names = os.listdir(folder)
    doc = fitz.open(folder / next(n for n in names if n.startswith("Packing Labels Only")))
    labelmap = json.load(open(folder / next(n for n in names if n.endswith(".labelmap.json"))))
    labels = [" ".join(p.get_text().split()) for p in doc if "PACKING SPEC" in " ".join(p.get_text().split()) and "PACK AUDIT" not in p.get_text()]
    assert labels and not any(" CODE " in t for t in labels)
    assert all("Code for Pack" in t for t in labels)
    assert any("Code for pack 2: 1-2" in t for t in labels)  # the sub-group code column too
    packs = [pk for store in labelmap["stores"].values() for pk in store]
    label_pages = {p for pk in packs for p in pk["pdf_pages"]}
    blocks = {serial_of(pk["open360_code"])[2:] for pk in packs}  # 'T1P2': one audit page per tab and pack
    assert len(doc) == labelmap.get("report_pages", 1) + len(blocks) + len(label_pages)  # nothing else: no dividers
    for block in blocks:  # within a pack, stores in sheet order (the serial's store number)
        firsts = sorted((pk["pdf_pages"][0], int(pk["open360_code"][:2])) for pk in packs if serial_of(pk["open360_code"])[2:] == block)
        assert [n for _, n in firsts] == sorted(n for _, n in firsts)
    rows = open(folder / next(n for n in names if n.endswith("Open360.csv")), encoding="utf-8").read()
    assert re.search(r"03T1P2[A-Z]{4} Chadstone", rows) and re.search(r"01T2P1[A-Z]{4} Corio Village", rows)


def test_courier_labels_rendered_in_workers_match_rendering_here(tmp_path, monkeypatch):
    """Stitch Labels renders big jobs in worker processes: the pictures must be exactly the ones rendered here."""
    from engine.packing_labels import label_render
    doc = fitz.open()
    for n in range(6):
        page = doc.new_page(width=288, height=432)  # a 4x6 courier label
        page.insert_text((20, 60), f"Item Ref: {n + 1:02d}T1P1 Store {n}", fontsize=14)
        page.draw_rect(fitz.Rect(20, 100, 268, 160), color=(0, 0, 0), fill=(0, 0, 0))
    data = doc.tobytes()
    jobs = [("labels.pdf", n + 1) for n in range(6)]
    monkeypatch.setattr(label_render, "WORKER_MIN_LABELS", 1)
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    in_workers = label_render.render_many([("labels.pdf", data)], jobs, (200, 300), root)
    path = tmp_path / "labels.pdf"
    path.write_bytes(data)
    assert in_workers == label_render._render_pages({"labels.pdf": str(path)}, jobs, (200, 300))


def test_problems_name_the_store_on_their_row(tmp_path, monkeypatch):
    """'Row 13 Carindale', not only 'Row 13': on the problem itself, the pack's audit page and the preview's grouped
    line."""
    from engine.packing_labels.packing_label_generator import issue_where, group_issues
    from engine.packing_labels.horizontal_labels import generate_horizontal_labels
    final, sig = _spec_without_quantities(tmp_path, monkeypatch)
    groups, _, _, warnings = read(final, sig=sig)
    found = {w['rows']: w for w in warnings if w['text'].startswith("Packing Spec but no quantities")}
    assert found['13']['stores'] == [[13, "Carindale"]] and issue_where(found['13']) == "Row 13 Carindale"
    assert issue_where({'rows': '13-15', 'stores': [[13, "Carindale"], [15, "Garden City"]]}) == "Rows 13-15: Carindale, Garden City"
    assert "Pack 2: row 13 (Carindale)" in group_issues(warnings)[0]['where']
    pdf = tmp_path / "labels.pdf"
    notes = {(1, 2): [found['13']], (1, 1): [found['15']]}
    generate_horizontal_labels(groups, str(pdf), ['job_no', 'qty'], serials={k: g['serial'] for k, g in groups.items()},
                               add_dividers=False, notes=notes)
    audits = [p.get_text() for p in fitz.open(pdf) if "PACK AUDIT" in p.get_text()]
    assert any("Row 13 Carindale" in t for t in audits) and any("Row 15 Garden City" in t for t in audits)


def test_vertical_labels_get_a_pack_audit_page_per_tab(tmp_path):
    """Vertical files: 'PACK AUDIT' over 'Tab 1: VIC' before each tab's labels (instead of the plain divider), with
    the tab's figures, its Packing Specs and what needs a look by row and store."""
    import openpyxl
    from engine.packing_labels.packing_label_generator import parse_packing_data, generate_packing_labels
    path = tmp_path / "vertical.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for tab, rows in (("VIC", [("Corio Village", "5 Bacchus Marsh Rd", "Corio", "VIC", "3214", "OB170170170"),
                               ("Highpoint", "", "", "", "", "OB170170170"),
                               ("Chadstone", "1341 Dandenong Rd", "Chadstone", "VIC", "3148", "2 x CS450320150")]),
                      ("NSW", [("Parramatta", "159 Church St", "Parramatta", "NSW", "2150", "OB170170170")])):
        ws = wb.create_sheet(tab)
        ws.append(["Packing Spec", "Store Name", "Address Line 1", "Suburb", "State", "Postcode", "Job Number", "Qty", "Install"])
        for i, (store, line1, suburb, state, pc, spec) in enumerate(rows, 1):
            ws.append([spec, store, line1, suburb, state, pc, f"J900200-{i:02d}", 3, "N"])
    wb.save(path)
    groups, notes, tabs = {}, {}, {}
    for n, tab in enumerate(("VIC", "NSW"), 1):
        g, _, _, warnings = parse_packing_data(str(path), 1, tab)
        notes[n] = warnings
        for k, v in g.items():
            groups[f"{tab} - {k}"] = v
            tabs[f"{tab} - {k}"] = (n, tab)
    assert any(w.get('stores') == [[3, "Highpoint"]] for w in notes[1])  # 'Row 3 Highpoint': No address
    pdf = tmp_path / "labels.pdf"
    info = generate_packing_labels(groups, str(pdf), ['job_no', 'qty'], tabs=tabs, tab_notes=notes)
    texts = [p.get_text() for p in fitz.open(pdf)]
    audits = [i for i, t in enumerate(texts) if "PACK AUDIT" in t]
    assert len(audits) == 2 and audits[0] == 0  # one per tab, the first before any label
    assert "Tab 1: VIC" in texts[0] and "Tab 2: NSW" in texts[audits[1]]
    # The Packing Specs as badges, as the preview shows them; Items counts the item rows
    assert "PACKING SPECS" in texts[0] and "OB170170170 x 2" in texts[0] and "2 x CS450320150 x 1" in texts[0]
    assert "ITEMS" + chr(10) + "3" in texts[0]
    assert "Row 3 Highpoint" in texts[0] and not any("DIVIDER - NOT A LABEL" in t for t in texts)
    label_pages = sorted(p for v in info['pages'].values() for p in v)
    assert not set(label_pages) & set(audits)  # the label map points at labels, never at an audit page


def test_tab_card_lists_each_packs_packing_specs(files):
    final, sig, _ = files
    _, _, found, _ = read(final, sig=sig)
    by_pack = {p['pack']: dict(p['specs']) for p in found['specs_by_pack']}
    assert by_pack["Pack 1"] == {"OB600100100": 7}
    assert by_pack["Pack 2"] == {"CS450320150": 6, "2 x CS450320150": 1}


def test_preview_problems_not_fixed_go_on_the_audit_pages(client, tmp_path, monkeypatch):
    """A consignment problem still there at Generate (here: no suburb) is on the audit page of each pack it has
    labels in, with its row and store."""
    import fixtures
    tabs = copy.deepcopy(fixtures.TABS)
    stores = tabs["Spring Sale"]["stores"]
    for i, st in enumerate(stores):
        if st[0] == "Highpoint":
            stores[i] = st[:3] + ("",) + st[4:]  # no suburb
    monkeypatch.setattr(fixtures, "TABS", tabs)
    codes = build_final_file(tmp_path / FINAL)
    build_signature_links(tmp_path / SIG, codes)
    folder = run_page(client, tmp_path / FINAL, tmp_path / SIG)
    doc = fitz.open(folder / next(n for n in os.listdir(folder) if n.startswith("Packing Labels Only")))
    audits = [p.get_text() for p in doc if "PACK AUDIT" in p.get_text() and "Spring" not in p.get_text()]
    hits = [t for t in audits if "Address incomplete: no suburb (not fixed in the preview)" in t]
    assert len(hits) == 2  # Highpoint has a label in pack 1 and in pack 2 of the first tab
    assert all("Row 10 Highpoint" in t and "Consignment #" in t for t in hits)
    assert all("Check before packing" in t for t in hits)


def test_dividers_carry_no_barcodes(client, files):
    """The items on these labels have their own barcodes: the code dividers (installer tabs too) have none."""
    final, sig, _ = files
    folder = run_page(client, final, sig)
    pdf = next(n for n in os.listdir(folder) if n.startswith("Packing Labels Only"))
    dividers = [p for p in fitz.open(folder / pdf) if "Matched Labels" in p.get_text()]
    assert dividers and not any("J4767" in p.get_text() for p in dividers)  # Easter is an installer tab


def _sheets_project(projects, final, sig, stage=False):
    """A Packing Sheets project as the dashboard sees it: its Final file (and a Stage file), its Signature Links."""
    import shutil
    folder = projects / "Spring rollout_Job-J476699-01_261017_1000"
    folder.mkdir()
    shutil.copy(final, folder / final.name)
    if stage:
        shutil.copy(final, folder / f"Stage - 1 - {final.name}")
    shutil.copy(sig, folder / sig.name)
    (folder / "Packing Sheet_J476699-01.xlsx").write_bytes(b"")
    return folder


def test_dashboard_offers_generate_labels_and_the_shuffler_for_packing_sheets(client, files):
    c, app_module, projects = client
    final, sig, _ = files
    folder = _sheets_project(projects, final, sig)
    page = c.get("/").data.decode()
    assert 'name="from_project" value="Spring rollout_Job-J476699-01_261017_1000"' in page
    assert f'name="final_file" value="{final.name}"' in page and "Label Shuffler" in page
    (folder / f"Stage - 1 - {final.name}").write_bytes(final.read_bytes())  # a second Final file: pick one
    page = c.get("/").data.decode()
    assert 'id="labels-Spring rollout_Job-J476699-01_261017_1000"' in page and 'type="radio" name="final_file"' in page
    shuffler = c.get("/pdf?project=Spring rollout_Job-J476699-01_261017_1000").data.decode()
    assert 'value="Spring rollout_Job-J476699-01_261017_1000" selected' in shuffler


def test_labels_from_a_packing_sheets_project_use_its_final_file_and_signature_links(client, files):
    c, app_module, projects = client
    final, sig, _ = files
    folder = _sheets_project(projects, final, sig, stage=True)
    before = sorted(os.listdir(folder))
    r = c.post("/packing-labels/horizontal", data={"from_project": folder.name}, follow_redirects=True)
    assert "Final file to make the labels from" in r.data.decode()  # two Final files: which one has to be said
    r = c.post("/packing-labels/horizontal", data={"from_project": folder.name, "final_file": final.name},
               follow_redirects=True)
    page = r.data.decode()
    assert 'value="8"' in page and "Spring Sale" in page  # tabs and header rows from its Label Plugin
    upload = projects / app_module.secure_filename(final.name)
    assert upload.exists() and os.path.exists(app_module._signature_path(str(upload)))
    assert sorted(os.listdir(folder)) == before  # the project itself is as it was
    picker = c.get("/packing-labels/horizontal").data.decode()
    assert 'id="pl-project-pick"' in picker and folder.name in picker


def test_kinds_and_item_number_rows_go_in_the_item_box_by_default(client, tmp_path):
    """Horizontal files: a 'Kinds' row after the description and an 'Item Number' row after the dimensions, when the
    sheet has them; other rows stay in the pool on the left."""
    from openpyxl import load_workbook
    c, app_module, projects = client
    path = tmp_path / FINAL
    codes = build_final_file(path)
    build_signature_links(tmp_path / SIG, codes)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        jobs = [x for x in range(1, ws.max_column + 1) if re.match(r"J\d", str(ws.cell(row=7, column=x).value or ""))]
        for label, row in (("Kinds", 5), ("Item Number", 2), ("Finish", 1)):
            ws.cell(row=row, column=7, value=label)
            for n, x in enumerate(jobs, 1):
                ws.cell(row=row, column=x, value=f"{label[0]}{n}")
    wb.save(path)
    with open(path, "rb") as f, open(tmp_path / SIG, "rb") as s:
        c.post("/packing-labels/horizontal", data={"file": (f, FINAL), "signature_file": (s, SIG)},
               content_type="multipart/form-data", follow_redirects=True)
    form = {"filename": app_module.secure_filename(FINAL), "all_tabs": ["Spring Sale"], "selected_tabs": ["Spring Sale"],
            "header_row_Spring Sale": "8", "preview": "1"}
    page = c.post("/packing-labels/horizontal", data=form, follow_redirects=True).data.decode()
    layout = page[page.index('id="attribute-list"'):page.index('id="attribute_order"')]
    order = re.findall(r'data-id="([^"]+)"', layout)
    assert order.index("desc") + 1 == order.index("row_5") and order.index("dimension") + 1 == order.index("row_2")
    assert "row_1" not in order and 'data-id="row_1"' in page  # Finish waits on the left


def test_packing_labels_left_without_a_courier_label_are_listed_with_their_pages(client, files):
    """Three courier labels never came: the result and the report name those packing labels and the page each is on."""
    import csv
    from engine.packing_labels.label_stitcher import stitch
    final, sig, _ = files
    folder = run_page(client, final, sig)
    names = os.listdir(folder)
    labelmap = json.load(open(folder / next(n for n in names if n.endswith(".labelmap.json"))))
    rows = list(csv.DictReader(open(folder / next(n for n in names if n.endswith("Open360.csv")), encoding="utf-8")))
    refs = [r["Item Reference"] for r in rows for _ in range(int(r["Item Quantity"] or 1))]
    courier = fitz.open()
    for ref in refs[3:]:  # the first three never came
        page = courier.new_page(width=288, height=432)
        page.insert_text((20, 300), f"Item Ref: {ref}", fontsize=10)
    out = folder / "Complete.pdf"
    result = stitch(str(folder), labelmap, [("courier.pdf", courier.tobytes())], str(out))
    missing = [m for m in result["missing_pages"] if not m["not_in_csv"]]
    assert len(missing) == 3 and all(m["page"] for m in missing)
    doc = fitz.open(out)
    for m in missing:  # the page named is that packing label, with no courier label on it
        text = " ".join(doc[m["page"] - 1].get_text().split())
        assert m["store"] in text and "PACKING SPEC" in text
    report = "".join(doc[n].get_text() for n in range(result["report_pages"]))
    assert all(f"{m['slot']} - page {m['page']}" in report.replace("\n", " ") for m in missing)


def test_packs_to_be_collected_are_never_without_a_courier_label(client, files):
    """Carindale's address says 'Kyle collect from warehouse': its labels get no courier label by design, so a stitch
    with every courier label that came counts them as 'Kyle to collect', not 'Without a courier label', and the
    matching page doesn't wait for them. A label map made before it said so is read off the labels' yellow panel."""
    import csv
    import openpyxl
    from engine.packing_labels.label_stitcher import stitch
    final, sig, _ = files
    wb = openpyxl.load_workbook(final)
    ws = wb["Spring Sale"]
    header = {str(ws.cell(row=8, column=c).value or '').strip(): c for c in range(1, ws.max_column + 1)}
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Carindale")
    ws.cell(row=row, column=header["Address 1"], value="Kyle collect from warehouse")
    for name in ("Suburb", "State", "Postcode"):
        ws.cell(row=row, column=header[name]).value = None
    wb.save(final)
    folder = run_page(client, final, sig)
    names = os.listdir(folder)
    map_path = folder / next(n for n in names if n.endswith(".labelmap.json"))
    labelmap = json.load(open(map_path))
    assert [x["collect"] for x in labelmap["not_in_courier_csv"]] == ["KYLE TO COLLECT"]
    rows = list(csv.DictReader(open(folder / next(n for n in names if n.endswith("Open360.csv")), encoding="utf-8")))
    courier = fitz.open()
    for ref in [r["Item Reference"] for r in rows for _ in range(int(r["Item Quantity"] or 1))]:
        courier.new_page(width=288, height=432).insert_text((20, 300), f"Item Ref: {ref}", fontsize=10)
    older = {**labelmap, "not_in_courier_csv": [{k: v for k, v in x.items() if k != "collect"}
                                                for x in labelmap["not_in_courier_csv"]]}
    for n, lm in enumerate((labelmap, older)):
        out = folder / f"Complete {n}.pdf"
        result = stitch(str(folder), lm, [("courier.pdf", courier.tobytes())], str(out))
        assert result["without_label"] == [] and result["complete"]
        assert not result["session"]["waiting"]
        report = " ".join(" ".join(fitz.open(out)[p].get_text().split()) for p in range(result["report_pages"]))
        assert "Kyle to collect 1 label" in report.replace(":", "")


def _courier_pdf(refs, note=""):
    """A courier label PDF: one page per Item Reference (note: more text, e.g. a new consignment number)."""
    doc = fitz.open()
    for ref in refs:
        doc.new_page(width=288, height=432).insert_text((20, 300), f"Item Ref: {ref} {note}", fontsize=10)
    return doc.tobytes()


def test_each_stitch_builds_on_an_earlier_one_and_a_pack_booked_again_is_asked_about(client, files):
    """Courier labels that come in batches: each stitch keeps its own Complete Labels and builds on the latest one by
    default (its labels stay, the new PDF fills the rest), with a history of what was placed and what's left; a pack
    an earlier stitch labelled that gets another courier label waits for 'new' or 'keep' before Finalise; a stitch
    can be built on an older one."""
    import csv
    import io as io_module
    c, app_module, projects = client
    final, sig, _ = files
    folder = run_page(client, final, sig)
    names = os.listdir(folder)
    rows = list(csv.DictReader(open(folder / next(n for n in names if n.endswith("Open360.csv")), encoding="utf-8")))
    refs = [r["Item Reference"] for r in rows for _ in range(int(r["Item Quantity"] or 1))]
    first, rest = refs[:4], refs[4:]

    def stitch_page(pdf_name, data, **form):
        return c.post(f"/stitch/{folder.name}", data={"labels": (io_module.BytesIO(data), pdf_name), **form},
                      content_type="multipart/form-data", follow_redirects=True).data.decode()

    def finalise(**body):
        r = c.post(f"/stitch/{folder.name}/finalise", json={"matches": {}, "add_unmatched": True, **body})
        if r.status_code == 202:
            assert app_module.STITCH_DONE[str(folder)].wait(30)
        return r

    def history():
        lm = app_module._stitch_project(folder.name)[2]
        return app_module._history(str(folder), lm)

    # 1. The first batch: the rest are still to come, so the matching page; Finalise makes stitch 1
    page = stitch_page("batch-a.pdf", _courier_pdf(first))
    assert "Match Courier Labels" in page
    assert finalise().status_code == 202
    (one,) = history()
    assert one["base"] is None and one["placed"] == 4 and len(one["remaining"]) == len(rest)
    # 2. The rest, on top of stitch 1 (the default): everything placed, straight away
    batch_b = _courier_pdf(rest)
    page = stitch_page("batch-b.pdf", batch_b)
    assert "placed</h2>" in page
    one, two = history()
    assert two["base"] == 1 and two["placed"] == len(refs) and two["new"] == len(rest) and two["remaining"] == []
    assert (folder / one["output"]).exists() and (folder / two["output"]).exists()  # every stitch's file is kept
    assert "Stitch 2" in c.get(f"/stitch/{folder.name}").data.decode()
    # 3. One pack booked again: asked about, and Finalise waits for the choice
    page = stitch_page("batch-c.pdf", _courier_pdf(refs[:1], "Consignment 99 rebooked"))
    assert "Booked again" in page
    session = json.load(open(app_module._run_file(str(folder), "session"), encoding="utf-8"))
    (again,) = session["rebooked"]
    assert again["earlier"]["file"] == "batch-a.pdf" and again["new"]["file"] == "batch-c.pdf" and again["choice"] is None
    r = finalise()
    assert r.status_code == 400 and "booked again" in r.get_json()["error"]
    assert finalise(rebook={again["new"]["key"]: "new"}).status_code == 202
    three = history()[-1]
    assert three["base"] == 2 and three["rebook"] == {again["new"]["key"]: "new"} and three["new"] == 1
    lm = app_module._stitch_project(folder.name)[2]
    plan = app_module.plan_stitch(str(folder), lm, app_module._courier_pdfs(str(folder), [f for l in three["layers"] for f in l]),
                                  None, three["layers"], three["rebook"])
    assert next(p for p in plan["placed"] if p["item_reference"] == again["slot"])["file"] == "batch-c.pdf"
    # 4. Built on stitch 1 instead: batch B again fills the rest; batch C isn't in it
    stitch_page("batch-b.pdf", batch_b, base="1")  # the same file again: kept once, under its own name
    four = history()[-1]
    assert four["base"] == 1 and four["layers"] == [["batch-a.pdf"], ["batch-b.pdf"]] and four["remaining"] == []


def test_units_on_no_label_say_which_cells(tmp_path):
    """A pack marked x while the store has quantities in it: the units-on-no-label error names each cell (J9 (2))."""
    from openpyxl import load_workbook
    build_final_file(tmp_path / FINAL)
    wb = load_workbook(tmp_path / FINAL)
    ws = wb["Spring Sale"]
    ws.cell(row=9, column=8, value="x")  # Corio Village (row 9): pack 1's Packing Spec, its quantities still there
    wb.save(tmp_path / FINAL)
    with pytest.raises(PackCheckError) as e:
        read(tmp_path / FINAL)
    issue = next(i for i in e.value.issues if "on no label" in i["text"])
    assert issue["rows"] == "9" and "Quantities" in issue["col"]
    cells = re.findall(r"([A-Z]+)9 \((\d+)\)", issue["detail"])
    assert cells and sum(int(q) for _, q in cells) == int(issue["text"].split()[0])  # every lost unit, cell by cell
    assert all(letter in issue["col"] for letter, _ in cells)


def test_the_other_courier_csv_is_made_later_from_the_details_kept(client, files):
    """Generate keeps what both courier CSVs are made from: the one not there is made from it in the same project,
    exactly as Generate would have written it, and the label map lists it."""
    c, app_module, projects = client
    final, sig, _ = files
    folder = run_page(client, final, sig)
    names = os.listdir(folder)
    recipe = next(n for n in names if n.endswith(".courier.json"))
    openfreight = next(n for n in names if n.endswith("OpenFreight.csv"))
    original = (folder / openfreight).read_bytes()
    (folder / openfreight).unlink()  # only the Open360 CSV is left
    page = c.get("/").data.decode()
    assert "Create OpenFreight CSV" in page and 'name="kind" value="openfreight"' in page
    r = c.post(f"/project/{folder.name}/courier-csv", data={"kind": "openfreight"})
    assert r.status_code == 302 and f"open={folder.name.replace(' ', '+')}" in r.headers["Location"].replace("%20", "+")
    made = next(n for n in os.listdir(folder) if n.endswith("OpenFreight.csv"))
    assert (folder / made).read_bytes() == original  # the same file Generate wrote
    labelmap = json.load(open(folder / next(n for n in os.listdir(folder) if n.endswith(".labelmap.json"))))
    assert made in labelmap["files"]["courier_csvs"]
    assert "Create OpenFreight CSV" not in c.get("/").data.decode()  # both kinds there now
    assert json.load(open(folder / recipe, encoding="utf-8"))["made"]["openfreight"] == made


def test_labels_from_a_packing_sheets_project_are_saved_in_that_project(client, files):
    c, app_module, projects = client
    final, sig, _ = files
    folder = _sheets_project(projects, final, sig)
    before = set(os.listdir(folder))
    c.post("/packing-labels/horizontal", data={"from_project": folder.name, "final_file": final.name}, follow_redirects=True)
    filename = app_module.secure_filename(final.name)
    form = {"filename": filename, "from_project": folder.name, "all_tabs": ["Spring Sale", "Easter"],
            "selected_tabs": ["Spring Sale", "Easter"], "header_row_Spring Sale": "8", "header_row_Easter": "8"}
    preview = c.post("/packing-labels/horizontal", data={**form, "preview": "1"}, follow_redirects=True).data.decode()
    assert "Saved in project" in preview and 'name="project_name"' not in preview
    gen = {**form, "generate": "1", "who_pays": "S", "service_code": app_module.SERVICE_STORE.services()[0]['code'],
           "reference": "J476699", "courier_csv": ["open360"], "code_dividers": "1",
           "attribute_order": "thumbnail,desc,dimension,job_no,barcode,qty"}
    r = c.post("/packing-labels/horizontal", data=gen, follow_redirects=True)
    assert "Labels Successfully Generated" in r.data.decode()
    assert [p.name for p in projects.iterdir() if p.is_dir()] == [folder.name]  # no new project
    after = set(os.listdir(folder))
    assert before <= after  # the packing sheets files are all still there
    new = after - before
    assert any(n.startswith("Packing Labels Only") for n in new) and any(n.endswith(".labelmap.json") for n in new)
    assert any(n.endswith("Open360.csv") for n in new) and any(n.endswith(".courier.json") for n in new)
    assert not (projects / filename).exists()  # the copy made for the page went
    assert not any(n.startswith("Used for labels") for n in new)  # unchanged: no copy kept
    assert app_module.sheets_runs == []  # the project has its Packing Sheet and Signature Links already
    page = c.get("/").data.decode()
    assert "Stitch Labels" in page and "Create OpenFreight CSV" in page and "Label Shuffler" in page


def _run_without_signature(client, final, project, extra=None):
    """The page with a distribution file and no Signature Links file: scan, preview (header row 8), Generate.
    Returns (project folder, preview page, result page)."""
    c, app_module, projects = client
    with open(final, "rb") as f:
        c.post("/packing-labels/horizontal", data={"file": (f, final.name)}, content_type="multipart/form-data",
               follow_redirects=True)
    form = {"filename": app_module.secure_filename(final.name), "all_tabs": ["Spring Sale", "Easter"],
            "selected_tabs": ["Spring Sale", "Easter"], "header_row_Spring Sale": "8", "header_row_Easter": "8"}
    preview = c.post("/packing-labels/horizontal", data={**form, "preview": "1"}, follow_redirects=True).data.decode()
    gen = {**form, "generate": "1", "who_pays": "S", "service_code": app_module.SERVICE_STORE.services()[0]['code'],
           "reference": "J476699", "project_name": project, "courier_csv": ["open360"],
           "attribute_order": "thumbnail,desc,dimension,job_no,barcode,qty", **(extra or {})}
    result = c.post("/packing-labels/horizontal", data=gen, follow_redirects=True).data.decode()
    assert "Labels Successfully Generated" in result, result[:1500]
    folder = next(p for p in projects.iterdir() if p.is_dir() and p.name.startswith(project))
    return folder, preview, result


@pytest.mark.parametrize("grouping", ["code", "store"])
def test_labels_without_signature_links_also_make_the_packing_sheet_and_signature_links(client, tmp_path, grouping):
    """No Signature Links file sent: Packing Sheets runs on the file too (every pack of each tab, from the layout the
    labels were read with), whatever the label order, and its Packing Sheet and Signature Links go into the project."""
    c, app_module, projects = client
    final = tmp_path / "Spring rollout.xlsx"
    build_final_file(final, codes=False)
    folder, _, result = _run_without_signature(client, final, f"Raw {grouping}", {"label_grouping": grouping})
    (run,) = app_module.sheets_runs
    assert run["tabs"] == ["Spring Sale", "Easter"]
    spring = run["inputs"]["Spring Sale"]
    assert (spring["start"], spring["job"], spring["store"]) == ("H8", "J7", "A")  # first Packing Spec, first job number
    assert spring["selected_packs"] == ["1", "2"] and run["inputs"]["Easter"]["job"] == "I7"
    names = os.listdir(folder)
    assert "Packing Sheet_J476699.xlsx" in names and "Signature links_J476699.xlsx" in names
    assert not any(n.startswith("Final ") for n in names)  # only the two files
    assert "Download Packing Sheet" in result and "Download Signature Links" in result
    assert any(n.startswith("Packing Labels Only") for n in names)  # the labels as ever


def test_labels_with_signature_links_make_no_packing_sheet(client, files):
    c, app_module, projects = client
    final, sig, _ = files
    folder = run_page(client, final, sig)
    assert app_module.sheets_runs == [] and not any(n.startswith("Packing Sheet") for n in os.listdir(folder))


def test_a_final_file_without_its_signature_links_isnt_coded_again(client, tmp_path):
    """A Final file (code columns in already) sent without its Signature Links: no Packing Sheet made from it (its codes
    would be coded again); the result page says why."""
    c, app_module, projects = client
    final = tmp_path / "Final J476699 Spring rollout.xlsx"
    build_final_file(final)
    _, _, result = _run_without_signature(client, final, "Final run")
    assert app_module.sheets_runs == []
    assert "already has packing-sheet code columns" in result


def test_blueprint_preview_comes_first_with_each_packs_columns(client, tmp_path):
    """Horizontal labels: a Blueprint Preview per tab, as Packing Sheets shows it, before the Data Previews."""
    c, app_module, projects = client
    final = tmp_path / "Spring rollout.xlsx"
    build_final_file(final, codes=False)
    _, preview, _ = _run_without_signature(client, final, "Blueprint")
    text = " ".join(re.sub(r"<[^>]+>", " ", preview).split())
    assert text.index("Blueprint Previews") < text.index("Data Previews")
    assert "Last Column: Q" in text and "Last Row: 16" in text and "Total Stores: 8" in text
    assert "Packaging stocks are in H to I." in text and "1 spans from J to N" in text and "2 spans from O to Q" in text
    assert "more than one pack group" not in text  # no job number is in two packs


SHARED = "J476699-01 is in Spring Sale | Pack 1, Spring Sale | Pack 2, Easter | Pack 1"


def _shared_job_file(path):
    """The sample file (no code columns) with J476699-01, pack 1's first item, also on pack 2's first item (O7) and on
    Easter's first item (I7): one job number in three pack groups, two tabs."""
    import openpyxl
    build_final_file(path, codes=False)
    wb = openpyxl.load_workbook(path)
    wb["Spring Sale"]["O7"] = "J476699-01"
    wb["Easter"]["I7"] = "J476699-01"
    wb.save(path)


def _text(page):
    return " ".join(re.sub(r"<[^>]+>", " ", page).split())


def test_blueprint_previews_warn_about_a_job_number_in_more_than_one_pack(client, tmp_path):
    """The download page's warning, on both Blueprint Previews before anything is made: horizontal labels and
    Packing Sheets name the same job number and the same pack groups, across tabs. Generate isn't locked by it."""
    c, app_module, projects = client
    final = tmp_path / "Spring rollout.xlsx"
    _shared_job_file(final)
    with open(final, "rb") as f:
        c.post("/packing-labels/horizontal", data={"file": (f, final.name)}, content_type="multipart/form-data",
               follow_redirects=True)
    form = {"filename": app_module.secure_filename(final.name), "all_tabs": ["Spring Sale", "Easter"],
            "selected_tabs": ["Spring Sale", "Easter"], "header_row_Spring Sale": "8", "header_row_Easter": "8"}
    labels = _text(c.post("/packing-labels/horizontal", data={**form, "preview": "1"}, follow_redirects=True).data.decode())
    assert "These job numbers appear in more than one pack group" in labels and SHARED in labels
    assert labels.index(SHARED) < labels.index("Data Previews")  # with the Blueprint Previews, at the top

    sheets = _text(c.post("/preview", data={
        "filename": app_module.secure_filename(final.name), "all_tabs": ["Spring Sale", "Easter"],
        "selected_tabs": ["Spring Sale", "Easter"], "start_Spring Sale": "H8", "job_Spring Sale": "J7",
        "store_Spring Sale": "A", "start_Easter": "H8", "job_Easter": "I7", "store_Easter": "A"},
        follow_redirects=True).data.decode())
    assert "These job numbers appear in more than one pack group" in sheets and SHARED in sheets
    assert sheets.index("Blueprint Previews") < sheets.index(SHARED) < sheets.index("Spring Sale Last Column")
    assert "Cannot Generate" not in sheets and "Generate Packing Sheet" in sheets


def test_pack_job_numbers_keep_each_packs_own_once():
    """Packing Sheets' job numbers per pack: blank cells skipped, a job number twice in one pack kept once (that's
    two kinds, not a shared job number); one in two packs is what's warned about."""
    from engine.packing_sheets.matrix_engine import pack_job_numbers, job_groups, find_shared_job_numbers
    row = [None, "J1-01", " J1-02 ", "", "J1-02", "J1-01", None]
    cols, packs = pack_job_numbers(row, [{"name": "1", "start": 2, "end": 5}, {"name": "2", "start": 6, "end": 8}])
    assert packs == [("1", ["J1-01", "J1-02"]), ("2", ["J1-01"])]
    assert cols == {2: "J1-01", 3: "J1-02", 5: "J1-02", 6: "J1-01"}
    assert find_shared_job_numbers(job_groups("Tab", packs)) == ["J1-01 is in Tab | Pack 1, Tab | Pack 2"]


def test_packing_sheets_file_names_use_the_job_number_itself():
    from engine.packing_sheets.matrix_engine import file_job
    assert file_job("J481189-01") == file_job("J481189-001") == file_job("J481189") == "J481189"
    assert file_job("Window decal") == "Window decal"  # not a job number: as written


def _into_project_form(c, app_module, folder, final):
    c.post("/packing-labels/horizontal", data={"from_project": folder.name, "final_file": final.name}, follow_redirects=True)
    form = {"filename": app_module.secure_filename(final.name), "from_project": folder.name,
            "all_tabs": ["Spring Sale", "Easter"], "selected_tabs": ["Spring Sale", "Easter"],
            "header_row_Spring Sale": "8", "header_row_Easter": "8"}
    gen = {**form, "generate": "1", "who_pays": "S", "service_code": app_module.SERVICE_STORE.services()[0]['code'],
           "reference": "J476699", "courier_csv": ["open360"], "code_dividers": "1",
           "attribute_order": "thumbnail,desc,dimension,job_no,barcode,qty"}
    return form, gen


def test_a_failed_generate_into_a_project_leaves_the_project_as_it_was(client, files):
    """A failure takes back only what the run wrote: the Packing Sheets project (Final file, Packing Sheet, Signature
    Links) stays, and so does the page's copy, so Generate can be tried again."""
    c, app_module, projects = client
    final, sig, _ = files
    folder = _sheets_project(projects, final, sig)
    before = set(os.listdir(folder))
    form, gen = _into_project_form(c, app_module, folder, final)
    r = c.post("/packing-labels/horizontal", data={**gen, "service_code": "NOPE1"}, follow_redirects=True)
    assert "Unknown service code" in r.data.decode()
    assert folder.is_dir() and set(os.listdir(folder)) == before
    assert (projects / form["filename"]).exists()  # the copy being worked on, for a retry


def test_only_a_real_packing_sheets_project_takes_labels(client, files):
    c, app_module, projects = client
    for bad in ("..", ".", "Not a project"):
        r = c.post("/packing-labels/horizontal", data={"from_project": bad, "final_file": "Final x.xlsx"}, follow_redirects=True)
        assert "on the dashboard any more: pick one from the list" in r.data.decode(), bad
    assert app_module._sheet_project("..") == "" and app_module._sheet_project("../projects") == ""


def test_an_edited_final_file_copy_is_kept_with_the_labels(client, files):
    """Fixes made with Open Excel on the page go into the labels; the copy they were made from stays in the project."""
    from openpyxl import load_workbook
    c, app_module, projects = client
    final, sig, _ = files
    folder = _sheets_project(projects, final, sig)
    form, gen = _into_project_form(c, app_module, folder, final)
    copy = projects / form["filename"]
    wb = load_workbook(copy)
    wb.worksheets[0]["AZ1"] = "checked by the warehouse"  # an edit away from the grid
    wb.save(copy)
    r = c.post("/packing-labels/horizontal", data=gen, follow_redirects=True)
    assert "Labels Successfully Generated" in r.data.decode()
    kept = [n for n in os.listdir(folder) if n.startswith("Used for labels - ")]
    assert len(kept) == 1 and kept[0] in r.data.decode() and not copy.exists()
    assert load_workbook(folder / kept[0]).worksheets[0]["AZ1"].value == "checked by the warehouse"
    assert app_module._sheet_outputs(str(folder))[0] == [final.name]  # not taken for a Final file


def test_the_other_csv_warns_when_the_project_csv_was_changed_and_never_overwrites(client, files):
    """An Open360 CSV edited after Generate (saved by Excel in its own code page): the card says the new OpenFreight
    CSV won't have those changes and makes it only when asked again, as a new file; the edited one stays as it was."""
    import time
    c, app_module, projects = client
    final, sig, _ = files
    folder = run_page(client, final, sig)
    open360 = next(n for n in os.listdir(folder) if n.endswith("Open360.csv"))
    (folder / next(n for n in os.listdir(folder) if n.endswith("OpenFreight.csv"))).unlink()
    edited = (folder / open360).read_text(encoding="utf-8-sig").replace("Corio Village", "Corio Village Café")
    (folder / open360).write_bytes(edited.encode("cp1252"))  # as Excel's "CSV (Comma delimited)" saves it
    later = time.time() + 30
    os.utime(folder / open360, (later, later))
    kinds, _ = app_module._courier_csvs(str(folder), os.listdir(folder))
    assert kinds == {"open360": open360}  # still known for what it is
    r = c.post(f"/project/{folder.name}/courier-csv", data={"kind": "openfreight"})
    assert "stale=" in r.headers["Location"] and not any(n.endswith("OpenFreight.csv") for n in os.listdir(folder))
    page = c.get(r.headers["Location"]).data.decode()
    assert "was changed after the labels were generated" in page and "Make it anyway" in page
    c.post(f"/project/{folder.name}/courier-csv", data={"kind": "openfreight", "confirm": "1"})
    assert any(n.endswith("OpenFreight.csv") for n in os.listdir(folder))
    assert (folder / open360).read_bytes() == edited.encode("cp1252")  # the edited file untouched


def test_unused_name_never_returns_a_file_that_is_there(tmp_path):
    from engine.core.core_math import unused_name
    assert unused_name(str(tmp_path), "J1 - OpenFreight.csv") == "J1 - OpenFreight.csv"
    (tmp_path / "J1 - OpenFreight.csv").write_text("x")
    (tmp_path / "J1 - OpenFreight (2).csv").write_text("x")
    assert unused_name(str(tmp_path), "J1 - OpenFreight.csv") == "J1 - OpenFreight (3).csv"


def test_courier_import_keeps_the_details_for_the_other_csv(client, tmp_path):
    """Courier Import with only the Open360 CSV: the OpenFreight one is made later from the project card, quickly."""
    import time
    from openpyxl import Workbook
    c, app_module, projects = client
    wb = Workbook()
    ws = wb.active
    ws.title = "Openings"
    ws["A1"] = "J476702 - Store openings"
    for col, h in enumerate(["Store Name", "Address 1", "Suburb", "State", "Postcode", "1"], 1):
        ws.cell(row=3, column=col, value=h)
    for r, (store, line, sub, st, pc) in enumerate([("Corio Village", "1 Bacchus Marsh Rd", "Corio", "VIC", "3214"),
                                                    ("Knox City", "425 Burwood Hwy", "Wantirna South", "VIC", "3152")], 4):
        for col, v in enumerate([store, line, sub, st, pc, "OB600100100"], 1):
            ws.cell(row=r, column=col, value=v)
    path = tmp_path / "openings.xlsx"
    wb.save(path)
    with open(path, "rb") as f:
        c.post("/courier-import", data={"file": (f, "openings.xlsx")}, content_type="multipart/form-data", follow_redirects=True)
    form = {"filename": "openings.xlsx", "all_tabs": ["Openings"], "selected_tabs": ["Openings"], "header_row_Openings": "3"}
    c.post("/courier-import", data={**form, "preview": "1"}, follow_redirects=True)
    r = c.post("/courier-import", data={**form, "generate": "1", "who_pays": "S", "reference": "J476702",
                                        "service_code": app_module.SERVICE_STORE.services()[0]['code'],
                                        "project_name": "Openings", "courier_csv": ["open360"],
                                        "progress_id": "courier-import-0001"}, follow_redirects=True)
    assert "Courier Import File Generated" in r.data.decode()
    # The loading overlay followed it: every step, its log lines (the same as the Logs page), 100 % at the end
    shown = c.get("/api/progress/courier-import-0001").get_json()
    lines = [x['text'] for x in shown['lines']]
    assert shown['finished'] and shown['percent'] == 100 and shown['counted']
    assert any(t.startswith("Generate openings.xlsx") for t in lines) and any(t.startswith("Open360 CSV") for t in lines)
    assert any(t.startswith("details for the other courier CSV") for t in lines)
    assert any(t.startswith("Generate done in") for t in lines)
    folder = next(p for p in projects.iterdir() if p.is_dir())
    assert "Create OpenFreight CSV" in c.get("/").data.decode()
    started = time.perf_counter()
    c.post(f"/project/{folder.name}/courier-csv", data={"kind": "openfreight"})
    assert time.perf_counter() - started < 1.0  # a moment, not a Generate
    made = next(n for n in os.listdir(folder) if n.endswith("OpenFreight.csv"))
    text = (folder / made).read_text(encoding="utf-8-sig")
    assert "T1P1 Corio Village" in text and "T1P1 Knox City" in text  # the same Item References as the Open360 CSV


def test_project_record_says_what_made_the_project_and_sets_the_card(client, files):
    """Each tool notes itself in the project's record; the dashboard takes the card's colours and buttons from it."""
    from engine.core import project_record
    c, app_module, projects = client
    final, sig, _ = files
    folder = _sheets_project(projects, final, sig)
    project_record.record(str(folder), 'packing_sheets', [final.name])
    c.post("/packing-labels/horizontal", data={"from_project": folder.name, "final_file": final.name}, follow_redirects=True)
    form = {"filename": app_module.secure_filename(final.name), "from_project": folder.name, "all_tabs": ["Spring Sale", "Easter"],
            "selected_tabs": ["Spring Sale", "Easter"], "header_row_Spring Sale": "8", "header_row_Easter": "8"}
    c.post("/packing-labels/horizontal", data={**form, "preview": "1"}, follow_redirects=True)
    c.post("/packing-labels/horizontal", data={**form, "generate": "1", "who_pays": "S", "reference": "J476699",
                                               "service_code": app_module.SERVICE_STORE.services()[0]['code'],
                                               "courier_csv": ["open360"], "attribute_order": "desc,job_no,barcode,qty"},
           follow_redirects=True)
    rec = project_record.read(str(folder))
    assert set(rec["features"]) == {"packing_sheets", "labels"}
    assert rec["history"][-1]["details"] == {"layout": "horizontal", "from_final_file": True}
    assert any(f.endswith(".labelmap.json") for f in rec["features"]["labels"]["files"])
    page = c.get("/").data.decode()
    card = page[page.index(f'id="content-{folder.name}"'):]
    card = card[:card.index("Delete Entire Project")]
    assert all(b in card for b in ("Stitch Labels", "Generate Labels", "Label Shuffler", "Create OpenFreight CSV"))
    assert "project-feature" not in rec  # nothing odd in it
    r = c.post(f"/project/{folder.name}/courier-csv", data={"kind": "openfreight"})
    assert r.status_code == 302 and "courier_csv" in project_record.read(str(folder))["features"]


def test_a_project_from_before_records_is_read_from_its_files(client, files):
    from engine.core import project_record
    c, app_module, projects = client
    final, sig, _ = files
    folder = _sheets_project(projects, final, sig)  # no record
    assert project_record.read(str(folder)) is None
    page = c.get("/").data.decode()
    assert "project-kind-sheets" in page and "Label Shuffler" in page


def test_same_name_at_different_addresses_can_be_kept_separate():
    """Two stores with one name at different addresses: grouped to merge, unless both were kept separate."""
    import app as app_module

    def con(n, line, sid):
        return {'number': n, 'destination': {'receiver': 'Quinn and Co - Kew', 'line1': line, 'line2': '', 'suburb': 'KEW',
                                             'state': 'VIC', 'postcode': '3101', 'country': 'AU'},
                'cartons': [{'source_id': sid}], 'service_code': '', 'service_code_used': 'IPECX'}
    grouped = app_module._receiver_groups([con(1, '212 High Street', 'a'), con(2, 'Shop D05, Central Street', 'b')])
    assert grouped[0]['group']['first'] and grouped[0]['group']['size'] == 2
    apart = app_module._receiver_groups([con(1, '212 High Street', 'a'), con(2, 'Shop D05, Central Street', 'b')], {'a', 'b'})
    assert all('group' not in c and c['kept_separate'] for c in apart)


def test_generate_reports_its_progress_for_the_loading_overlay(client, files, monkeypatch):
    """The overlay's bar (progress.py): only forward, the labels counted as they're drawn, finished at 100 % with
    the run's own log lines."""
    from engine.core import progress
    seen, real_tick = [], progress.tick

    def spy(done, total, step=None):
        real_tick(done, total, step)
        seen.append((progress.get(progress.current())['percent'], step))
    monkeypatch.setattr(progress, 'tick', spy)
    final, sig, _ = files
    run_page(client, final, sig, extra={"progress_id": "overlay-test-0001"})
    percents = [p for p, _ in seen]
    assert percents == sorted(percents) and percents[0] == 2 and percents[-1] < 90  # reading the tabs, then drawing
    assert seen[-1][1].startswith("Saving the PDF in packing-sheet order")
    assert any(s and s.startswith("Drawing the packing labels: 1 of ") for _, s in seen)
    c, *_ = client
    d = c.get("/api/progress/overlay-test-0001").get_json()
    assert d['finished'] and d['percent'] == 100 and d['counted']
    assert any(x['text'].startswith("Generate done in") for x in d['lines'])
    assert c.get("/api/progress/never-started-01").status_code == 404


def test_a_stopped_generate_shows_what_it_did_and_where_it_stopped(client, files):
    """The page that says Generate stopped lists the run's own log lines (progress.lines_now), the same as the Logs
    page, so what happened is there without opening the log."""
    c, app_module, _ = client
    final, sig, _ = files
    with open(final, "rb") as f, open(sig, "rb") as s:
        c.post("/packing-labels/horizontal", data={"file": (f, final.name), "signature_file": (s, sig.name)},
               content_type="multipart/form-data", follow_redirects=True)
    form = {"filename": app_module.secure_filename(final.name), "all_tabs": ["Spring Sale", "Easter"],
            "selected_tabs": ["Spring Sale", "Easter"], "header_row_Spring Sale": "8", "header_row_Easter": "8"}
    r = c.post("/packing-labels/horizontal", data={**form, "generate": "1", "who_pays": "S", "service_code": "NOPE1",
                                                   "reference": "J476699", "project_name": "Stopped run",
                                                   "courier_csv": ["open360"], "progress_id": "stopped-run-0001",
                                                   "attribute_order": "thumbnail,desc,dimension,job_no,barcode,qty"},
               follow_redirects=True)
    page = r.data.decode()
    assert "Unknown service code" in page and "What the app did before it stopped" in page
    assert "Stopped in the app at:</strong> app.py, line" in page and "in _distribution_page" in page
    assert "Generate FAILED after" in page and "Technical detail (the full trace)" in page


def test_stitch_page_picks_the_courier_csv_and_keeps_an_attached_one_beside_it(client, files):
    """The Stitch page lists the job's courier CSVs (the changed one marked) to pick the one with the current
    services; an attached CSV is kept beside the project's (nothing replaced) and remembered for Finalise; another
    job's CSV is refused."""
    import io as io_module
    import time
    c, app_module, projects = client
    final, sig, _ = files
    folder = run_page(client, final, sig)
    open360 = next(n for n in os.listdir(folder) if n.endswith("Open360.csv"))
    original = (folder / open360).read_bytes()
    later = time.time() + 600
    os.utime(folder / open360, (later, later))
    page = c.get(f"/stitch/{folder.name}").data.decode()
    assert "Which courier CSV has the current services?" in page and "changed after Generate" in page
    before = set(os.listdir(folder))
    edited = original.decode("utf-8-sig").replace("DHLWPXOD", "TNTN").encode("utf-8-sig")
    r = c.post(f"/stitch/{folder.name}", data={"courier_csv": (io_module.BytesIO(edited), open360)},
               content_type="multipart/form-data", follow_redirects=True)
    assert "Choose the courier label PDF" in r.data.decode()  # no labels yet; the CSV is kept all the same
    kept = [n for n in set(os.listdir(folder)) - before if n.endswith(".csv")]
    assert len(kept) == 1 and (folder / open360).read_bytes() == original
    labelmap = json.load(open(folder / next(n for n in os.listdir(folder) if n.endswith(".labelmap.json"))))
    assert labelmap["files"]["service_csv"] == kept[0] and kept[0] in labelmap["files"]["courier_csvs"]
    assert set(app_module.pack_services(str(folder), labelmap, kept[0]).values()) == {"TNTN"}
    other = b"Receiver Name,Service,Item Reference\r\nSomeone,TNTN,99T9P9 Elsewhere\r\n"
    r = c.post(f"/stitch/{folder.name}", data={"courier_csv": (io_module.BytesIO(other), "other.csv")},
               content_type="multipart/form-data", follow_redirects=True)
    assert "this job" in r.data.decode() and not (folder / "other.csv").exists()


def test_two_label_runs_in_one_project_are_stitched_apart(client, files):
    """Labels made twice in one Packing Sheets project (e.g. the Final file, then a Stage file): the first run's
    stitching (a session from before runs were kept apart, its Complete Labels) stays the first run's; the Stitch page
    lists both and stitches the one picked; the CSV buttons follow the newest run's own CSVs."""
    c, app_module, projects = client
    final, sig, _ = files
    folder = _sheets_project(projects, final, sig)
    _, gen = _into_project_form(c, app_module, folder, final)
    assert "Labels Successfully Generated" in c.post("/packing-labels/horizontal", data=gen, follow_redirects=True).data.decode()
    first = app_module._label_maps(str(folder))[0]
    (folder / "stitch_session.json").write_text('{"marker": "first run"}', encoding="utf-8")  # as an older stitch left it
    (folder / "Complete Labels 1 - first.pdf").write_bytes(b"%PDF-1.4")
    lm = json.load(open(folder / first, encoding="utf-8"))
    lm["files"]["complete_labels"] = ["Complete Labels 1 - first.pdf"]
    json.dump(lm, open(folder / first, "w", encoding="utf-8"))
    import time
    time.sleep(1.1)  # the second run's files are newer
    _, gen = _into_project_form(c, app_module, folder, final)
    assert "Labels Successfully Generated" in c.post("/packing-labels/horizontal",
                                                     data={**gen, "courier_csv": ["openfreight"]}, follow_redirects=True).data.decode()
    maps = app_module._label_maps(str(folder))
    assert len(maps) == 2 and maps[1] == first and app_module._current_run(str(folder)) == maps[0]
    earlier = folder / f"stitch_session - {first[:-len('.labelmap.json')]}.json"
    assert earlier.exists() and not (folder / "stitch_session.json").exists()  # the first run's, by name
    page = c.get(f"/stitch/{folder.name}").data.decode()
    assert "This project has labels from 2 runs" in page and "stitching this one" in page
    c.get(f"/stitch/{folder.name}?run={first}")
    assert app_module._current_run(str(folder)) == first and app_module._run_file(str(folder), "session") == str(earlier)
    assert (folder / "Complete Labels 1 - first.pdf").exists()  # another run's Complete Labels: never removed
    kinds, _ = app_module._courier_csvs(str(folder), os.listdir(folder))
    assert list(kinds) == ["openfreight"]  # the newest run made only OpenFreight: Open360 is offered for it



def test_an_unexpected_failure_behind_the_overlay_lands_on_the_logs_page(client, files, monkeypatch):
    """Generate (behind the loading overlay) failing for a reason that isn't a check: the browser goes to the Logs
    page at that failure, with a way back to the page as it was. A check (an unknown service code) stays on the page."""
    c, app_module, _ = client
    final, sig, _ = files

    def broken(*a, **k):
        raise RuntimeError("the PDF couldn't be written")
    monkeypatch.setattr(app_module, "generate_horizontal_labels", broken)
    with pytest.raises(AssertionError):  # run_page expects success
        run_page(client, final, sig, extra={"progress_id": "failing-run-0001"})
    form = {"filename": app_module.secure_filename(final.name), "all_tabs": ["Spring Sale", "Easter"],
            "selected_tabs": ["Spring Sale", "Easter"], "header_row_Spring Sale": "8", "header_row_Easter": "8",
            "generate": "1", "who_pays": "S", "service_code": app_module.SERVICE_STORE.services()[0]['code'],
            "reference": "J476699", "project_name": "Broken run", "courier_csv": ["open360"], "progress_id": "failing-run-0002",
            "attribute_order": "thumbnail,desc,dimension,job_no,barcode,qty"}
    r = c.post("/packing-labels/horizontal", data=form)
    assert r.status_code == 303 and r.headers["Location"].startswith("/logs?at=")
    from urllib.parse import parse_qs, urlsplit
    q = parse_qs(urlsplit(r.headers["Location"]).query)
    assert q["find"][0].startswith("Generate FAILED") and q["back"][0].startswith("/packing-labels/horizontal?view=")
    back = c.get(q["back"][0]).data.decode()
    assert "the PDF couldn" in back and "What the app did before it stopped" in back  # the page as it was, to try again
    page = c.post("/packing-labels/horizontal", data={**form, "service_code": "NOPE1"}, follow_redirects=True).data.decode()
    assert "Unknown service code" in page  # a check: fixed on the page, no trip to the Logs page


def test_a_script_task_behind_the_overlay_gets_the_logs_address_others_the_usual_error(client, monkeypatch):
    c, app_module, _ = client

    def broken(name):
        raise RuntimeError("label map unreadable")
    monkeypatch.setattr(app_module, "_stitch_project", broken)
    r = c.post("/stitch/any/finalise", json={}, headers={"X-Progress-Id": "script-run-0001"})
    assert r.status_code == 500 and r.headers["X-Log-Url"].startswith("/logs?at=") and r.get_json()["log_url"]
    r = c.post("/stitch/any/finalise", json={})  # not behind the overlay: the usual error, nothing to follow
    assert r.status_code == 500 and "X-Log-Url" not in r.headers


def test_the_dashboard_shows_the_job_number_itself(client, tmp_path):
    """J477130, never J477130-03 or J477130-001: from the Packing Sheet's name, the folder, the CSVs."""
    from engine.core.core_math import job_series
    for text, job in [("J476523-01", "J476523"), ("J476523-001", "J476523"), ("J476523", "J476523"),
                      ("Final J476523 Spring", "J476523"), ("Spring_Job-J476523-01_261017_1000", "J476523"),
                      ("j476523-01", "J476523"), ("UNKNOWN", "UNKNOWN"), ("", "")]:
        assert job_series(text) == job, text
    c, app_module, projects = client
    folder = projects / "Spring_Job-J477130-03_261018_1000"
    folder.mkdir()
    (folder / "Packing Sheet_J477130-03.xlsx").write_bytes(b"")
    assert app_module._project_job(str(folder), os.listdir(folder), "Spring") == "J477130"


def test_a_new_label_shuffler_project_needs_its_job_number(client, files):
    """Creating a project on the Label Shuffler asks for its job number (the dashboard shows it, J476523 kept as the
    number itself); without one the page says so and keeps what was typed."""
    from engine.core import project_record
    c, app_module, projects = client
    _, sig, _ = files
    with open(sig, "rb") as f:
        r = c.post("/pdf", data={"step": "1", "new_project": "Shuffle test", "excel_file": (f, sig.name)},
                   content_type="multipart/form-data", follow_redirects=True)
    page = r.data.decode()
    assert "Type the job number of this project" in page and 'value="Shuffle test"' in page
    assert not any(p.is_dir() for p in projects.iterdir())  # nothing made yet
    with open(sig, "rb") as f:
        c.post("/pdf", data={"step": "1", "new_project": "Shuffle test", "job_number": " j476523-01 ",
                             "excel_file": (f, sig.name)}, content_type="multipart/form-data", follow_redirects=True)
    folder = next(p for p in projects.iterdir() if p.is_dir())
    assert project_record.job(str(folder)) == "J476523"
    assert app_module._project_job(str(folder), os.listdir(folder), "Shuffle test") == "J476523"


def test_a_stitch_puts_logos_only_beside_courier_labels(client, files):
    """A packing label left without a courier label (Macpac Kent St in a real run) gets no courier logo: the logo is
    on exactly the pages a courier label went on."""
    import io as io_module
    from PIL import Image
    c, app_module, projects = client
    final, sig, _ = files
    buf = io_module.BytesIO()
    Image.new('RGBA', (777, 123), (200, 0, 0, 255)).save(buf, 'PNG')  # a size nothing else on the labels has
    code = app_module.SERVICE_STORE.services()[0]['code']
    app_module.SERVICE_STORE.set_logo(app_module.SERVICE_STORE.group_of(code)['id'], buf.getvalue())
    folder = run_page(client, final, sig)
    lm = json.load(open(folder / next(n for n in os.listdir(folder) if n.endswith(".labelmap.json")), encoding="utf-8"))
    packs = [pk for pks in lm["stores"].values() for pk in pks]
    courier = fitz.open()
    for pk in packs[1:]:  # every pack but the first gets its courier label
        for _ in range(int(pk.get("boxes") or 1)):
            courier.new_page(width=288, height=432).insert_text((20, 40), f"Item Ref: {pk['open360_item_reference']}")
    data = courier.tobytes()
    (folder / "courier.pdf").write_bytes(data)
    out, result, _ = app_module._finalise(str(folder), app_module._stitch_project(folder.name)[2], [("courier.pdf", data)],
                                          add_unmatched=False)
    shift = result["report_pages"] - (lm.get("project") or {}).get("report_pages", 0)
    with_logo = {n + 1 for n, page in enumerate(fitz.open(folder / out)) if any(img[2] == 777 for img in page.get_images())}
    assert with_logo == {p["packing_page"] + shift for p in result["placed"]}
    assert not with_logo & {p + shift for p in packs[0]["box_pages"]}


def test_a_horizontal_code_on_another_stores_label_is_not_placed_by_it(client, files):
    """Horizontal Item References carry this Generate's own code ('03T1P2QWXK Chadstone'), as vertical ones do: the
    serial alone repeats in every job, the code doesn't. A courier label of another job (same serial, another code)
    isn't placed; this job's labels are, the store written in full, cut short or not at all. A project made before
    codes (bare '03T1P2') still checks the store written on the label."""
    import copy as copy_module
    c, app_module, projects = client
    final, sig, _ = files
    folder = run_page(client, final, sig)
    lm = json.load(open(folder / next(n for n in os.listdir(folder) if n.endswith(".labelmap.json")), encoding="utf-8"))
    packs = [(store, pk) for store, pks in lm["stores"].items() for pk in pks if int(pk.get("boxes") or 1) == 1]
    (s1, right), (s2, cut), (s3, bare), (s4, other) = packs[:4]
    address = "\n1 SAMPLE ST\nCORIO VIC 3214"  # the label's text runs on after the Item Ref: not part of the store
    another_job = serial_of(other["open360_code"]) + ("ZZZZ" if not other["open360_code"].endswith("ZZZZ") else "YYYY")

    def plan(label_map, texts):
        courier = fitz.open()
        for text in texts:
            courier.new_page(width=288, height=432).insert_text((20, 40), text + address)
        return app_module.plan_stitch(str(folder), label_map, [("c.pdf", courier.tobytes())])
    got = plan(lm, [f"Item Ref: {right['open360_code']} {s1}", f"Item Ref: {cut['open360_code']} {s2[:6]}",
                    f"Item Ref: {bare['open360_code']}", f"Item Ref: {another_job} {s4}"])
    placed = {p["pack_reference"] for p in got["placed"]}
    assert {right["item_reference"], cut["item_reference"], bare["item_reference"]} <= placed
    assert other["item_reference"] not in placed and len(got["unmatched"]) == 1  # another job's: not placed
    # A project generated before the codes: its bare serials, and the store on the label checked
    old = copy_module.deepcopy(lm)
    for pks in old["stores"].values():
        for pk in pks:
            pk["open360_code"] = serial_of(pk["open360_code"])
    got = plan(old, [f"Item Ref: {serial_of(right['open360_code'])} {s1}",
                     f"Item Ref: {serial_of(other['open360_code'])} Somewhere Else Entirely"])
    assert {p["pack_reference"] for p in got["placed"]} == {right["item_reference"]}
    (u,) = got["unmatched"]
    assert f"is {s4}'s in this job, but the label is for Somewhere Else Entirely" in u["reason"]


def test_a_background_stitch_stays_on_its_run_while_another_is_picked(client, files, monkeypatch):
    """Two runs in one project; run A stitches in the background (the matching page's Finalise). Picking run B
    meanwhile is refused, and even if B became the current run (another tab, a race), A's stitch writes only A's label
    map and session: B's label map is left as it was."""
    import threading
    import time
    c, app_module, projects = client
    final, sig, _ = files
    folder = _sheets_project(projects, final, sig)
    for csv_kind in (["open360"], ["openfreight"]):
        _, gen = _into_project_form(c, app_module, folder, final)
        c.post("/packing-labels/horizontal", data={**gen, "courier_csv": csv_kind}, follow_redirects=True)
        time.sleep(1.1)
    run_b, run_a = app_module._label_maps(str(folder))  # newest first
    c.get(f"/stitch/{folder.name}?run={run_a}")
    found = app_module._stitch_project(folder.name)
    assert found[1] == run_a
    b_before = (folder / run_b).read_bytes()
    lm = found[2]
    pk = next(p for pks in lm["stores"].values() for p in pks)
    courier = fitz.open()
    courier.new_page(width=288, height=432).insert_text((20, 40), f"Item Ref: {pk['open360_item_reference']}")
    data = courier.tobytes()
    release, rendering = threading.Event(), threading.Event()
    real = app_module.render_stitch

    def slow(*a, **k):
        rendering.set()
        release.wait(10)
        return real(*a, **k)
    monkeypatch.setattr(app_module, "render_stitch", slow)
    app_module._start_finalise(str(folder), lm, [("c.pdf", data)], True)
    assert rendering.wait(10)
    page = c.get(f"/stitch/{folder.name}?run={run_b}", follow_redirects=True).data.decode()
    assert "Still stitching the other labels of this project" in page
    assert app_module._current_run(str(folder)) == run_a  # the switch was refused
    app_module._write_json(str(folder / app_module.CURRENT_RUN_FILE), {"label_map": run_b})  # as if it raced through
    release.set()
    assert app_module.STITCH_DONE[str(folder)].wait(30)
    assert app_module.STITCH_JOBS[str(folder)]["state"] == "done"
    assert (folder / run_b).read_bytes() == b_before  # B untouched
    a_map = json.load(open(folder / run_a, encoding="utf-8"))
    assert a_map["files"]["complete_labels"] == [app_module.STITCH_JOBS[str(folder)]["output"]]
    stem = run_a[:-len(".labelmap.json")]
    assert json.load(open(folder / f"stitch_session - {stem}.json", encoding="utf-8")).get("finalised")


def test_sub_group_reads_packs_whose_names_have_quotes(client, monkeypatch):
    """Tab and pack names come from the client's file ('12" Poster'): the page names its fields with them as they are
    (built as elements, not HTML), and Sub-Group reads every range back; none is dropped."""
    c, app_module, projects = client
    folder = projects / "Quotes_Job-J1_261018_2200"
    folder.mkdir()
    tab, pack = 'Spring "Sale"', '12" Poster'
    json.dump({"job_id": "J1", "file1_name": "Final x.xlsx", "tabs": {tab: {"packs": {pack: {"is_selected": True}}}}},
              open(folder / "Final x.json", "w"))
    got = {}
    monkeypatch.setattr(app_module, "execute_subgroups", lambda d, meta, instructions: got.update(instructions))
    r = c.post(f"/subgroup/{folder.name}", data={
        "target_json": "Final x.json", "selected_tabs[]": [tab], f"item_row_{tab}": "9", f"target_pack_{tab}[]": [pack],
        f"start_item_{tab}_{pack}[]": ["1", "4"], f"end_item_{tab}_{pack}[]": ["3", "6"]})
    assert r.status_code == 302
    assert got == {tab: {"item_row": 9, "packs": {pack: [[1, 3], [4, 6]]}}}
