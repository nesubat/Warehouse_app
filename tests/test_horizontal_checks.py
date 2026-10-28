"""Horizontal files: every unit in the grid lands on exactly one label, or Generate is refused with the Excel row and
column of what's wrong. Each case changes one thing in the sample files (tests/fixtures.py) that once lost units
without a word.

    pytest -q tests/test_horizontal_checks.py
"""
import os
import re
import sys

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fixtures import build_final_file, build_signature_links, JOB_ROW, PACK_ROW  # noqa: E402
from engine.packing_labels.horizontal_distribution import read_horizontal, read_signature_file  # noqa: E402
from engine.packing_labels.packing_label_generator import PackCheckError  # noqa: E402

FINAL = "Final J476699 Spring rollout.xlsx"
SIG = "Signature links_J476699.xlsx"
TAB = "Spring Sale"
SPRING_UNITS = 187  # every quantity in the sample's Spring Sale grid


@pytest.fixture()
def files(tmp_path):
    codes = build_final_file(tmp_path / FINAL)
    build_signature_links(tmp_path / SIG, codes)
    return tmp_path / FINAL, tmp_path / SIG


def read(final, sig=None, tab=TAB, tab_no=1):
    return read_horizontal(str(final), PACK_ROW, tab, None, tab_no, read_signature_file(str(sig)) if sig else None)


def units(groups):
    return sum(i["units"] for g in groups.values() for i in g["items"] if not i.get("is_code"))


def refused(final, sig=None):
    with pytest.raises(PackCheckError) as e:
        read(final, sig)
    return e.value.issues


def edit(path, change, tab=TAB):
    """Opens a sample file, lets change(ws) edit the tab, saves it."""
    wb = openpyxl.load_workbook(path)
    change(wb[tab])
    wb.save(path)


def job_cols(ws):
    """The item columns, found by their job numbers (never fixed letters)."""
    return [c for c in range(1, ws.max_column + 1) if re.match(r"J\d", str(ws.cell(row=JOB_ROW, column=c).value or ""))]


def store_row(ws, name):
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == name)


def test_unchanged_sample_reads_every_unit_with_no_problems(files):
    final, sig = files
    groups, _, _, warnings = read(final, sig)
    assert units(groups) == SPRING_UNITS and not warnings
    alone, *_ = read(final)
    assert units(alone) == SPRING_UNITS
    easter, *_ = read(final, sig, "Easter", 2)
    assert units(easter) == 10


def test_two_packs_with_one_name_refused(files):
    """Pack 2's header renamed '1': both took the Label Plugin's pack 1 and one overwrote the other's labels."""
    final, sig = files
    cells = {}

    def rename(ws):
        # The pack headers are the cells merged over each pack's items (the Packing Spec columns are headed '1', '2' too)
        heads = sorted(m.min_col for m in ws.merged_cells.ranges if m.min_row == PACK_ROW)
        cells["at"] = [f"{get_column_letter(c)}{PACK_ROW}" for c in heads]
        ws.cell(row=PACK_ROW, column=heads[1], value="1")
    edit(final, rename)
    issues = refused(final, sig)
    issue = next(i for i in issues if "packs named '1'" in i["text"])
    assert issue["col"] == ", ".join(cells["at"])


def test_two_packs_given_one_pack_order_refused(files):
    """The Label Plugin gives packs '1' and '2' the same Pack Order: their serials would be the same."""
    final, sig = files
    wb = openpyxl.load_workbook(sig)
    ws = wb["Label Plugin"]
    head = [str(c.value or "").strip() for c in ws[1]]
    tab, pack, order = head.index("Tab"), head.index("Pack"), head.index("Pack Order")
    for row in ws.iter_rows(min_row=2):
        if row[tab].value == TAB and str(row[pack].value) == "2":
            row[order].value = 1
    wb.save(sig)
    issues = refused(final, sig)
    issue = next(i for i in issues if "are all pack 1" in i["text"])
    assert re.fullmatch(rf"[A-Z]+{PACK_ROW}, [A-Z]+{PACK_ROW}", issue["col"])


def test_quantities_on_a_row_with_no_store_name_refused(files):
    final, sig = files
    at = {}

    def clear(ws):
        at["row"] = store_row(ws, "Highpoint")
        at["cols"] = [get_column_letter(c) for c in job_cols(ws) if ws.cell(row=at["row"], column=c).value]
        ws.cell(row=at["row"], column=1).value = None  # (ws.cell(..., value=None) leaves the cell as it was)
    edit(final, clear)
    for s in (sig, None):
        issue = next(i for i in refused(final, s) if i["text"] == "Quantities but no store name")
        assert issue["rows"] == str(at["row"]) and issue["col"] == ", ".join(at["cols"])
        assert issue["detail"].startswith("25 units in ")


def test_item_column_with_no_job_number_refused(files):
    final, sig = files
    at = {}

    def clear(ws):
        col = job_cols(ws)[1]  # J476699-02, the A2 poster
        at["cell"] = f"{get_column_letter(col)}{JOB_ROW}"
        ws.cell(row=JOB_ROW, column=col).value = None
    edit(final, clear)
    for s in (sig, None):
        issue = next(i for i in refused(final, s) if i["text"].startswith("No job number"))
        assert issue["text"] == f"No job number in {at['cell']}: 16 units in this column"
        assert issue["col"] == at["cell"] and issue["rows"] == "9-14, 16"  # Garden City (15) has none


def test_store_below_the_label_plugins_last_row_refused(files):
    """A store added under the last one after Packing Sheets ran: the Label Plugin still says the stores end above it."""
    final, sig = files
    at = {}

    def add(ws):
        total = store_row(ws, "TOTAL")
        at["row"] = total
        ws.cell(row=total + 1, column=1, value="TOTAL")
        for c in range(1, 10):  # the store's name, address and Packing Specs, copied from the store above
            ws.cell(row=total, column=c, value=ws.cell(row=total - 1, column=c).value)
        ws.cell(row=total, column=1, value="Doncaster")
        for c, q in zip(job_cols(ws), (5, 5, 5, 5, 5, 0, 0, 0)):
            ws.cell(row=total, column=c, value=q or None)
    edit(final, add)
    issue = next(i for i in refused(final, sig) if i["text"].startswith("Stores below the Label Plugin's last store row"))
    assert issue["rows"] == str(at["row"]) and "25 units" in issue["detail"]
    assert issue["stores"] == [[at["row"], "Doncaster"]]


@pytest.mark.parametrize("value, shown", [(-2, "-2"), (1.5, "1.5")])
def test_negative_or_part_quantities_go_on_the_label_as_written(files, value, shown):
    """Nothing in the grid but blank, '-' or 0 is left off a label: -2 and 1.5 are printed as the file has them."""
    final, sig = files
    at = {}

    def put(ws):
        col = job_cols(ws)[0]  # J476699-01: Corio Village has 2
        at["row"] = store_row(ws, "Corio Village")
        ws.cell(row=at["row"], column=col, value=value)
    edit(final, put)
    groups, *_ = read(final, sig)
    group = next(g for g in groups.values() if g["excel_rows"] == [at["row"]] and g["pack_no"] == 1)
    assert next(i["qty"] for i in group["items"] if i["job_no"] == "J476699-01") == shown
    assert units(groups) == SPRING_UNITS - 2 + value


def test_a_quantity_with_a_note_counts_its_units_and_prints_the_note(files, tmp_path):
    """'6 (2 of each kind)': 6 units on the label, the note under its quantity box (it was refused as not a number,
    and the store's pack reported as having no quantities)."""
    import pymupdf as fitz
    from engine.packing_labels.packing_label_generator import generate_packing_labels
    final, sig = files
    at = {}

    def put(ws):
        col = job_cols(ws)[5]  # J476699-05, pack 2's counter card: Chadstone has 6
        at["row"] = store_row(ws, "Chadstone")
        ws.cell(row=at["row"], column=col, value="6 (2 of each kind)")
    edit(final, put)
    groups, *_ = read(final, sig)
    assert units(groups) == SPRING_UNITS
    key, group = next((k, g) for k, g in groups.items() if g["excel_rows"] == [at["row"]] and g["pack_no"] == 2)
    card = next(i for i in group["items"] if i["job_no"] == "J476699-05")
    assert card["qty"] == "6" and card["qty_note"] == "2 of each kind"
    pdf = tmp_path / "labels.pdf"
    generate_packing_labels({key: group}, str(pdf), ['job_no', 'qty'])
    assert "(2 of each kind)" in fitz.open(pdf)[0].get_text()


def test_words_in_the_grid_go_on_the_label_as_written_and_dash_or_zero_is_empty(files, tmp_path):
    """'TBC' gets its item box with 'TBC' in the quantity box (no units counted); '-' and '0' are no item, as the Label
    Shuffler and Packing Sheets take them."""
    import pymupdf as fitz
    from engine.packing_labels.packing_label_generator import generate_packing_labels
    final, sig = files
    at = {}

    def put(ws):
        cols = job_cols(ws)[5:]  # pack 2: Knox City has 3, 5, 0
        at["row"] = store_row(ws, "Knox City")
        ws.cell(row=at["row"], column=cols[0], value="TBC")
        ws.cell(row=at["row"], column=cols[1], value="-")
        ws.cell(row=at["row"], column=cols[2], value="0")
    edit(final, put)
    groups, *_ = read(final, sig)
    assert units(groups) == SPRING_UNITS - 8  # Knox City's 3 and 5 are gone; TBC counts none
    key, group = next((k, g) for k, g in groups.items() if g["excel_rows"] == [at["row"]] and g["pack_no"] == 2)
    assert [(i["job_no"], i["qty"]) for i in group["items"]] == [("J476699-05", "TBC")]
    pdf = tmp_path / "labels.pdf"
    generate_packing_labels({key: group}, str(pdf), ['job_no', 'qty'])
    assert "TBC" in fitz.open(pdf)[0].get_text()


def test_signature_links_store_code_tabs_skip_the_tabs_that_arent_stores(files):
    from engine.core.core_math import NOT_STORE_TABS
    _, sig = files
    tabs = set(read_signature_file(str(sig))["codes"])
    assert tabs == {"Spring Sale", "Easter"} and not tabs & set(NOT_STORE_TABS)


def test_a_size_split_over_rows_prints_with_its_letters(files):
    """Width and Length in rows of their own (a misspelt label too): each item's Dimensions line on the label is
    'W 600mm × L 900mm', as vertical files show their Width / Height columns; a size already in its own words stays."""
    final, sig = files
    label_col = {}

    def split(ws):
        label_col["c"] = next(c for c in range(1, ws.max_column + 1) if ws.cell(row=4, column=c).value == "Size")
        ws.cell(row=4, column=label_col["c"]).value = "Width (mm)"
        ws.cell(row=5, column=label_col["c"]).value = "Lenght (mm)"
        for n, c in enumerate(job_cols(ws)):
            ws.cell(row=4, column=c).value = 600 + n
            ws.cell(row=5, column=c).value = "90cm" if n == 0 else 900 + n
    edit(final, split)
    groups, *_ = read(final, sig)
    dims = {i["dimension"] for g in groups.values() for i in g["items"] if not i.get("is_code")}
    assert "W 600mm × L 90cm" in dims and "W 601mm × L 901mm" in dims
    assert not any(k.startswith("row_") and i.get(k) for g in groups.values() for i in g["items"] for k in i)  # not extra blocks


def test_dimension_labels_are_read_by_their_words():
    from engine.packing_labels.packing_label_generator import dim_part, dim_unit, join_dimensions, _header_field
    for label, part in [("Width", "W"), ("Size W (mm)", "W"), ("Bredth", "B"), ("Item Length", "L"), ("L", "L"),
                        ("Height (cm)", "H"), ("Depth", "D"), ("Long", "L")]:
        assert dim_part(label) == part, label
    for label in ("Description", "Long sleeve tee", "Dimensions", "Size", "Wall art", "Material", "Brand"):
        assert dim_part(label) is None, label
    assert dim_unit("Width (mm)") == "mm" and dim_unit("Width") == ""
    assert join_dimensions({"L": ("900", "mm"), "W": ("600", "mm"), "B": ("", "mm")}) == "W 600mm × L 900mm"
    assert _header_field("Breadth (mm)") == "dim_b" and _header_field("Width") == "dim_w"  # vertical columns too


def test_two_stores_with_one_name_keep_their_own_signature_links_codes(files):
    """Two stores of the same name (another branch, written the same way) with different codes: each takes its own
    code from Signature Links (the first of that name the first code, and so on), not the last row's code for both."""
    final, sig = files
    book = openpyxl.load_workbook(sig)
    ws = book[TAB]
    head = [c.value for c in ws[1]]
    col = next(i for i, h in enumerate(head, 1) if h and str(h).startswith("Code for"))
    rows = [(r, ws.cell(row=r, column=1).value, ws.cell(row=r, column=col).value) for r in range(2, ws.max_row + 1)]
    (r1, name1, code1), (r2, name2, code2) = next(
        ((a, b) for a in rows for b in rows if a[0] < b[0] and a[2] and b[2] and a[2] != b[2]))
    ws.cell(row=r2, column=1).value = name1  # the second store now has the first one's name
    book.save(sig)

    def rename(sheet):
        row = store_row(sheet, name2)
        sheet.cell(row=row, column=1).value = name1
    edit(final, rename)
    pack = str(head[col - 1])[len("Code for "):].strip()
    groups, *_ = read(final, sig)
    codes = sorted(g["pack_code"] for g in groups.values() if g["store_name"] == name1 and pack.endswith(str(g["pack_name"])))
    assert codes == sorted([code1, code2])  # each its own code, not one code twice
