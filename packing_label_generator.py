import re
from collections import Counter
import pandas as pd
import openpyxl
import io
import math
import zipfile
import xml.etree.ElementTree as ET
import pymupdf as fitz
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
from pdf_engine import code128_modules, _draw_code128, BARCODE_QUIET_MODULES, BARCODE_MAX_MODULE

def extract_rich_value_images(excel_path, sheet_name):
    ns = {
        'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
        'xlrd': 'http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel'
    }
    image_map = {}
    try:
        with zipfile.ZipFile(excel_path, 'r') as z:
            wb_xml = z.read('xl/workbook.xml')
            wb_root = ET.fromstring(wb_xml)
            sheet_id = None
            for sheet in wb_root.findall('.//main:sheet', ns):
                if sheet.get('name') == sheet_name:
                    sheet_id = sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                    break
            if not sheet_id: return {}
            
            wb_rels_xml = z.read('xl/_rels/workbook.xml.rels')
            wb_rels_root = ET.fromstring(wb_rels_xml)
            sheet_path = None
            for rel in wb_rels_root.findall('.//{http://schemas.openxmlformats.org/package/2006/relationships}Relationship'):
                if rel.get('Id') == sheet_id:
                    sheet_path = 'xl/' + rel.get('Target')
                    break
            if not sheet_path: return {}
            
            sheet_xml = z.read(sheet_path)
            sheet_root = ET.fromstring(sheet_xml)
            cell_vm_map = {} 
            for row in sheet_root.findall('.//main:row', ns):
                r_idx = int(row.get('r'))
                for c in row.findall('main:c', ns):
                    vm = c.get('vm')
                    if vm:
                        ref = c.get('r')
                        col_letter = "".join(filter(str.isalpha, ref))
                        cell_vm_map[(r_idx, col_letter)] = vm
            if not cell_vm_map: return {}
                
            meta_xml = z.read('xl/metadata.xml')
            meta_root = ET.fromstring(meta_xml)
            vm_nodes = meta_root.find('.//main:valueMetadata', ns)
            vm_to_rvb = {}
            if vm_nodes is not None:
                for i, bk in enumerate(vm_nodes.findall('main:bk', ns)):
                    rc = bk.find('main:rc', ns)
                    if rc is not None: vm_to_rvb[str(i+1)] = rc.get('v')
                        
            rv_xml = z.read('xl/richData/richValueRel.xml')
            rv_root = ET.fromstring(rv_xml)
            rvb_to_rid = {}
            for i, rel in enumerate(rv_root.findall('.//xlrd:rel', ns)):
                rvb_to_rid[str(i)] = rel.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                
            rv_rels_xml = z.read('xl/richData/_rels/richValueRel.xml.rels')
            rv_rels_root = ET.fromstring(rv_rels_xml)
            rid_to_path = {}
            for rel in rv_rels_root.findall('.//{http://schemas.openxmlformats.org/package/2006/relationships}Relationship'):
                rid_to_path[rel.get('Id')] = rel.get('Target').replace('../', 'xl/')
                
            for (r, col_letter), vm in cell_vm_map.items():
                rvb = vm_to_rvb.get(vm)
                if rvb:
                    rid = rvb_to_rid.get(rvb)
                    if rid:
                        media_path = rid_to_path.get(rid)
                        if media_path:
                            col_idx = column_index_from_string(col_letter)
                            image_map[(r, col_idx)] = io.BytesIO(z.read(media_path))
    except Exception as e:
        pass
    return image_map

EMU_PER_PT = 12700
EMU_PER_PX = 9525


def _row_height_emu(ws, row):
    h = ws.row_dimensions[row].height if row in ws.row_dimensions else None
    return (h or ws.sheet_format.defaultRowHeight or 15) * EMU_PER_PT


def _col_width_emu(ws, col):
    letter = get_column_letter(col)
    w = ws.column_dimensions[letter].width if letter in ws.column_dimensions else None
    w = w or ws.sheet_format.defaultColWidth or 8.43
    return (w * 7 + 5) * EMU_PER_PX


def _cell_at(offset_emu, size_of):
    """1-based row/col index containing a sheet-absolute EMU offset."""
    idx, edge = 1, 0
    while edge + size_of(idx) <= offset_emu and idx < 100000:
        edge += size_of(idx)
        idx += 1
    return idx


def _start_emu(index0, offset, size_of):
    return sum(size_of(i) for i in range(1, index0 + 1)) + offset


def _floating_image_cell(ws, anchor):
    """Cell holding the image's centre, so an image that pokes into a neighbouring row still maps to its own row."""
    row_size = lambda r: _row_height_emu(ws, r)
    col_size = lambda c: _col_width_emu(ws, c)
    top = _start_emu(anchor._from.row, anchor._from.rowOff, row_size)
    left = _start_emu(anchor._from.col, anchor._from.colOff, col_size)
    if getattr(anchor, 'to', None) is not None:
        bottom = _start_emu(anchor.to.row, anchor.to.rowOff, row_size)
        right = _start_emu(anchor.to.col, anchor.to.colOff, col_size)
    elif getattr(anchor, 'ext', None) is not None:
        bottom = top + anchor.ext.height
        right = left + anchor.ext.width
    else:
        return anchor._from.row + 1, anchor._from.col + 1
    return _cell_at((top + bottom) / 2, row_size), _cell_at((left + right) / 2, col_size)


def _row_ranges(rows):
    """[5, 6, 7, 10] -> '5-7, 10'"""
    parts, start, prev = [], None, None
    for r in sorted(rows):
        if start is None:
            start = prev = r
        elif r == prev + 1:
            prev = r
        else:
            parts.append(f"{start}-{prev}" if prev > start else str(start))
            start = prev = r
    if start is not None:
        parts.append(f"{start}-{prev}" if prev > start else str(start))
    return ", ".join(parts)


ADDRESS_FIELDS = ('address_1', 'address_2', 'suburb', 'state', 'postcode', 'country')
INSTALL_YES = {'y', 'yes', 'true'}


def is_install_flag(value):
    return str(value or '').strip().lower() in INSTALL_YES


def normalize_postcode(postcode, country=''):
    """Australian and New Zealand postcodes are always 4 digits. Excel drops the leading zero
    (0803 -> 803, or 803.0), and the courier portal rejects a 3-digit postcode, so pad it back."""
    value = str(postcode if postcode is not None else '').strip()
    if value.endswith('.0') and value[:-2].isdigit():
        value = value[:-2]
    if value.isdigit() and len(value) == 3 and str(country or 'AU').strip().upper() in ('AU', 'NZ', 'AUSTRALIA', 'NEW ZEALAND'):
        value = value.zfill(4)
    return value


def _norm(text):
    """Comparison key: ignores capitals, spacing and punctuation ('12 Bourke Rd.' == '12  BOURKE RD')."""
    return " ".join(re.sub(r"[^0-9a-z]+", " ", str(text or '').lower()).split())


NOT_A_PLACE = ('email', 'e-mail', 'phone', 'mobile', 'contact', 'attn', 'attention', 'fax')
RECEIVER_WORDS = ('store', 'retailer', 'receiver', 'consignee', 'ship to', 'deliver to', 'delivery name')


def _header_field(header):
    """Which column a header is: 'Store Name', 'Retailer', 'Receiver Name' or 'Consignee' are all the store
    (receiver); 'Address Line 1', 'Receiver Address Line 1' or 'Store Address' the street. Address parts
    are checked before the store/receiver words, so 'Receiver Suburb' is the suburb, not the receiver."""
    h = " ".join(str(header).strip().lower().replace('_', ' ').split())
    words = set(re.findall(r'[a-z]+', h))
    place_ok = not any(w in h for w in NOT_A_PLACE)
    if 'packing spec' in h: return 'packing_spec'
    if 'job' in h: return 'job_no'
    if 'qty' in h or 'quantity' in h: return 'qty'
    if 'desc' in h: return 'desc'
    if 'thumb' in h or 'image' in h or 'picture' in h or 'art' in words or 'artwork' in words: return 'thumbnail'
    if 'dimension' in h: return 'dim_combined'
    if 'width' in h or h == 'w': return 'dim_w'
    if 'height' in h or h == 'h': return 'dim_h'
    if place_ok and ('address' in h or 'street' in h):
        return 'address_2' if re.search(r'(line|address|addr)\s*2\b', h) else 'address_1'
    if place_ok and ('suburb' in h or 'town' in words or 'city' in words or 'locality' in words): return 'suburb'
    if place_ok and 'state' in words: return 'state'
    if place_ok and ('postcode' in h or 'post code' in h or 'postal code' in h or 'zip' in words): return 'postcode'
    if place_ok and 'country' in h: return 'country'
    if 'material' in h: return 'material'
    if 'note' in h: return 'notes'
    if 'install' in h: return 'install'
    if place_ok and any(w in h for w in RECEIVER_WORDS): return 'store_name'
    return None


class PackCheckError(ValueError):
    """Blocking problems found in the sheet; carries the structured issues for the preview table."""
    def __init__(self, issues, warnings):
        self.issues, self.warnings = issues, warnings
        super().__init__("\n".join(_issue_line(i) for i in issues))


def _issue(rows, col, text, detail=''):
    return {'rows': rows, 'col': col, 'text': text, 'detail': detail}


def _issue_line(i):
    return " · ".join(x for x in (f"Row {i['rows']}" if i['rows'] else '', i['col'], i['text']) if x)


def _span(rows):
    return _row_ranges(list(range(rows[0]['row'], rows[-1]['row'] + 1)))


def _short(text, limit=40):
    return text if len(text) <= limit else text[:limit - 1].rstrip() + "…"


STREET_TYPES = {
    'street': 'st', 'road': 'rd', 'avenue': 'ave', 'boulevard': 'blvd', 'highway': 'hwy',
    'drive': 'dr', 'parade': 'pde', 'place': 'pl', 'court': 'ct', 'crescent': 'cres',
    'lane': 'ln', 'terrace': 'tce', 'arcade': 'arc',
}


def _address_key(text):
    """Like _norm, plus 'Street' == 'St' etc. Only after a name, so '12 St Kilda Rd' keeps its Saint."""
    words = _norm(text).split()
    return " ".join(
        STREET_TYPES.get(w, w) if i and words[i - 1].isalpha() else w
        for i, w in enumerate(words)
    )


def _distinct(rows, key):
    found = {}
    for r in rows:
        if r[key]:
            found.setdefault(r[key], []).append(r['row'])
    return found


def _variants(rows, key, display):
    """'4-5: 49 Church Street… · 58: 49 Church St…'"""
    by_value = {}
    for r in rows:
        if r[key]:
            by_value.setdefault(r[key], (r[display], []))[1].append(r['row'])
    # Skip the leading words a variant shares with the first one, so the detail shows where they differ
    def shared(a, b):
        n = 0
        while n < min(len(a), len(b)) and _norm(a[n]) == _norm(b[n]):
            n += 1
        return max(0, n - 1)

    variants = [(text.split(), rs) for text, rs in by_value.values()]
    ref = variants[0][0]
    parts = []
    for idx, (words, rs) in enumerate(variants):
        same = max(shared(ref, w) for w, _ in variants[1:]) if idx == 0 else shared(ref, words)
        parts.append(f"{_row_ranges(rs)}: {'…' if same else ''}{_short(' '.join(words[same:]))}")
    return " · ".join(parts)


def _check_pack_consistency(pack_groups, pack_rows, has_address, col):
    errors, warnings = [], []
    no_address_packs = []

    for pack_id, rows in pack_rows.items():
        span = _span(rows)
        pack = pack_groups[pack_id]['pack_spec_name']

        stores = _distinct(rows, 'store_key')
        if len(stores) > 1:
            errors.append(_issue(span, col['store_name'], f"Store name differs in pack {pack}", _variants(rows, 'store_key', 'store')))
        elif stores and any(not r['store_key'] for r in rows):
            blank = [r['row'] for r in rows if not r['store_key']]
            errors.append(_issue(_row_ranges(blank), col['store_name'], f"Store name blank in pack {pack}"))

        if len({r['install'] for r in rows}) > 1:
            flagged = [r['row'] for r in rows if r['install']]
            errors.append(_issue(span, col['install'], f"Install mixed in pack {pack}", f"Y on {_row_ranges(flagged)}"))

        if has_address:
            addresses = _distinct(rows, 'address_key')
            if len(addresses) > 1:
                errors.append(_issue(span, col['address'], f"Address differs in pack {pack}", _variants(rows, 'address_key', 'address')))
            elif addresses and any(not r['address_key'] for r in rows):
                blank = [r['row'] for r in rows if not r['address_key']]
                errors.append(_issue(_row_ranges(blank), col['address'], f"Address blank in pack {pack}"))
            elif not addresses:
                no_address_packs.append(rows[0]['row'])

            if any(r['crosses_pack'] for r in rows):
                warnings.append(_issue(span, col['address'], f"Merged address shared with next pack"))

    if no_address_packs:
        warnings.append(_issue(_row_ranges(no_address_packs), col['address'], "No address"))

    # Across packs of the same store: store deliveries must agree, installer deliveries may differ
    if has_address:
        by_store = {}
        for rows in pack_rows.values():
            first = rows[0]
            if first['store_key'] and first['address_key']:
                by_store.setdefault(first['store_key'], []).append(first)
        for store_rows in by_store.values():
            for install in (False, True):
                group = [r for r in store_rows if r['install'] == install]
                if len(_distinct(group, 'address_key')) > 1:
                    rows_text = _row_ranges([r['row'] for r in group])
                    name = group[0]['store']
                    if install:
                        warnings.append(_issue(rows_text, col['address'], f"{name}: installer addresses differ", _variants(group, 'address_key', 'address')))
                    else:
                        errors.append(_issue(rows_text, col['address'], f"{name}: store address differs between packs", _variants(group, 'address_key', 'address')))

    # Neighbouring packs with the same spec and store usually mean the Packing Spec cells weren't merged.
    # No store column: the address (which then names the receiver) stands in for the store.
    ids = list(pack_rows)
    who = lambda pid: _norm(pack_groups[pid]['store_name']) or pack_rows[pid][0]['address_key']
    for prev_id, next_id in zip(ids, ids[1:]):
        prev, nxt = pack_groups[prev_id], pack_groups[next_id]
        if _norm(prev['pack_spec_name']) == _norm(nxt['pack_spec_name']) and who(prev_id) == who(next_id):
            span = _row_ranges(list(range(pack_rows[prev_id][0]['row'], pack_rows[next_id][-1]['row'] + 1)))
            warnings.append(_issue(span, col['packing_spec'], f"Same spec {nxt['pack_spec_name']} split into separate packs — merge if one box"))

    return errors, warnings


def parse_packing_data(excel_path, header_row, sheet_name=None):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    
    # Every image in a cell is kept: {(row, col): [(reading-order key, image bytes), ...]}
    image_map = {cell: [((0, 0, 0, 0), data)] for cell, data in extract_rich_value_images(excel_path, ws.title).items()}

    for image in getattr(ws, '_images', []):
        try:
            r, c, order = None, None, (0, 0, 0, 0)
            if hasattr(image, 'anchor'):
                if hasattr(image.anchor, '_from'):
                    r, c = _floating_image_cell(ws, image.anchor)
                    f = image.anchor._from
                    order = (f.row, f.rowOff, f.col, f.colOff)  # top-to-bottom, then left-to-right
                elif isinstance(image.anchor, str):
                    c_letter, r_str = coordinate_from_string(image.anchor)
                    c = column_index_from_string(c_letter)
                    r = int(r_str)

            if r and c:
                img_bytes = None
                if hasattr(image, 'ref') and hasattr(image.ref, 'getvalue'):
                    img_bytes = image.ref.getvalue()
                elif hasattr(image, '_data'):
                    img_bytes = image._data() if callable(image._data) else image._data

                if img_bytes:
                    image_map.setdefault((r, c), []).append((order, io.BytesIO(img_bytes)))
        except Exception:
            continue

    cols = {}
    for col_idx in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=col_idx).value
        if not val: continue
        field = _header_field(str(val))
        if not field:
            continue
        # As before, a later matching column wins; but a 'Store'/'Retailer' column beats a 'Receiver'-type one
        is_store = field == 'store_name' and ('store' in str(val).lower() or 'retailer' in str(val).lower())
        if field == 'store_name' and not is_store and cols.get('_store_is_named'):
            continue
        cols[field] = col_idx
        if is_store:
            cols['_store_is_named'] = True
    cols.pop('_store_is_named', None)

    found_headers = {
        'Packing Spec': 'packing_spec' in cols,
        'Store Name': 'store_name' in cols,
        'Job Number': 'job_no' in cols,
        'Quantity': 'qty' in cols,
        'Description': 'desc' in cols,
        'Thumbnail': 'thumbnail' in cols,
        'Dimensions': 'dim_combined' in cols or 'dim_w' in cols or 'dim_h' in cols,
        'Address Info': any(k in cols for k in ADDRESS_FIELDS),
        'Material': 'material' in cols,
        'Notes': 'notes' in cols,
        'Install': 'install' in cols
    }

    if 'packing_spec' not in cols:
        raise ValueError(f"Could not find a 'Packing Spec' column in Row {header_row}.")
    if 'job_no' not in cols:
        raise ValueError(f"Could not find a 'Job Number' column in Row {header_row}.")
    if 'install' not in cols:
        raise ValueError(f"Could not find an 'Install' column in Row {header_row}. Add one: Y for packs going to an installer, N or blank for the store.")

    def get_cell_info(row, col):
        for merged_range in ws.merged_cells.ranges:
            if merged_range.min_col <= col <= merged_range.max_col and merged_range.min_row <= row <= merged_range.max_row:
                val = ws.cell(row=merged_range.min_row, column=merged_range.min_col).value
                return val, merged_range.min_row, merged_range.max_row
        return ws.cell(row=row, column=col).value, row, row

    pack_groups = {}
    pack_rows = {}
    rows_missing_job = []
    rows_no_barcode = []
    rows_without_image = []
    address_keys = [k for k in ADDRESS_FIELDS if k in cols]
    col_ref = {k: f"{get_column_letter(c)} · {str(ws.cell(row=header_row, column=c).value or '').strip().title()}" for k, c in cols.items()}
    addr_letters = [get_column_letter(cols[k]) for k in address_keys]
    col_ref['address'] = f"{addr_letters[0]}–{addr_letters[-1]} · Address" if len(addr_letters) > 1 else (f"{addr_letters[0]} · Address" if addr_letters else "Address")
    col_ref.setdefault('store_name', "Store Name")
    col_ref.setdefault('install', "Install")
    col_ref.setdefault('thumbnail', "Thumbnail")
    current_row = header_row + 1
    last_processed_row = current_row

    while True:
        pack_spec_val, pack_top_row, pack_bottom_row = get_cell_info(current_row, cols['packing_spec'])

        if pack_spec_val is None or str(pack_spec_val).strip() == "": break

        packing_spec = str(pack_spec_val).strip()
        pack_id = f"ROW_{pack_top_row}_{packing_spec}"

        def get_val(key):
            if key in cols:
                val, _, _ = get_cell_info(current_row, cols[key])
                return str(val).strip() if val is not None else ""
            return ""

        store_name = get_val('store_name')
        address_parts = {k: get_val(k) for k in address_keys}
        install = is_install_flag(get_val('install'))

        # Address cells merged past this pack's edges (shared with a neighbouring pack)
        crosses_pack = any(
            lo < pack_top_row or hi > pack_bottom_row
            for _, lo, hi in (get_cell_info(current_row, cols[k]) for k in address_keys)
        )

        if pack_id not in pack_groups:
            pack_groups[pack_id] = {
                'pack_spec_name': packing_spec,
                'store_name': store_name,
                'install': install,
                **{k: address_parts.get(k, '') for k in ADDRESS_FIELDS},
                'excel_rows': [],
                'items': []
            }
            pack_rows[pack_id] = []

        pack_rows[pack_id].append({
            'row': current_row,
            'store': store_name,
            'store_key': _norm(store_name),
            'address': ", ".join(v for v in address_parts.values() if v),
            'address_key': _address_key(" ".join(address_parts.values())),
            'install': install,
            'crosses_pack': crosses_pack,
        })

        dim_w = get_val('dim_w')
        dim_h = get_val('dim_h')
        dim_combined = get_val('dim_combined')

        if dim_combined: dimension = dim_combined
        elif dim_w and dim_h: dimension = f"{dim_w} x {dim_h}"
        elif dim_w: dimension = dim_w
        elif dim_h: dimension = dim_h
        else: dimension = ""

        thumbnails = []
        if 'thumbnail' in cols:
            tc = cols['thumbnail']
            # Images come from the thumbnail cell itself, never a neighbouring cell. A merged
            # thumbnail cell is one cell: every row it covers gets all the images inside it.
            area = next((m for m in ws.merged_cells.ranges
                         if m.min_row <= current_row <= m.max_row and m.min_col <= tc <= m.max_col), None)
            cells = ([(r, c) for r in range(area.min_row, area.max_row + 1) for c in range(area.min_col, area.max_col + 1)]
                     if area else [(current_row, tc)])
            found = sorted((entry for cell in cells for entry in image_map.get(cell, [])), key=lambda e: e[0])
            thumbnails = [data for _, data in found]
            if not thumbnails:
                rows_without_image.append(current_row)

        item = {
            'job_no': get_val('job_no'),
            'qty': get_val('qty') or "1",
            'desc': get_val('desc'),
            'dimension': dimension,
            'material': get_val('material'),
            'install': "INSTALLER" if install else "",
            'notes': get_val('notes'),
            'thumbnails': thumbnails
        }

        if item['job_no'] and code128_modules(item['job_no']) is None:
            rows_no_barcode.append(current_row)
        if not item['job_no']:
            rows_missing_job.append(current_row)
        pack_groups[pack_id]['items'].append(item)
        pack_groups[pack_id]['excel_rows'].append(current_row)
        last_processed_row = current_row
        current_row += 1

    rows_missing_spec = []
    for row in range(current_row, ws.max_row + 1):
        spec_val, _, _ = get_cell_info(row, cols['packing_spec'])
        job_val, _, _ = get_cell_info(row, cols['job_no'])
        if (spec_val is None or str(spec_val).strip() == "") and job_val is not None and str(job_val).strip():
            rows_missing_spec.append(row)

    wb.close()

    errors = []
    if rows_missing_spec:
        errors.append(_issue(_row_ranges(rows_missing_spec), col_ref['packing_spec'], "Packing Spec empty"))
    if rows_missing_job:
        errors.append(_issue(_row_ranges(rows_missing_job), col_ref['job_no'], "Job Number empty"))
    if not pack_groups and not errors:
        errors.append(_issue('', '', f"No data found below row {header_row}"))

    warnings = []
    if rows_without_image:
        warnings.append(_issue(_row_ranges(rows_without_image), col_ref['thumbnail'], "No image"))
    if rows_no_barcode:
        warnings.append(_issue(_row_ranges(rows_no_barcode), col_ref['job_no'],
                               "No barcode: Job Number has a line break or a character a barcode can't hold"))

    pack_errors, pack_warnings = _check_pack_consistency(pack_groups, pack_rows, bool(address_keys), col_ref)
    errors += pack_errors
    warnings += pack_warnings

    if errors:
        raise PackCheckError(errors, warnings)

    return pack_groups, last_processed_row, found_headers, warnings

TEXT_STYLES = {
    'desc': (9, "helv", (0, 0, 0)),
    'dimension': (9, "hebo", (0, 0, 0)),
    'material': (9, "helv", (0, 0, 0)),
    'install': (9, "hebo", (0.8, 0, 0)),
    'notes': (9, "helv", (0, 0, 0)),
}
LINE = 1.2


def _clean(val):
    val = str(val or '').strip()
    return '' if val.lower() == 'none' else val


def _wrap(text, fontname, size, width):
    lines = []
    for paragraph in text.splitlines() or ['']:
        line = ''
        for word in paragraph.split():
            while fitz.get_text_length(word, fontname=fontname, fontsize=size) > width:
                cut = len(word)
                while cut > 1 and fitz.get_text_length(word[:cut], fontname=fontname, fontsize=size) > width:
                    cut -= 1
                if line:
                    lines.append(line)
                    line = ''
                lines.append(word[:cut])
                word = word[cut:]
            candidate = f"{line} {word}".strip()
            if line and fitz.get_text_length(candidate, fontname=fontname, fontsize=size) > width:
                lines.append(line)
                line = word
            else:
                line = candidate
        if line:
            lines.append(line)
    return lines


def _draw_lines(page, lines, x0, x1, y, fontname, size, color):
    for line in lines:
        w = fitz.get_text_length(line, fontname=fontname, fontsize=size)
        page.insert_text((x0 + (x1 - x0 - w) / 2, y + size * 0.92), line,
                         fontname=fontname, fontsize=size, color=color)
        y += size * LINE
    return y


def _fit_rect(slot, data):
    """The largest rect with the image's own proportions that fits inside slot, centred in it."""
    pix = fitz.Pixmap(data.getvalue())
    scale = min(slot.width / pix.width, slot.height / pix.height)
    w, h = pix.width * scale, pix.height * scale
    x0, y0 = slot.x0 + (slot.width - w) / 2, slot.y0 + (slot.height - h) / 2
    return fitz.Rect(x0, y0, x0 + w, y0 + h)


def _layout_cell(item, attribute_order, cell, text_scale, thumb_scale, draw_page=None):
    """Stacks the chosen blocks top-down; returns the total height used. Draws only when draw_page is given."""
    pad = 4
    inner_x0, inner_x1 = cell.x0 + pad, cell.x1 - pad
    inner_w = inner_x1 - inner_x0
    gap = 3 * text_scale
    y = cell.y0 + pad
    for attr in attribute_order:
        if attr == 'thumbnail':
            images = item.get('thumbnails') or []
            if not images:
                continue
            h = cell.height * 0.35 * thumb_scale
            if draw_page:
                # One image: centred, 70% of the box width. Several: share the full inner width,
                # up to 3 side by side, then further rows; each keeps its own proportions.
                n = len(images)
                per_row = min(n, 3)
                rows = math.ceil(n / per_row)
                area_w = cell.width * 0.7 if n == 1 else inner_w
                x0 = cell.x0 + (cell.width - area_w) / 2
                slot_w = (area_w - (per_row - 1) * 2) / per_row
                slot_h = (h - (rows - 1) * 2) / rows
                for i, data in enumerate(images):
                    sx = x0 + (i % per_row) * (slot_w + 2)
                    sy = y + (i // per_row) * (slot_h + 2)
                    rect = fitz.Rect(sx, sy, sx + slot_w, sy + slot_h)
                    try:
                        draw_page.insert_image(_fit_rect(rect, data), stream=data.getvalue())
                    except Exception:
                        draw_page.insert_textbox(rect, "Image Error", fontsize=8, align=1)
            y += h + gap

        elif attr == 'job_no':
            val = _clean(item.get('job_no'))
            if not val:
                continue
            size = 14 * text_scale
            lines = _wrap(val, "hebo", size, inner_w - 4)
            h = len(lines) * size * LINE + 6 * text_scale
            if draw_page:
                _round_rect(draw_page, fitz.Rect(cell.x0 + 3, y, cell.x1 - 3, y + h), 3, fill=(0, 0, 0))
                _draw_lines(draw_page, lines, inner_x0, inner_x1, y + 3 * text_scale, "hebo", size, (1, 1, 1))
            y += h + gap

        elif attr == 'barcode':
            val = item.get('barcode') or _clean(item.get('job_no'))
            modules = code128_modules(val, compact=True) if val else None
            if not modules:
                continue
            # Code 128 needs 10 blank modules either side; the bars get as wide as the box allows
            module_w = min(BARCODE_MAX_MODULE, inner_w / (len(modules) + 2 * BARCODE_QUIET_MODULES))
            bar_h = 22 * text_scale
            size = 6.5 * text_scale
            h = bar_h + 2 + size * LINE
            if draw_page:
                bx = cell.x0 + (cell.width - len(modules) * module_w) / 2
                _draw_code128(draw_page, bx, y, module_w, bar_h, modules)
                _draw_lines(draw_page, [val], inner_x0, inner_x1, y + bar_h + 2, "helv", size, (0, 0, 0))
            y += h + gap

        elif attr == 'qty':
            val = _clean(item.get('qty'))
            if not val:
                continue
            size = 20 * text_scale
            border = 2.5 * text_scale
            lines = _wrap(val, "hebo", size, inner_w - 2 * border - 8)
            widest = max(fitz.get_text_length(l, fontname="hebo", fontsize=size) for l in lines)
            box_w = min(inner_w, max(70 * text_scale, widest + 16 * text_scale))
            h = len(lines) * size * LINE + 6 * text_scale
            if draw_page:
                bx0 = cell.x0 + (cell.width - box_w) / 2
                half = border / 2
                draw_page.draw_rect(fitz.Rect(bx0 + half, y + half, bx0 + box_w - half, y + h - half), color=(0, 0, 0), width=border)
                _draw_lines(draw_page, lines, bx0, bx0 + box_w, y + 3 * text_scale, "hebo", size, (0, 0, 0))
            y += h + gap

        elif attr in TEXT_STYLES:
            val = _clean(item.get(attr))
            if not val:
                continue
            base, fontname, color = TEXT_STYLES[attr]
            size = base * text_scale
            lines = _wrap(val, fontname, size, inner_w)
            if draw_page:
                _draw_lines(draw_page, lines, inner_x0, inner_x1, y, fontname, size, color)
            y += len(lines) * size * LINE + gap
    return y - gap - cell.y0 + pad


def _draw_cell(page, item, attribute_order, cell):
    # Shrink text first (down to 65%), then the thumbnail, until the whole stack fits inside the cell
    text_scale = thumb_scale = 1.0
    for step in range(71):
        s = 1.0 - step * 0.01
        text_scale = max(s, 0.65)
        thumb_scale = 1.0 if s >= 0.65 else s / 0.65
        if _layout_cell(item, attribute_order, cell, text_scale, thumb_scale) <= cell.height:
            break
    _layout_cell(item, attribute_order, cell, text_scale, thumb_scale, draw_page=page)


def _round_rect(page, rect, r, color=None, fill=None, width=0):
    radius = min(0.5, r / max(1, min(rect.width, rect.height)))
    page.draw_rect(rect, color=color, fill=fill, width=width, radius=radius)


def _spec_panel(page, rect, spec, max_size=22):
    """Black rounded panel with the Packing Spec in large white type. Returns its bottom y."""
    inner_w = rect.width - 16
    size = max_size
    lines = _wrap(spec, "hebo", size, inner_w)
    while len(lines) > 2 and size > 11:
        size -= 1
        lines = _wrap(spec, "hebo", size, inner_w)
    lines = lines[:3]
    h = 14 + len(lines) * size * LINE + 6
    panel = fitz.Rect(rect.x0, rect.y0, rect.x1, rect.y0 + h)
    _round_rect(page, panel, 6, fill=(0, 0, 0))
    page.insert_text((panel.x0 + 8, panel.y0 + 11), "PACKING SPEC", fontname="hebo", fontsize=7, color=(1, 1, 1))
    _draw_lines(page, lines, panel.x0 + 8, panel.x0 + 8 + inner_w, panel.y0 + 15, "hebo", size, (1, 1, 1))
    return panel.y1


def _counter_chip(page, x0, y0, text, size=11):
    w = fitz.get_text_length(text, fontname="hebo", fontsize=size) + 16
    h = size * LINE + 8
    chip = fitz.Rect(x0, y0, x0 + w, y0 + h)
    _round_rect(page, chip, h / 2, color=(0, 0, 0), width=1.5)
    page.insert_text((x0 + 8, y0 + 4 + size * 0.92), text, fontname="hebo", fontsize=size)
    return chip


def barcode_text(job, kind):
    """What an item's barcode encodes: 'J477161-26 K3', or just 'J477161-54' for a job number on one row."""
    return f"{job} K{kind}" if kind else job


def assign_barcodes(pack_groups):
    """Gives every item its barcode text. Same idea as the Distribution Mapper's kinds, with rows in
    place of columns: a job number on more than one row gets one kind per row, numbered K1, K2, K3...
    in reading order (first tab to last, then top to bottom), so K1 is its first row in the first tab.
    Call it with all selected tabs combined, in tab order."""
    items = [item for g in pack_groups.values() for item in g['items']]
    key = lambda item: str(item.get('job_no') or '').strip().upper()
    totals = Counter(key(i) for i in items if key(i))
    seen = Counter()
    for item in items:
        job = str(item.get('job_no') or '').strip()
        if not job:
            item['barcode'] = ''
            continue
        seen[key(item)] += 1
        item['barcode'] = barcode_text(job, seen[key(item)] if totals[key(item)] > 1 else None)


def assign_label_numbers(pack_groups):
    """'Label X of Y': each store's pack groups numbered in file order (store names compared loosely)."""
    totals = Counter(_norm(g['store_name']) for g in pack_groups.values())
    seen = Counter()
    for g in pack_groups.values():
        key = _norm(g['store_name'])
        seen[key] += 1
        g['label_no'], g['label_total'] = seen[key], totals[key]


REPORT_COLOURS = {'bad': (0.75, 0.15, 0.13), 'warn': (0.85, 0.55, 0.0), 'info': (0.16, 0.44, 0.64)}


def _latin(text):
    """The built-in PDF fonts only have Latin-1: swap the few other characters the app writes."""
    swaps = {'—': '-', '–': '-', '→': '->', '✓': 'OK', '✏️': '', '’': "'", '‘': "'", '“': '"', '”': '"', '…': '...', '×': 'x'}
    for a, b in swaps.items():
        text = str(text).replace(a, b)
    return text.encode('latin-1', 'replace').decode('latin-1')


def draw_report_pages(doc, report, at=0):
    """Inserts the job report in front of the labels (A4 landscape), as many pages as it needs. Returns the count.

    report: {'title': 'Packing Labels Only', 'job': 'J477161', 'project': 'Lux Test 30', 'when': '05/10/2026 20:25',
             'facts': [(label, value) or (label, value, 'bad'|'warn')...],
             'sections': [{'title', 'level': 'bad'|'warn'|'info', 'items': [text...]}]}
    Compact: the job number in a white-on-black panel with the project beside it, the facts as one row of chips
    (red / amber when they need a look), then each finding as a coloured heading with its items run together."""
    W, H = fitz.paper_size("a4-l")
    M = 30
    pages = []

    def new_page():
        page = doc.new_page(pno=at + len(pages), width=W, height=H)
        pages.append(page)
        page.insert_text((M, M + 10), _latin(f"{report['title'].upper()}  ·  REPORT"), fontname="hebo", fontsize=9,
                         color=(0.35, 0.35, 0.35))
        page.insert_text((W - M - 40, M + 10), _latin(f"Page {len(pages)}"), fontname="helv", fontsize=8, color=(0.5, 0.5, 0.5))
        return page, M + 20

    page, y = new_page()
    # Job number: the thing to check first. The project name beside it needs no caption.
    job = _latin(report.get('job') or '(no job number)')
    size = 30
    job_w = fitz.get_text_length(job, fontname="hebo", fontsize=size) + 30
    panel = fitz.Rect(M, y, M + job_w, y + 48)
    _round_rect(page, panel, 8, fill=(0, 0, 0))
    page.insert_text((M + 8, y + 12), "JOB NUMBER", fontname="hebo", fontsize=6.5, color=(1, 1, 1))
    page.insert_text((M + 15, y + 40), job, fontname="hebo", fontsize=size, color=(1, 1, 1))
    x = panel.x1 + 16
    name = (_wrap(_latin(report.get('project') or ''), "hebo", 20, W - M - x) or [''])[0]
    page.insert_text((x, y + 24), name, fontname="hebo", fontsize=20)
    page.insert_text((x, y + 42), _latin(report.get('when', '')), fontname="helv", fontsize=9, color=(0.35, 0.35, 0.35))
    y = panel.y1 + 10

    # Key facts as chips; the ones that need a look are coloured
    cx = M
    for fact in report.get('facts', []):
        label, value = fact[0], fact[1]
        colour = REPORT_COLOURS.get(fact[2]) if len(fact) > 2 and fact[2] else None
        text = _latin(f"{label}: {value}")
        w = fitz.get_text_length(text, fontname="hebo", fontsize=8.5) + 14
        if cx + w > W - M:
            cx, y = M, y + 20
        _round_rect(page, fitz.Rect(cx, y, cx + w, y + 16), 8, color=colour or (0, 0, 0), width=1.2 if colour else 0.8,
                    fill=colour)
        page.insert_text((cx + 7, y + 11.5), text, fontname="hebo", fontsize=8.5, color=(1, 1, 1) if colour else (0, 0, 0))
        cx += w + 6
    y += 28

    sections = [s for s in report.get('sections', []) if s.get('items')]
    if not sections:
        page.insert_text((M, y + 12), "Nothing needs a look: every label, address and courier detail checked out.",
                         fontname="hebo", fontsize=11, color=(0.12, 0.52, 0.29))
    for section in sections:
        colour = REPORT_COLOURS.get(section.get('level'), REPORT_COLOURS['info'])
        # The items run on as one paragraph, separated by dots, instead of a line each
        body = _wrap(_latin("   ·   ".join(section['items'])), "helv", 8.5, W - 2 * M - 12)
        lines = [(True, _latin(f"{section['title']} ({len(section['items'])})"))] + [(False, t) for t in body]
        top = y
        for bold, text in lines:
            if y + 12 > H - M:
                page.draw_rect(fitz.Rect(M, top, M + 3, y), color=None, fill=colour)
                page, y = new_page()
                top = y
            if bold:
                page.insert_text((M + 9, y + 9.5), text, fontname="hebo", fontsize=9.5, color=colour)
                y += 13
            else:
                page.insert_text((M + 9, y + 9), text, fontname="helv", fontsize=8.5)
                y += 11
        page.draw_rect(fitz.Rect(M, top, M + 3, y), color=None, fill=colour)
        y += 7
    return len(pages)


def generate_packing_labels(pack_groups, output_pdf_path="packing_labels.pdf", attribute_order=None, addresses=None,
                            report=None):
    """Draws the PDF. Returns where each pack landed: {'pages': {pack_key: [0-based page indexes]}, ...}.

    addresses: {pack_key: {'address', 'receiver', 'contact', 'sent'}} as finalised in the consignment preview
    (address book, edits); a plain address string also works. A pack going to an installer gets a red
    "Installer: <receiver> · Attn <contact>" line; any other pack whose
    receiver isn't the store (or has an Attn) gets "Deliver to: …". A pack left out of the courier CSV
    (sent=False) says so in the courier label region. A pack missing from addresses shows its Excel address.
    report: the job report (draw_report_pages) put in front as page 1; the page numbers returned count it."""
    addresses = addresses or {}
    doc = fitz.open()
    MM2PT = 2.83465
    MARGIN = 4 * MM2PT
    GAP = 2.5 * MM2PT
    CELL_RADIUS = 3 * MM2PT
    A4_W, A4_H = fitz.paper_size("a4-l")

    ZONE_A_W = 107 * MM2PT
    ZONE_A_H = 150 * MM2PT
    CHIP_H = 11 * LINE + 8
    ROWS = 3

    # One box size for every page: width from page 1's 4 columns beside the courier region,
    # height from the later pages, which lose the footer strip at the bottom
    FIRST_COLS = 4
    CELL_W = (A4_W - 2 * MARGIN - ZONE_A_W - GAP - (FIRST_COLS - 1) * GAP) / FIRST_COLS
    CELL_H = (A4_H - 2 * MARGIN - CHIP_H - GAP - (ROWS - 1) * GAP) / ROWS
    NEXT_COLS = int((A4_W - 2 * MARGIN + GAP) // (CELL_W + GAP))
    FIRST_PAGE_CELLS, NEXT_PAGE_CELLS = FIRST_COLS * ROWS, NEXT_COLS * ROWS

    if not attribute_order:
        attribute_order = ['thumbnail', 'desc', 'dimension', 'job_no', 'barcode', 'qty']

    assign_label_numbers(pack_groups)
    assign_barcodes(pack_groups)
    page_map = {}

    for pack_key, group_data in pack_groups.items():
        items = group_data['items']
        total_items = len(items)
        store = group_data['store_name']
        spec = group_data.get('pack_spec_name', '')
        label_text = f"LABEL {group_data['label_no']} OF {group_data['label_total']}"
        info = addresses.get(pack_key) or {}
        if isinstance(info, str):
            info = {'address': info}
        installer = bool(group_data.get('install'))
        first_page = len(doc)

        total_pages = 1 if total_items <= FIRST_PAGE_CELLS else 1 + math.ceil((total_items - FIRST_PAGE_CELLS) / NEXT_PAGE_CELLS)
        current_item_idx = 0
        page_num = 1

        while current_item_idx < total_items or (total_items == 0 and page_num == 1):
            page = doc.new_page(width=A4_W, height=A4_H)
            page_text = f"PAGE {page_num} OF {total_pages}"

            if page_num == 1:
                zone_a_rect = fitz.Rect(MARGIN, MARGIN, MARGIN + ZONE_A_W, MARGIN + ZONE_A_H)
                page.draw_rect(zone_a_rect, color=(0.8, 0.8, 0.8), width=0.5, dashes="[3] 0")
                if info.get('sent', True):
                    page.insert_textbox(zone_a_rect, "Courier Label Region\n(107mm x 150mm)", fontsize=10, color=(0.6, 0.6, 0.6), align=1)
                else:
                    # Left out of the courier CSV: no courier label will come for this box
                    note = fitz.Rect(zone_a_rect.x0 + 10, zone_a_rect.y0 + zone_a_rect.height / 2 - 30, zone_a_rect.x1 - 10, zone_a_rect.y1)
                    page.insert_textbox(note, "NOT IN COURIER CSV\nAddress incomplete: no courier label for this box",
                                        fontname="hebo", fontsize=11, color=(0.75, 0.1, 0.1), align=1)

                zone_b = fitz.Rect(MARGIN, zone_a_rect.y1 + GAP, MARGIN + ZONE_A_W, A4_H - MARGIN)
                y = _spec_panel(page, zone_b, spec) + 8

                if store and store.lower() != 'none':
                    for line in _wrap(store, "hebo", 13, zone_b.width)[:2]:
                        page.insert_text((zone_b.x0, y + 12), line, fontname="hebo", fontsize=13)
                        y += 13 * LINE

                # Who it's really going to, when that isn't the store itself (installers, Attn names)
                receiver, contact = _clean(info.get('receiver')), _clean(info.get('contact'))
                if installer or (receiver and _norm(receiver) != _norm(store)) or contact:
                    to = ("Installer: " if installer else "Deliver to: ") + (receiver or store or '')
                    to += f" · Attn {contact}" if contact else ''
                    for line in _wrap(to, "hebo", 9.5, zone_b.width)[:2]:
                        page.insert_text((zone_b.x0, y + 9.5), line, fontname="hebo", fontsize=9.5,
                                         color=(0.75, 0.1, 0.1) if installer else (0, 0, 0))
                        y += 9.5 * LINE

                address = info.get('address')
                if not address:
                    addr_parts = [_clean(group_data.get(k)) for k in ('address_1', 'address_2', 'suburb', 'state', 'postcode', 'country')]
                    address = ", ".join(p for p in addr_parts if p)
                if address:
                    address_lines = _wrap(address, "helv", 9, zone_b.width)
                    room = int((zone_b.y1 - CHIP_H - 4 - y) // (9 * LINE))
                    for line in address_lines[:max(0, min(3, room))]:
                        page.insert_text((zone_b.x0, y + 9), line, fontname="helv", fontsize=9)
                        y += 9 * LINE

                chip_y = zone_b.y1 - (11 * LINE + 8)
                chip = _counter_chip(page, zone_b.x0, chip_y, label_text)
                _counter_chip(page, chip.x1 + 6, chip_y, page_text)

                grid_x0 = zone_a_rect.x1 + GAP
                cols = FIRST_COLS
            else:
                footer_y = A4_H - MARGIN - CHIP_H
                x = MARGIN
                if store and store.lower() != 'none':
                    page.insert_text((x, footer_y + 4 + 12 * 0.92), store, fontname="hebo", fontsize=12)
                    x += fitz.get_text_length(store, fontname="hebo", fontsize=12) + 10
                chip = _counter_chip(page, x, footer_y, label_text)
                _counter_chip(page, chip.x1 + 6, footer_y, page_text)

                grid_x0 = MARGIN
                cols = NEXT_COLS


            for r in range(ROWS):
                for c in range(cols):
                    if current_item_idx >= total_items: break
                    cx0 = grid_x0 + c * (CELL_W + GAP)
                    cy0 = MARGIN + r * (CELL_H + GAP)
                    cell = fitz.Rect(cx0, cy0, cx0 + CELL_W, cy0 + CELL_H)
                    _round_rect(page, cell, CELL_RADIUS, color=(0, 0, 0), width=0.75)
                    _draw_cell(page, items[current_item_idx], attribute_order, cell)
                    current_item_idx += 1
            page_num += 1
            if total_items == 0: break
        page_map[pack_key] = list(range(first_page, len(doc)))

    report_pages = draw_report_pages(doc, report, at=0) if report else 0
    if report_pages:
        page_map = {k: [p + report_pages for p in v] for k, v in page_map.items()}
    doc.save(output_pdf_path)
    doc.close()
    return {
        'pages': page_map,
        'report_pages': report_pages,
        'page_size_mm': [round(A4_W / MM2PT, 1), round(A4_H / MM2PT, 1)],
        'courier_region_mm': {'x': round(MARGIN / MM2PT, 1), 'y': round(MARGIN / MM2PT, 1),
                              'width': round(ZONE_A_W / MM2PT, 1), 'height': round(ZONE_A_H / MM2PT, 1)},
    }
